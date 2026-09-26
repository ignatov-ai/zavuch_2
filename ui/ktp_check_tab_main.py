# -*- coding: utf-8 -*-
import os
import sys
from datetime import datetime
from PySide6.QtWidgets import *
from PySide6.QtCore import QThread, Signal, Qt
from PySide6.QtGui import QColor

from collector_ktp_main import KTPMainCollector


class KTPMainLoadThread(QThread):
    """Поток для загрузки списка групп основного расписания"""
    finished = Signal(list)
    error = Signal(str)
    log_message = Signal(str)

    def __init__(self, auth, academic_year_id=14):
        super().__init__()
        self.auth = auth
        self.academic_year_id = academic_year_id

    def run(self):
        try:
            collector = KTPMainCollector(self.auth)
            collector.log_callback = lambda text: self.log_message.emit(text)
            groups = collector.get_all_groups()

            # Добавляем URL журнала для каждой группы
            for group in groups:
                group_id = group.get('id')
                class_unit_id = group.get('class_unit_id', '')

                if group_id and class_unit_id:
                    group['journal_url'] = (
                        f"https://dnevnik.mos.ru/manage/journal"
                        f"?from=journals&group_id={group_id}&class_unit_id={class_unit_id}"
                    )
                elif group_id:
                    group['journal_url'] = (
                        f"https://dnevnik.mos.ru/manage/journal?group_id={group_id}"
                    )
                else:
                    group['journal_url'] = ''

            self.finished.emit(groups)
        except Exception as e:
            import traceback
            error_text = f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}"
            self.error.emit(error_text)
            self.finished.emit([])


class KTPMainCheckThread(QThread):
    """Поток для проверки КТП основного расписания"""
    progress_update = Signal(int, int, str)
    row_checked = Signal(dict)
    finished = Signal(int)
    error = Signal(str)
    log_message = Signal(str)

    def __init__(self, auth, groups_to_check, academic_year_id=14):
        super().__init__()
        self.auth = auth
        self.groups_to_check = groups_to_check
        self.academic_year_id = academic_year_id
        self._is_running = True

    def stop(self):
        self._is_running = False

    def run(self):
        try:
            collector = KTPMainCollector(self.auth)
            collector.log_callback = lambda text: self.log_message.emit(text)

            total = len(self.groups_to_check)
            checked = 0

            for i, group in enumerate(self.groups_to_check):
                if not self._is_running:
                    self.log_message.emit("[KTP-MAIN] ⏹️ Остановка проверки...")
                    break

                group_id = group.get('id')
                group_name = group.get('name', '')

                self.progress_update.emit(i + 1, total, group_name)

                try:
                    has_ktp, lessons_with_names, total_lessons = collector.has_ktp(
                        group_id, self.academic_year_id
                    )

                    group['has_ktp'] = has_ktp
                    group['lessons_with_names'] = lessons_with_names
                    group['total_lessons'] = total_lessons
                    group['is_checked'] = True

                    self.row_checked.emit(group.copy())
                    checked += 1
                except Exception as e:
                    self.log_message.emit(f"[KTP-MAIN] ❌ Ошибка для {group_name}: {e}")
                    group['is_checked'] = True
                    group['has_ktp'] = False
                    group['lessons_with_names'] = 0
                    group['total_lessons'] = 0
                    self.row_checked.emit(group.copy())

            self.finished.emit(checked)

        except Exception as e:
            import traceback
            error_text = f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}"
            self.error.emit(error_text)
            self.finished.emit(0)


class KTPMainCheckTab(QWidget):
    """Вкладка проверки КТП основного расписания (ОЧ+ФЧ)"""

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.auth = None
        self.academic_year_id = 14
        self.all_groups = []
        self.filtered_groups = []
        self.initUI()

    def on_academic_year_updated(self, aid):
        self.academic_year_id = aid
        if hasattr(self, 'year_spin'):
            self.year_spin.setValue(aid)

    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(10)
        main_layout.setContentsMargins(15, 15, 15, 15)

        # === ПАНЕЛЬ УПРАВЛЕНИЯ ===
        control_group = QGroupBox("Управление")
        control_layout = QHBoxLayout(control_group)

        control_layout.addWidget(QLabel("Учебный год:"))
        self.year_spin = QSpinBox()
        self.year_spin.setRange(1, 99)
        self.year_spin.setValue(self.academic_year_id)
        self.year_spin.setMaximumWidth(80)
        self.year_spin.setToolTip("13 = 2025-2026, 14 = 2026-2027")
        control_layout.addWidget(self.year_spin)

        control_layout.addSpacing(20)

        self.load_btn = QPushButton("📋 1. Загрузить список групп")
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(32)
        self.load_btn.clicked.connect(self.load_groups)
        control_layout.addWidget(self.load_btn)

        self.check_btn = QPushButton("🔍 2. Проверить выбранные")
        self.check_btn.setEnabled(False)
        self.check_btn.setMinimumHeight(32)
        self.check_btn.setStyleSheet("""
            QPushButton { background-color: #4CAF50; color: white;
                font-weight: bold; border-radius: 5px; }
            QPushButton:hover { background-color: #45a049; }
            QPushButton:disabled { background-color: #cccccc; color: #666666; }
        """)
        self.check_btn.clicked.connect(self.check_ktp_selected)
        control_layout.addWidget(self.check_btn)

        self.check_unchecked_btn = QPushButton("🔍 Проверить не проверенные")
        self.check_unchecked_btn.setEnabled(False)
        self.check_unchecked_btn.setMinimumHeight(32)
        self.check_unchecked_btn.setStyleSheet("""
            QPushButton { background-color: #FF9800; color: white;
                font-weight: bold; border-radius: 5px; }
            QPushButton:hover { background-color: #F57C00; }
            QPushButton:disabled { background-color: #cccccc; color: #666666; }
        """)
        self.check_unchecked_btn.clicked.connect(self.check_ktp_unchecked)
        control_layout.addWidget(self.check_unchecked_btn)

        self.check_all_btn = QPushButton("🔍 Проверить все")
        self.check_all_btn.setEnabled(False)
        self.check_all_btn.setMinimumHeight(32)
        self.check_all_btn.setStyleSheet("""
            QPushButton { background-color: #2196F3; color: white;
                font-weight: bold; border-radius: 5px; }
            QPushButton:hover { background-color: #1976D2; }
            QPushButton:disabled { background-color: #cccccc; color: #666666; }
        """)
        self.check_all_btn.clicked.connect(self.check_ktp_all)
        control_layout.addWidget(self.check_all_btn)

        self.stop_btn = QPushButton("⏹️ Остановить")
        self.stop_btn.setEnabled(False)
        self.stop_btn.setMaximumWidth(120)
        self.stop_btn.clicked.connect(self.stop_check)
        control_layout.addWidget(self.stop_btn)

        control_layout.addStretch()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("font-weight: bold;")
        control_layout.addWidget(self.stats_label)

        main_layout.addWidget(control_group)

        # === ПРОГРЕСС-БАР ===
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.progress_bar.setStyleSheet("""
            QProgressBar { border: 1px solid #cccccc; border-radius: 4px;
                text-align: center; height: 24px; }
            QProgressBar::chunk { background-color: #4CAF50; border-radius: 4px; }
        """)
        main_layout.addWidget(self.progress_bar)

        # === ФИЛЬТРЫ ===
        filter_group = QGroupBox("Фильтры")
        filter_layout = QHBoxLayout(filter_group)

        filter_layout.addWidget(QLabel("Параллель:"))
        self.level_filter = QComboBox()
        self.level_filter.addItem("Все параллели", None)
        for level in range(1, 12):
            self.level_filter.addItem(f"{level} класс", level)
        self.level_filter.setMinimumWidth(120)
        self.level_filter.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.level_filter)

        filter_layout.addWidget(QLabel("Поиск:"))
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Название группы или предмета...")
        self.search_edit.setMinimumWidth(200)
        self.search_edit.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.search_edit)

        filter_layout.addWidget(QLabel("Статус КТП:"))
        self.status_filter = QComboBox()
        self.status_filter.addItem("Все", "all")
        self.status_filter.addItem("✅ Есть КТП", "has_ktp")
        self.status_filter.addItem("❌ Нет КТП", "no_ktp")
        self.status_filter.addItem("❓ Не проверено", "not_checked")
        self.status_filter.setMinimumWidth(140)
        self.status_filter.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.status_filter)

        self.clear_filters_btn = QPushButton("Очистить")
        self.clear_filters_btn.clicked.connect(self.clear_filters)
        filter_layout.addWidget(self.clear_filters_btn)

        filter_layout.addStretch()
        main_layout.addWidget(filter_group)

        # === ПАНЕЛЬ ВЫБОРА ===
        select_group = QGroupBox("Выбор групп")
        select_layout = QHBoxLayout(select_group)

        self.select_all_btn = QPushButton("✓ Выбрать всё")
        self.select_all_btn.setEnabled(False)
        self.select_all_btn.clicked.connect(self.select_all)
        select_layout.addWidget(self.select_all_btn)

        self.deselect_all_btn = QPushButton("✗ Снять всё")
        self.deselect_all_btn.setEnabled(False)
        self.deselect_all_btn.clicked.connect(self.deselect_all)
        select_layout.addWidget(self.deselect_all_btn)

        select_layout.addSpacing(20)

        self.select_visible_btn = QPushButton("✓ Выбрать видимые")
        self.select_visible_btn.setEnabled(False)
        self.select_visible_btn.clicked.connect(self.select_visible)
        select_layout.addWidget(self.select_visible_btn)

        self.select_unchecked_btn = QPushButton("✓ Выбрать не проверенные")
        self.select_unchecked_btn.setEnabled(False)
        self.select_unchecked_btn.clicked.connect(self.select_unchecked)
        select_layout.addWidget(self.select_unchecked_btn)

        select_layout.addStretch()

        self.selected_count_label = QLabel("Выбрано: 0")
        self.selected_count_label.setStyleSheet("font-weight: bold; color: #2196F3;")
        select_layout.addWidget(self.selected_count_label)

        main_layout.addWidget(select_group)

        # === ТАБЛИЦА ===
        self.groups_table = QTableWidget()
        self.groups_table.setAlternatingRowColors(False)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.groups_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.groups_table.horizontalHeader().setStretchLastSection(True)
        self.groups_table.setColumnCount(8)
        self.groups_table.setHorizontalHeaderLabels([
            '✓', 'Группа', 'Предмет',
            'Параллель', 'Класс(ы)',
            'Учеников', 'КТП', 'Статус'
        ])
        self.groups_table.verticalHeader().setDefaultSectionSize(70)
        self.groups_table.doubleClicked.connect(self.on_row_double_clicked)
        main_layout.addWidget(self.groups_table)

        # === ЭКСПОРТ ===
        export_layout = QHBoxLayout()
        export_layout.addStretch()
        self.export_btn = QPushButton("📊 Экспорт в Excel")
        self.export_btn.setEnabled(False)
        self.export_btn.clicked.connect(self.export_to_excel)
        export_layout.addWidget(self.export_btn)
        main_layout.addLayout(export_layout)

        self.set_controls_enabled(False)

    def set_controls_enabled(self, enabled):
        self.level_filter.setEnabled(enabled)
        self.search_edit.setEnabled(enabled)
        self.status_filter.setEnabled(enabled)
        self.clear_filters_btn.setEnabled(enabled)
        self.select_all_btn.setEnabled(enabled)
        self.deselect_all_btn.setEnabled(enabled)
        self.select_visible_btn.setEnabled(enabled)
        self.select_unchecked_btn.setEnabled(enabled)
        self.export_btn.setEnabled(enabled and bool(self.all_groups))

    # ==================== ЭТАП 1: ЗАГРУЗКА ====================

    def load_groups(self):
        if not self.auth:
            QMessageBox.warning(self, "Ошибка", "Сначала авторизуйтесь на вкладке 'Настройки'")
            return

        self.load_btn.setEnabled(False)
        self.check_btn.setEnabled(False)
        self.check_all_btn.setEnabled(False)
        self.check_unchecked_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFormat("Загрузка списка групп...")
        self.groups_table.setRowCount(0)
        self.all_groups = []

        academic_year_id = self.year_spin.value()
        self.academic_year_id = academic_year_id

        self.load_thread = KTPMainLoadThread(self.auth, academic_year_id)
        self.load_thread.finished.connect(self.on_load_finished)
        self.load_thread.error.connect(self.on_load_error)
        self.load_thread.log_message.connect(self.on_log_message)
        self.load_thread.start()

    def on_load_finished(self, groups):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)

        if not groups:
            QMessageBox.warning(self, "Внимание", "Не удалось получить список групп.")
            return

        for group in groups:
            group['is_checked'] = False
            group['has_ktp'] = None
            group['lessons_with_names'] = 0
            group['total_lessons'] = 0
            group['is_selected'] = False

        self.all_groups = groups
        self.apply_filter()
        self.update_stats()
        self.set_controls_enabled(True)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)

        QMessageBox.information(
            self, "Готово",
            f"✅ Загружено групп: {len(groups)}\n\n"
            f"Выберите группы и нажмите 'Проверить выбранные'"
        )

    def on_load_error(self, error):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)
        QMessageBox.critical(self, "Ошибка загрузки", f"Ошибка:\n{error}")

    # ==================== ЭТАП 2: ПРОВЕРКА ====================

    def check_ktp_selected(self):
        selected = [g for g in self.all_groups if g.get('is_selected')]
        if not selected:
            QMessageBox.warning(self, "Ошибка", "Не выбрано ни одной группы.")
            return
        self._start_check(selected, "выбранных")

    def check_ktp_unchecked(self):
        unchecked = [g for g in self.all_groups if not g.get('is_checked')]
        if not unchecked:
            QMessageBox.information(self, "Внимание", "Все группы уже проверены!")
            return
        self._start_check(unchecked, "не проверенных")

    def check_ktp_all(self):
        if not self.all_groups:
            QMessageBox.warning(self, "Ошибка", "Сначала загрузите список групп")
            return
        self._start_check(self.all_groups, "всех")

    def _start_check(self, groups, label):
        reply = QMessageBox.question(
            self, "Подтверждение",
            f"Проверить КТП для {len(groups)} {label} групп?\n\nЭто может занять несколько минут.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        self.check_btn.setEnabled(False)
        self.check_all_btn.setEnabled(False)
        self.check_unchecked_btn.setEnabled(False)
        self.load_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)

        academic_year_id = self.year_spin.value()

        self.check_thread = KTPMainCheckThread(self.auth, groups, academic_year_id)
        self.check_thread.progress_update.connect(self.on_check_progress)
        self.check_thread.row_checked.connect(self.on_row_checked)
        self.check_thread.finished.connect(self.on_check_finished)
        self.check_thread.error.connect(self.on_check_error)
        self.check_thread.log_message.connect(self.on_log_message)
        self.check_thread.start()

    def stop_check(self):
        if hasattr(self, 'check_thread') and self.check_thread:
            self.check_thread.stop()

    def on_check_progress(self, current, total, name):
        if total > 0:
            percent = int(current / total * 100)
            self.progress_bar.setValue(percent)
            self.progress_bar.setFormat(f"Проверка {current}/{total}: {name[:40]} ({percent}%)")
            QApplication.processEvents()

    def on_row_checked(self, group):
        for g in self.all_groups:
            if g['id'] == group['id']:
                g.update({
                    'has_ktp': group['has_ktp'],
                    'lessons_with_names': group['lessons_with_names'],
                    'total_lessons': group['total_lessons'],
                    'is_checked': True,
                })
                break
        self._update_row_in_table(group)
        self.update_stats()

    def _update_row_in_table(self, group):
        for i, g in enumerate(self.filtered_groups):
            if g['id'] == group['id']:
                g.update({
                    'has_ktp': group['has_ktp'],
                    'lessons_with_names': group['lessons_with_names'],
                    'total_lessons': group['total_lessons'],
                    'is_checked': True,
                })

                green_bg = QColor(220, 255, 220)
                red_bg = QColor(255, 220, 220)

                ktp_item = self.groups_table.item(i, 6)
                if ktp_item:
                    ktp_item.setText(f"{group['lessons_with_names']} / {group['total_lessons']}")

                status_item = self.groups_table.item(i, 7)
                if status_item:
                    if group['has_ktp']:
                        status_item.setText("✅ Есть КТП")
                        status_item.setBackground(green_bg)
                        row_color = green_bg
                    else:
                        status_item.setText("❌ Нет КТП")
                        status_item.setBackground(red_bg)
                        row_color = red_bg

                    for col in [1, 2, 3, 4, 5, 6]:
                        item = self.groups_table.item(i, col)
                        if item:
                            item.setBackground(row_color)
                break

    def on_check_finished(self, checked_count):
        self.progress_bar.setVisible(False)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)

        self.apply_filter()
        self.update_stats()

        total_unchecked = sum(1 for g in self.all_groups if not g.get('is_checked'))
        if total_unchecked > 0:
            QMessageBox.information(
                self, "Проверка остановлена",
                f"✅ Проверено групп: {checked_count}\n"
                f"❓ Осталось не проверено: {total_unchecked}"
            )
        else:
            QMessageBox.information(self, "Проверка завершена", f"✅ Проверено групп: {checked_count}")

    def on_check_error(self, error):
        self.progress_bar.setVisible(False)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        QMessageBox.critical(self, "Ошибка проверки", f"Ошибка:\n{error}")

    def on_log_message(self, text):
        try:
            sys.__stdout__.write(text + '\n')
            sys.__stdout__.flush()
        except Exception:
            pass

    # ==================== СТАТИСТИКА И ФИЛЬТРЫ ====================

    def update_stats(self):
        total = len(self.all_groups)
        checked = sum(1 for g in self.all_groups if g.get('is_checked'))
        with_ktp = sum(1 for g in self.all_groups if g.get('has_ktp') is True)
        without_ktp = sum(1 for g in self.all_groups if g.get('has_ktp') is False)

        self.stats_label.setText(
            f"Всего: {total} | Проверено: {checked} | ✅ {with_ktp} | ❌ {without_ktp}"
        )
        self.update_selected_count()

    def update_selected_count(self):
        selected = sum(1 for g in self.all_groups if g.get('is_selected'))
        self.selected_count_label.setText(f"Выбрано: {selected}")

    def apply_filter(self):
        level_filter = self.level_filter.currentData()
        search_text = self.search_edit.text().lower().strip()
        status_filter = self.status_filter.currentData()

        filtered = []
        for group in self.all_groups:
            if level_filter is not None:
                if level_filter not in group.get('class_levels', [group.get('_class_level')]):
                    continue

            if search_text:
                searchable = (
                    group.get('name', '').lower() +
                    ' ' + group.get('activity_name', '').lower() +
                    ' ' + group.get('activity_short_name', '').lower()
                )
                if search_text not in searchable:
                    continue

            if status_filter == 'has_ktp':
                if group.get('has_ktp') is not True:
                    continue
            elif status_filter == 'no_ktp':
                if group.get('has_ktp') is not False:
                    continue
            elif status_filter == 'not_checked':
                if group.get('is_checked'):
                    continue

            filtered.append(group)

        self.filtered_groups = filtered
        self.populate_table(filtered)

    def clear_filters(self):
        self.level_filter.setCurrentIndex(0)
        self.search_edit.clear()
        self.status_filter.setCurrentIndex(0)
        self.apply_filter()

    def populate_table(self, groups):
        self.groups_table.setSortingEnabled(False)
        self.groups_table.setUpdatesEnabled(False)
        self.groups_table.setRowCount(len(groups))

        green_bg = QColor(220, 255, 220)
        red_bg = QColor(255, 220, 220)
        gray_bg = QColor(245, 245, 245)

        for i, group in enumerate(groups):
            # Колонка 0: чекбокс
            cb_widget = QWidget()
            cb_layout = QHBoxLayout(cb_widget)
            cb_layout.setContentsMargins(0, 0, 0, 0)
            cb_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
            checkbox = QCheckBox()
            checkbox.setChecked(group.get('is_selected', False))
            checkbox.stateChanged.connect(lambda state, g=group: self.on_checkbox_changed(state, g))
            cb_layout.addWidget(checkbox)
            self.groups_table.setCellWidget(i, 0, cb_widget)

            # Колонка 1: Название группы
            name_item = QTableWidgetItem(group.get('name', ''))
            name_item.setToolTip(f"ID группы: {group['id']}")
            self.groups_table.setItem(i, 1, name_item)

            # Колонка 2: Предмет
            self.groups_table.setItem(i, 2, QTableWidgetItem(group.get('activity_name', '')))

            # Колонка 3: Параллель
            class_level = group.get('_class_level') or group.get('class_level_id')
            parallel_text = str(class_level) if class_level else ''
            parallel_item = QTableWidgetItem(parallel_text)
            parallel_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 3, parallel_item)

            # Колонка 4: Класс(ы)
            class_names = group.get('class_unit_names', [])
            if not class_names:
                cu_name = group.get('class_unit_name')
                if cu_name:
                    class_names = [cu_name]

            class_text = '\n'.join(sorted(class_names)) if class_names else ''
            class_item = QTableWidgetItem(class_text)
            class_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 4, class_item)

            # Колонка 5: Учеников
            students_item = QTableWidgetItem(str(group.get('student_count', 0)))
            students_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 5, students_item)

            # Колонка 6: КТП
            if group.get('is_checked'):
                ktp_text = f"{group['lessons_with_names']} / {group['total_lessons']}"
            else:
                ktp_text = "—"
            ktp_item = QTableWidgetItem(ktp_text)
            ktp_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 6, ktp_item)

            # Колонка 7: Статус
            has_ktp = group.get('has_ktp')
            is_checked = group.get('is_checked', False)

            if not is_checked:
                status_item = QTableWidgetItem("❓ Не проверено")
                status_item.setBackground(gray_bg)
                row_color = gray_bg
            elif has_ktp:
                status_item = QTableWidgetItem("✅ Есть КТП")
                status_item.setBackground(green_bg)
                row_color = green_bg
            else:
                status_item = QTableWidgetItem("❌ Нет КТП")
                status_item.setBackground(red_bg)
                row_color = red_bg

            self.groups_table.setItem(i, 7, status_item)

            for col in [1, 2, 3, 4, 5, 6]:
                item = self.groups_table.item(i, col)
                if item:
                    item.setBackground(row_color)

        self.groups_table.setUpdatesEnabled(True)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setColumnWidth(0, 40)
        self.groups_table.setColumnWidth(1, 250)
        self.groups_table.setColumnWidth(2, 200)
        self.groups_table.setColumnWidth(3, 90)
        self.groups_table.setColumnWidth(4, 130)
        self.groups_table.setColumnWidth(5, 80)
        self.groups_table.setColumnWidth(6, 100)
        self.groups_table.setColumnWidth(7, 130)

    def on_checkbox_changed(self, state, group):
        group['is_selected'] = (state == Qt.CheckState.Checked.value)
        self.update_selected_count()

    # ==================== ВЫБОР ====================

    def select_all(self):
        for group in self.all_groups:
            group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def deselect_all(self):
        for group in self.all_groups:
            group['is_selected'] = False
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def select_visible(self):
        for group in self.filtered_groups:
            group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def select_unchecked(self):
        for group in self.all_groups:
            if not group.get('is_checked'):
                group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def on_row_double_clicked(self, index):
        row = index.row()
        if row < 0 or row >= len(self.filtered_groups):
            return
        group = self.filtered_groups[row]
        url = group.get('journal_url', '')
        if url:
            try:
                import webbrowser
                webbrowser.open(url)
            except Exception as e:
                QMessageBox.warning(self, "Ошибка", f"Не удалось открыть браузер:\n{e}")

    # ==================== ЭКСПОРТ ====================

    def export_to_excel(self):
        if not self.missing_data:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Ошибка", "Модуль openpyxl не установлен")
            return

        class_name = self.selected_class["name"] if self.selected_class else "unknown"
        student_name = self.selected_student_name.replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Пропуски_{class_name}_{student_name}_{timestamp}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить Excel", filename, "Excel files (*.xlsx)"
        )

        if not file_path:
            return

        try:
            from collections import defaultdict

            wb = Workbook()

            # ==========================================================
            #  ЛИСТ 1: Детализация пропусков
            # ==========================================================
            ws = wb.active
            ws.title = "Детализация"

            headers = ["Класс", "ФИО ученика", "Дата", "Предмет",
                       "Учитель", "Тема урока", "Причина пропуска"]
            ws.append(headers)

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50",
                                      fill_type="solid")
            header_alignment = Alignment(horizontal='center', vertical='center',
                                         wrap_text=True)

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for item in self.missing_data:
                ws.append([
                    self.selected_class["name"] if self.selected_class else "",
                    self.selected_student_name,
                    item["date"],
                    item["subject"],
                    item["teacher"],
                    item["topic"],
                    item["reason"]
                ])

            col_widths = {'A': 12, 'B': 35, 'C': 12, 'D': 25,
                          'E': 25, 'F': 40, 'G': 20}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = 'A2'

            # ==========================================================
            #  ЛИСТ 2: Сводка по предметам
            # ==========================================================
            ws2 = wb.create_sheet("Сводка по предметам")

            # Заголовок листа — информация об ученике и периоде
            ws2.cell(row=1, column=1, value="Сводка пропусков по предметам")
            ws2.cell(row=1, column=1).font = Font(bold=True, size=13)
            ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

            ws2.cell(row=2, column=1,
                     value=f"Класс: {self.selected_class['name'] if self.selected_class else '-'}")
            ws2.cell(row=2, column=1).font = Font(bold=True)
            ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)

            ws2.cell(row=3, column=1, value=f"Ученик: {self.selected_student_name}")
            ws2.cell(row=3, column=1).font = Font(bold=True)
            ws2.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2)

            period_start = self.start_date_edit.date().toString("dd.MM.yyyy")
            period_end = self.end_date_edit.date().toString("dd.MM.yyyy")
            ws2.cell(row=4, column=1, value=f"Период: {period_start} — {period_end}")
            ws2.cell(row=4, column=1).font = Font(bold=True)
            ws2.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)

            ws2.cell(row=5, column=1, value=f"Всего пропусков: {len(self.missing_data)}")
            ws2.cell(row=5, column=1).font = Font(bold=True, color="C00000")
            ws2.merge_cells(start_row=5, start_column=1, end_row=5, end_column=2)

            # Заголовки таблицы сводки
            summary_headers = ["Предмет", "Количество пропусков"]
            summary_header_row = 7

            for col_idx, header in enumerate(summary_headers, start=1):
                cell = ws2.cell(row=summary_header_row, column=col_idx,
                                value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            # Считаем пропуски по предметам
            subjects_count = defaultdict(int)
            for item in self.missing_data:
                subjects_count[item["subject"]] += 1

            # Сортируем: сначала по убыванию количества, потом по алфавиту
            sorted_subjects = sorted(
                subjects_count.items(),
                key=lambda x: (-x[1], x[0])
            )

            row_idx = summary_header_row + 1
            for subject, count in sorted_subjects:
                ws2.cell(row=row_idx, column=1, value=subject)
                ws2.cell(row=row_idx, column=2, value=count)
                row_idx += 1

            # Итоговая строка
            ws2.cell(row=row_idx, column=1, value="ИТОГО")
            ws2.cell(row=row_idx, column=1).font = Font(bold=True)
            ws2.cell(row=row_idx, column=2, value=len(self.missing_data))
            ws2.cell(row=row_idx, column=2).font = Font(bold=True)

            # Границы для таблицы сводки
            from openpyxl.styles import Border, Side
            thin = Side(style='thin', color="999999")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for r in range(summary_header_row, row_idx + 1):
                for c in (1, 2):
                    cell = ws2.cell(row=r, column=c)
                    cell.border = border
                    if r > summary_header_row:
                        cell.alignment = Alignment(
                            horizontal='left' if c == 1 else 'center',
                            vertical='center'
                        )

            # Ширина колонок
            ws2.column_dimensions['A'].width = 45
            ws2.column_dimensions['B'].width = 22

            # ==========================================================
            #  СОХРАНЯЕМ
            # ==========================================================
            wb.save(file_path)
            self.log_signal.emit(f"✅ Excel экспортирован: {file_path}")
            self.log_signal.emit(
                f"   📄 Лист 1: Детализация ({len(self.missing_data)} записей)"
            )
            self.log_signal.emit(
                f"   📄 Лист 2: Сводка по предметам ({len(sorted_subjects)} предметов)"
            )

            reply = QMessageBox.question(
                self, "Готово",
                f"Файл сохранён:\n{file_path}\n\n"
                f"• Лист «Детализация»: {len(self.missing_data)} пропусков\n"
                f"• Лист «Сводка по предметам»: {len(sorted_subjects)} предметов\n\n"
                "Открыть файл?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)

        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка экспорта: {str(e)}")
            QMessageBox.critical(self, "Ошибка",
                                 f"Не удалось сохранить файл:\n{str(e)}")