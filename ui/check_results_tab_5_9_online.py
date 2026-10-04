# -*- coding: utf-8 -*-
"""
Вкладка «Проверка итогов (5-11) Online».
Данные тянутся из API dnevnik.mos.ru без скачивания Excel-журналов.
"""
from collections import defaultdict
from datetime import datetime
import os
from PySide6.QtWidgets import *
from PySide6.QtCore import *
from PySide6.QtGui import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Маппинги периодов аттестации
PERIOD_MAPPING_TRIMESTERS = {83558: 'Т1', 83559: 'Т2', 83560: 'Т3'}
PERIOD_MAPPING_SEMESTERS = {83561: 'П1', 83562: 'П2'}


# ============================================================
# Единый стиль (зелёный акцент — «online»)
# ============================================================
TAB_STYLE = """
    QWidget {
        font-size: 10pt;
        color: #1e293b;
    }
    QGroupBox {
        font-size: 11pt;
        font-weight: 700;
        color: #059669;
        border: 1px solid #6ee7b7;
        border-radius: 10px;
        margin-top: 12px;
        padding: 18px 14px 14px 14px;
        background: #ffffff;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 12px;
        padding: 0 6px;
    }
    QLineEdit {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QLineEdit:focus { border: 1px solid #059669; }
    QPushButton {
        min-height: 30px;
        padding: 4px 12px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #f8fafc;
        color: #1e293b;
        font-weight: 500;
    }
    QPushButton:hover { background: #e2e8f0; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
    QComboBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QTableWidget {
        border: 1px solid #e2e8f0;
        border-radius: 6px;
        gridline-color: #e2e8f0;
        alternate-background-color: #f8fafc;
    }
    QTableWidget::item { padding: 3px 6px; }
    QHeaderView::section {
        background: #f1f5f9;
        color: #1e293b;
        padding: 5px;
        border: none;
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
        font-weight: 600;
    }
    QProgressBar {
        min-height: 22px;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        background: #f8fafc;
        text-align: center;
        color: #1e293b;
        font-weight: 600;
    }
    QProgressBar::chunk {
        background-color: #059669;
        border-radius: 5px;
    }
"""

PRIMARY_BTN = """
    QPushButton {
        background-color: #059669;
        color: #ffffff;
        font-weight: 700;
        border: none;
        border-radius: 7px;
    }
    QPushButton:hover { background-color: #047857; }
    QPushButton:pressed { background-color: #065f46; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

SUCCESS_BTN = """
    QPushButton {
        background-color: #0284c7;
        color: #ffffff;
        font-weight: 600;
        border: none;
        border-radius: 6px;
    }
    QPushButton:hover { background-color: #0369a1; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""


# ============================================================
# Поток: загрузка групп параллели
# ============================================================
class OnlineGroupsLoadThread(QThread):
    finished = Signal(list)
    error = Signal(str)
    log = Signal(str)

    def __init__(self, auth, class_level):
        super().__init__()
        self.auth = auth
        self.class_level = class_level

    def run(self):
        try:
            self.log.emit(f"[Online] Загружаю группы параллели {self.class_level}...")
            data = self.auth.fetch(
                "https://dnevnik.mos.ru/jersey/api/groups",
                {
                    "academic_year_id": self.auth.aid,
                    "class_level_id": self.class_level,
                    "pid": self.auth.pid,
                    "with_lessons_only": "false",
                },
            ) or []
            self.log.emit(f"[Online] Получено групп: {len(data)}")
            self.finished.emit(data)
        except Exception as e:
            import traceback
            self.error.emit(f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}")
            self.finished.emit([])


# ============================================================
# Поток: загрузка учеников + отметок для группы
# ============================================================
class OnlineStudentsMarksLoadThread(QThread):
    finished = Signal(dict)
    error = Signal(str)
    log = Signal(str)

    def __init__(self, auth, group_id, class_unit_id):
        super().__init__()
        self.auth = auth
        self.group_id = group_id
        self.class_unit_id = class_unit_id

    def run(self):
        try:
            self.log.emit(f"[Online] Загружаю учеников группы {self.group_id}...")
            students = self.auth.fetch(
                "https://dnevnik.mos.ru/core/api/student_profiles",
                {
                    'academic_year_id': self.auth.aid,
                    'class_unit_ids': self.class_unit_id,
                    'group_ids': self.group_id,
                    'per_page': 1000,
                    'pid': self.auth.pid,
                    'with_archived_groups': 'true',
                    'with_deleted': 'true',
                    'with_final_marks': 'true',
                    'with_groups': 'true',
                    'with_home_based': 'true',
                },
            ) or []
            self.log.emit(f"[Online] Получено учеников: {len(students)}")

            self.log.emit(f"[Online] Загружаю отметки группы {self.group_id}...")
            all_marks = []
            page = 1
            per_page = 1000

            now = datetime.now()
            if now.month >= 9:
                start_date = f"01.09.{now.year}"
                end_date = f"31.08.{now.year + 1}"
            else:
                start_date = f"01.09.{now.year - 1}"
                end_date = f"31.08.{now.year}"

            while True:
                page_data = self.auth.fetch(
                    "https://dnevnik.mos.ru/core/api/marks",
                    {
                        'created_at_from': start_date,
                        'created_at_to': end_date,
                        'group_ids': self.group_id,
                        'per_page': per_page,
                        'page': page,
                        'pid': self.auth.pid,
                    },
                ) or []
                if not page_data:
                    break
                all_marks.extend(page_data)
                if len(page_data) < per_page:
                    break
                page += 1

            self.log.emit(f"[Online] Получено отметок: {len(all_marks)}")
            self.finished.emit({"students": students, "marks": all_marks})
        except Exception as e:
            import traceback
            self.error.emit(f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}")
            self.finished.emit({"students": [], "marks": []})


# ============================================================
# Вкладка
# ============================================================
class CheckResultsTab5_9Online(QWidget):
    """Проверка итоговых отметок 5-11 онлайн + средние баллы по периодам."""
    log_signal = Signal(str)

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.auth = None
        self.all_groups_by_level = defaultdict(list)
        self.current_level_groups = []
        self.current_class_groups = []
        self.all_results = []
        self._is_loading = False
        self.initUI()
        self.log_signal.connect(self.append_to_status)

    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(16, 14, 16, 14)

        # Заголовок
        title = QLabel("🎯  Проверка итогов Online (5-11)")
        title.setStyleSheet(
            "font-size: 16pt; font-weight: 700; color: #0f172a; padding: 2px 0 4px 0;"
        )
        main_layout.addWidget(title)

        subtitle = QLabel(
            "Данные загружаются напрямую из ЭЖД МЭШ. "
            "Выберите параллель, класс и журнал."
        )
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet("color: #64748b; margin-bottom: 4px;")
        main_layout.addWidget(subtitle)

        # === Панель управления ===
        settings_group = QGroupBox("Онлайн-проверка итогов")
        settings_layout = QGridLayout(settings_group)
        settings_layout.setVerticalSpacing(10)
        settings_layout.setHorizontalSpacing(12)

        settings_layout.addWidget(QLabel("Параллель:"), 0, 0)
        self.level_combo = QComboBox()
        self.level_combo.setMinimumWidth(160)
        self.level_combo.addItem("Выберите параллель", None)
        for level in range(5, 12):
            self.level_combo.addItem(f"{level} класс", level)
        self.level_combo.currentIndexChanged.connect(self.on_level_changed)
        settings_layout.addWidget(self.level_combo, 0, 1)

        self.load_groups_btn = QPushButton("📋  1. Загрузить журналы параллели")
        self.load_groups_btn.setMinimumHeight(38)
        self.load_groups_btn.setEnabled(False)
        self.load_groups_btn.setStyleSheet("""
            QPushButton {
                background: #ecfdf5; color: #059669;
                border: 1px solid #a7f3d0; font-weight: 600;
            }
            QPushButton:hover { background: #d1fae5; }
            QPushButton:disabled {
                background: #f1f5f9; color: #94a3b8;
                border-color: #e2e8f0;
            }
        """)
        self.load_groups_btn.clicked.connect(self.load_groups)
        settings_layout.addWidget(self.load_groups_btn, 0, 2, 1, 2)

        settings_layout.addWidget(QLabel("Класс:"), 1, 0)
        self.class_combo = QComboBox()
        self.class_combo.setMinimumWidth(220)
        self.class_combo.addItem("— выберите класс —", None)
        self.class_combo.setEnabled(False)
        self.class_combo.currentIndexChanged.connect(self.on_class_changed)
        settings_layout.addWidget(self.class_combo, 1, 1)

        settings_layout.addWidget(QLabel("Журнал:"), 1, 2)
        self.group_combo = QComboBox()
        self.group_combo.setMinimumWidth(400)
        self.group_combo.addItem("— выберите журнал —", None)
        self.group_combo.setEnabled(False)
        self.group_combo.currentIndexChanged.connect(self._update_check_btn)
        settings_layout.addWidget(self.group_combo, 1, 3)

        self.check_btn = QPushButton("🔍  2. Проверить итоговые отметки")
        self.check_btn.setEnabled(False)
        self.check_btn.setMinimumHeight(42)
        self.check_btn.setStyleSheet(PRIMARY_BTN + "QPushButton { font-size: 11pt; }")
        self.check_btn.clicked.connect(self.check_results)
        settings_layout.addWidget(self.check_btn, 2, 0, 1, 2)

        self.export_btn = QPushButton("📊  Экспорт в Excel")
        self.export_btn.setEnabled(False)
        self.export_btn.setStyleSheet(SUCCESS_BTN)
        self.export_btn.clicked.connect(self.export_to_excel)
        settings_layout.addWidget(self.export_btn, 2, 2, 1, 2)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        settings_layout.addWidget(self.progress_bar, 3, 0, 1, 4)

        main_layout.addWidget(settings_group)

        # === Фильтры ===
        filter_group = QGroupBox("🔎  Фильтры")
        filter_group.setStyleSheet("""
            QGroupBox {
                font-size: 11pt; font-weight: 700;
                color: #0369a1; border: 1px solid #7dd3fc;
                border-radius: 10px; margin-top: 12px;
                padding: 18px 14px 14px 14px; background: #ffffff;
            }
            QGroupBox::title {
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
            }
        """)
        filter_layout = QHBoxLayout(filter_group)
        filter_layout.setSpacing(10)
        filter_layout.setContentsMargins(10, 6, 10, 10)

        filter_layout.addWidget(QLabel("Ученик:"))
        self.filter_student = QLineEdit()
        self.filter_student.setPlaceholderText("ФИО")
        self.filter_student.setFixedWidth(180)
        self.filter_student.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_student)

        filter_layout.addWidget(QLabel("Статус:"))
        self.filter_status = QComboBox()
        self.filter_status.addItems([
            "Все статусы", "✅ СОВПАДАЕТ", "❌ НЕ СОВПАДАЕТ",
            "⚠️ НЕТ ГПА", "⚠️ ДОЛГ (А/З)", "⚠️ НПА", "⚠️ НЕТ ДАННЫХ",
        ])
        self.filter_status.setFixedWidth(180)
        self.filter_status.currentTextChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_status)

        self.clear_filters_btn = QPushButton("Очистить")
        self.clear_filters_btn.setStyleSheet("""
            QPushButton {
                background: #fff1f2; color: #be123c;
                border: 1px solid #fecdd3; font-weight: 600;
            }
            QPushButton:hover { background: #ffe4e6; }
        """)
        self.clear_filters_btn.clicked.connect(self.clear_filters)
        filter_layout.addWidget(self.clear_filters_btn)
        filter_layout.addStretch()

        main_layout.addWidget(filter_group)

        # === Таблица ===
        self.results_table = QTableWidget()
        self.results_table.setAlternatingRowColors(True)
        self.results_table.setSortingEnabled(False)
        self.results_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.results_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.results_table.horizontalHeader().setStretchLastSection(True)
        main_layout.addWidget(self.results_table, 1)

        # === Статус ===
        self.status_label = QLabel("Готов к работе")
        self.status_label.setStyleSheet("""
            color: #065f46; background: #ecfdf5;
            border-radius: 6px; padding: 7px 10px;
            font-weight: 600;
        """)
        main_layout.addWidget(self.status_label)

        self.setStyleSheet(TAB_STYLE)

    # ================================================================
    # Служебные методы
    # ================================================================
    def append_to_status(self, text):
        self.status_label.setText(str(text)[:200])
        QApplication.processEvents()

    # ================================================================
    # Авторизация
    # ================================================================
    def on_auth_updated(self, auth):
        self.auth = auth
        self.load_groups_btn.setEnabled(bool(auth))

    def on_academic_year_updated(self, aid):
        self.all_groups_by_level.clear()
        self.current_level_groups = []
        self.current_class_groups = []
        self.class_combo.clear()
        self.class_combo.addItem("— выберите класс —", None)
        self.group_combo.clear()
        self.group_combo.addItem("— выберите журнал —", None)
        self.class_combo.setEnabled(False)
        self.group_combo.setEnabled(False)
        self._update_check_btn()
        if self.auth:
            self.auth.aid = str(aid)
            self.auth.curr_aid = str(aid)

    # ================================================================
    # Загрузка параллели
    # ================================================================
    def on_level_changed(self, index):
        self.current_level_groups = []
        self.current_class_groups = []
        self.class_combo.clear()
        self.class_combo.addItem("— выберите класс —", None)
        self.class_combo.setEnabled(False)
        self.group_combo.clear()
        self.group_combo.addItem("— выберите журнал —", None)
        self.group_combo.setEnabled(False)
        self._update_check_btn()

        level = self.level_combo.currentData()
        if level and level in self.all_groups_by_level:
            self._populate_classes_from_groups(self.all_groups_by_level[level])

    def load_groups(self):
        if not self.auth:
            QMessageBox.warning(self, "Ошибка", "Нет авторизации.")
            return
        level = self.level_combo.currentData()
        if not level:
            QMessageBox.warning(self, "Ошибка", "Выберите параллель.")
            return
        if self._is_loading:
            return

        self._is_loading = True
        self.load_groups_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFormat("Загрузка списка журналов...")

        self.groups_thread = OnlineGroupsLoadThread(self.auth, level)
        self.groups_thread.log.connect(self.append_to_status)
        self.groups_thread.finished.connect(
            lambda g: self.on_groups_loaded(g, level)
        )
        self.groups_thread.error.connect(self.on_load_error)
        self.groups_thread.start()

    def on_groups_loaded(self, groups, level):
        self._is_loading = False
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_groups_btn.setEnabled(True)

        if not groups:
            QMessageBox.warning(self, "Внимание", "Список групп пуст.")
            return

        valid = []
        for g in groups:
            if g.get("is_metagroup") or g.get("is_ndo"):
                continue
            if not g.get("class_unit_id"):
                continue
            valid.append(g)

        if not valid:
            QMessageBox.warning(self, "Внимание", "Не найдено подходящих групп.")
            return

        self.all_groups_by_level[level] = valid
        self.current_level_groups = valid
        self._populate_classes_from_groups(valid)

        QMessageBox.information(
            self, "Готово",
            f"Загружено групп: {len(valid)}\n"
            f"Теперь выберите класс и предмет."
        )

    def _populate_classes_from_groups(self, groups):
        classes = {}
        for g in groups:
            cu_id = g.get("class_unit_id")
            cu_name = g.get("class_unit_name") or ""
            if cu_id and cu_id not in classes:
                classes[cu_id] = cu_name

        self.class_combo.clear()
        self.class_combo.addItem("— выберите класс —", None)
        sorted_classes = sorted(classes.items(), key=lambda x: (x[1] or ""))
        for cu_id, cu_name in sorted_classes:
            self.class_combo.addItem(cu_name or f"ID {cu_id}", cu_id)

        self.class_combo.setEnabled(bool(classes))
        if not classes:
            self.append_to_status("⚠️ В параллели нет классов.")
        else:
            self.append_to_status(f"Найдено классов: {len(classes)}")

    # ================================================================
    # Выбор класса → журналы
    # ================================================================
    def on_class_changed(self, index):
        cu_id = self.class_combo.currentData()
        self.group_combo.clear()
        self.group_combo.addItem("— выберите журнал —", None)
        self.group_combo.setEnabled(False)
        self._update_check_btn()

        if not cu_id:
            return

        class_groups = [
            g for g in self.current_level_groups
            if g.get("class_unit_id") == cu_id
        ]
        if not class_groups:
            self.append_to_status("⚠️ Для класса нет журналов.")
            return

        class_groups.sort(
            key=lambda x: (x.get("subject_name") or "", x.get("name") or "")
        )
        for g in class_groups:
            subj = g.get("subject_name") or ""
            gname = g.get("name") or ""
            self.group_combo.addItem(f"{subj} — {gname}", g)

        self.current_class_groups = class_groups
        self.group_combo.setEnabled(True)
        self.append_to_status(f"Найдено журналов: {len(class_groups)}")

    def _update_check_btn(self):
        has_group = self.group_combo.currentData() is not None
        self.check_btn.setEnabled(has_group and not self._is_loading)

    def on_load_error(self, err):
        self._is_loading = False
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_groups_btn.setEnabled(True)
        QMessageBox.critical(self, "Ошибка загрузки", err)

    # ================================================================
    # Проверка
    # ================================================================
    def check_results(self):
        if self._is_loading:
            return
        group = self.group_combo.currentData()
        if not group:
            QMessageBox.warning(self, "Ошибка", "Выберите журнал.")
            return

        group_id = group.get("id")
        class_unit_id = group.get("class_unit_id")
        if not group_id or not class_unit_id:
            QMessageBox.warning(self, "Ошибка", "У группы нет class_unit_id.")
            return

        self._is_loading = True
        self.check_btn.setEnabled(False)
        self.export_btn.setEnabled(False)
        self.results_table.clear()
        self.results_table.setRowCount(0)
        self.results_table.setColumnCount(0)
        self.all_results = []

        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFormat("Загрузка учеников и отметок...")

        self.data_thread = OnlineStudentsMarksLoadThread(
            self.auth, group_id, class_unit_id
        )
        self.data_thread.log.connect(self.append_to_status)
        self.data_thread.finished.connect(
            lambda payload: self.on_data_loaded(payload, group)
        )
        self.data_thread.error.connect(self.on_load_error)
        self.data_thread.start()

    def on_data_loaded(self, payload, group):
        try:
            students = payload.get("students") or []
            marks = payload.get("marks") or []

            if not students:
                QMessageBox.warning(self, "Внимание", "Нет учеников в группе.")
                return

            class_name = group.get("class_unit_name") or ""
            subject_name = group.get("subject_name") or ""
            class_level_id = group.get("class_level_id")

            if class_level_id in (10, 11):
                period_mapping = PERIOD_MAPPING_SEMESTERS
                is_high = True
            else:
                period_mapping = PERIOD_MAPPING_TRIMESTERS
                is_high = False

            marks_by_student = defaultdict(list)
            for m in marks:
                sid = m.get("student_profile_id")
                if sid is not None:
                    marks_by_student[sid].append(m)

            for student in students:
                sid = student.get("id")
                self._process_student(
                    student,
                    marks_by_student.get(sid, []),
                    class_name, subject_name,
                    period_mapping, is_high,
                )

            self.display_all_results()
            self.clear_filters()
            if self.all_results:
                self.export_btn.setEnabled(True)

            self.append_to_status(
                f"Учеников: {len(students)}, отметок: {len(marks)}, "
                f"строк: {len(self.all_results)}"
            )
        finally:
            self._is_loading = False
            self.check_btn.setEnabled(True)
            self.progress_bar.setVisible(False)
            self.progress_bar.setRange(0, 100)

    # ================================================================
    # Обработка одного ученика
    # ================================================================
    def _process_student(self, student, student_marks,
                         class_name, subject_name,
                         period_mapping, is_high_school):
        final_marks_by_period = {}
        gpa_mark = ''

        for mark in student.get('final_marks', []):
            period_id = mark.get('attestation_period_id')
            value_obj = mark.get('value', {})
            mark_type = mark.get('mark_type', '')
            if isinstance(value_obj, dict):
                value = value_obj.get('parsedValue') or value_obj.get('source')
            else:
                value = value_obj

            if value == 'А/З' or (isinstance(value, str) and 'А/З' in value):
                mark_value = 'А/З'
            elif value is not None:
                try:
                    if isinstance(value, (int, float)) or (
                        isinstance(value, str) and value.replace('.', '').isdigit()
                    ):
                        mark_value = str(int(float(value)))
                    else:
                        mark_value = str(value)
                except Exception:
                    mark_value = str(value)
            else:
                continue

            if mark_type == 'intermediate_attestation':
                gpa_mark = mark_value
                continue

            if period_id and period_id in period_mapping:
                final_marks_by_period[period_mapping[period_id]] = mark_value
            if mark.get('is_year_mark'):
                final_marks_by_period['Год'] = mark_value

        period_avgs = self._compute_period_averages(student_marks, period_mapping)
        fio = self._get_student_full_name(student)

        if is_high_school:
            result = self._process_result_high(
                class_name, fio, subject_name,
                final_marks_by_period.get('П1'),
                final_marks_by_period.get('П2'),
                gpa_mark,
                final_marks_by_period.get('Год'),
                period_avgs.get('П1'),
                period_avgs.get('П2'),
            )
        else:
            result = self._process_result_middle(
                class_name, fio, subject_name,
                final_marks_by_period.get('Т1'),
                final_marks_by_period.get('Т2'),
                final_marks_by_period.get('Т3'),
                gpa_mark,
                final_marks_by_period.get('Год'),
                period_avgs.get('Т1'),
                period_avgs.get('Т2'),
                period_avgs.get('Т3'),
            )

        if result:
            self.all_results.append(result)

    def _compute_period_averages(self, student_marks, period_mapping):
        periods = {}
        now = datetime.now()
        if now.month >= 9:
            y1 = now.year
            y2 = now.year + 1
        else:
            y1 = now.year - 1
            y2 = now.year

        for pid, name in period_mapping.items():
            if name == 'Т1':
                periods[name] = (datetime(y1, 9, 1), datetime(y1, 11, 30))
            elif name == 'Т2':
                periods[name] = (datetime(y1, 12, 1), datetime(y2, 2, 28))
            elif name == 'Т3':
                periods[name] = (datetime(y2, 3, 1), datetime(y2, 5, 31))
            elif name == 'П1':
                periods[name] = (datetime(y1, 9, 1), datetime(y1, 12, 31))
            elif name == 'П2':
                periods[name] = (datetime(y2, 1, 1), datetime(y2, 5, 31))

        period_values = defaultdict(list)
        for m in student_marks:
            date_str = m.get("date")
            name = m.get("name")
            weight = m.get("weight") or 1
            if not date_str or name is None:
                continue
            try:
                d, mo, y = date_str.split(".")
                dt = datetime(int(y), int(mo), int(d))
            except Exception:
                continue

            pname = None
            for pn, (start, end) in periods.items():
                if start <= dt <= end:
                    pname = pn
                    break
            if not pname:
                continue

            try:
                mark_int = int(name)
            except Exception:
                continue
            try:
                weight_int = int(weight)
            except Exception:
                weight_int = 1
            period_values[pname].append((mark_int, weight_int))

        result = {}
        for pname, lst in period_values.items():
            total_sum = sum(m * w for m, w in lst)
            total_w = sum(w for _, w in lst)
            if total_w > 0:
                result[pname] = round(total_sum / total_w, 2)
        return result

    def _get_student_full_name(self, student):
        last = student.get("last_name", "")
        first = student.get("first_name", "")
        middle = student.get("middle_name", "")
        if not last and student.get("user_name"):
            parts = student.get("user_name", "").split()
            if len(parts) >= 1:
                last = parts[0]
            if len(parts) >= 2:
                first = parts[1]
            if len(parts) >= 3:
                middle = parts[2]
        parts = [p for p in (last, first, middle) if p]
        return " ".join(parts) if parts else student.get("short_name", "")

    # ================================================================
    # Логика проверки
    # ================================================================
    def _safe_to_int_mark(self, val):
        if val is None:
            return None
        if isinstance(val, str):
            v = val.strip()
            if v in ('А/З', 'НПА'):
                return v
            if v in ('', '-'):
                return None
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return None

    def _get_expected_period_mark(self, avg_float):
        if avg_float >= 4.6:
            return 5
        elif avg_float >= 3.6:
            return 4
        elif avg_float >= 2.6:
            return 3
        else:
            return 'А/З'

    def _get_expected_year_mark(self, avg_float):
        if avg_float >= 4.5:
            return 5
        elif avg_float >= 3.5:
            return 4
        elif avg_float >= 2.5:
            return 3
        else:
            return 2

    def _process_result_middle(self, class_name, student_name, subject_name,
                               t1_val, t2_val, t3_val, gpa_val, year_val,
                               avg_t1, avg_t2, avg_t3):
        t1 = self._safe_to_int_mark(t1_val)
        t2 = self._safe_to_int_mark(t2_val)
        t3 = self._safe_to_int_mark(t3_val)
        gpa = self._safe_to_int_mark(gpa_val)
        year = self._safe_to_int_mark(year_val)
        return self._build_result(
            class_name, student_name, subject_name,
            period_fields=('Т1', 'Т2', 'Т3'),
            period_values=(t1, t2, t3),
            period_avgs=(avg_t1, avg_t2, avg_t3),
            gpa=gpa, year=year,
            raw_period_values=(t1_val, t2_val, t3_val),
            gpa_raw=gpa_val, year_raw=year_val,
        )

    def _process_result_high(self, class_name, student_name, subject_name,
                             p1_val, p2_val, gpa_val, year_val,
                             avg_p1, avg_p2):
        p1 = self._safe_to_int_mark(p1_val)
        p2 = self._safe_to_int_mark(p2_val)
        gpa = self._safe_to_int_mark(gpa_val)
        year = self._safe_to_int_mark(year_val)
        return self._build_result(
            class_name, student_name, subject_name,
            period_fields=('П1', 'П2'),
            period_values=(p1, p2),
            period_avgs=(avg_p1, avg_p2),
            gpa=gpa, year=year,
            raw_period_values=(p1_val, p2_val),
            gpa_raw=gpa_val, year_raw=year_val,
        )

    def _build_result(self, class_name, student_name, subject_name,
                      period_fields, period_values, period_avgs,
                      gpa, year, raw_period_values, gpa_raw, year_raw):
        missing_period = any(v is None for v in period_values)
        has_debt = has_npa = missing_gpa = False

        if gpa is None:
            missing_gpa = True
        elif isinstance(gpa, str) and gpa == 'А/З':
            has_debt = True
        elif isinstance(gpa, str) and gpa == 'НПА':
            has_npa = True

        valid_values = [v for v in period_values if isinstance(v, int)]
        if isinstance(gpa, int):
            valid_values.append(gpa)

        if missing_period:
            status_text, status_type = '⚠️ НЕТ ДАННЫХ', 'missing_data'
            calculated = calculated_mark = '-'
        elif missing_gpa:
            status_text, status_type = '⚠️ НЕТ ГПА', 'missing_gpa'
            calculated = calculated_mark = '-'
        elif has_debt:
            status_text, status_type = '⚠️ ДОЛГ (А/З)', 'debt'
            calculated = calculated_mark = '-'
        elif has_npa:
            status_text, status_type = '⚠️ НПА', 'npa'
            calculated = calculated_mark = '-'
        elif not valid_values:
            status_text, status_type = '⚠️ НЕТ ДАННЫХ', 'missing_data'
            calculated = calculated_mark = '-'
        else:
            calculated = round(sum(valid_values) / len(valid_values), 2)
            calculated_mark = self._get_expected_year_mark(calculated)
            if isinstance(year, int):
                if calculated_mark != year:
                    status_text, status_type = '❌ НЕ СОВПАДАЕТ', 'mismatch'
                else:
                    status_text, status_type = '✅ СОВПАДАЕТ', 'match'
            else:
                status_text, status_type = '⚠️ НЕТ ГОДОВОЙ', 'missing_data'
                calculated_mark = '-'

        result = {
            'class': class_name,
            'student': student_name,
            'subject': subject_name,
            'gpa': str(gpa_raw) if gpa_raw is not None else '-',
            'calculated': str(calculated) if calculated != '-' else '-',
            'calculated_mark': str(calculated_mark) if calculated_mark != '-' else '-',
            'actual_mark': str(year_raw) if year_raw is not None else '-',
            'status': status_text,
            'status_type': status_type,
        }

        for name, raw, avg in zip(period_fields, raw_period_values, period_avgs):
            result[f"{name.lower()}_avg"] = f"{avg:.2f}" if avg is not None else '-'
            result[f"{name.lower()}_mark"] = str(raw) if raw is not None else '-'
            if avg is not None and raw is not None:
                expected = self._get_expected_period_mark(avg)
                try:
                    actual_int = int(float(raw))
                except (ValueError, TypeError):
                    actual_int = None
                if isinstance(expected, int) and actual_int is not None and expected != actual_int:
                    result[f"{name.lower()}_mismatch"] = True
                else:
                    result[f"{name.lower()}_mismatch"] = False
            else:
                result[f"{name.lower()}_mismatch"] = False

        return result

    # ================================================================
    # Таблица
    # ================================================================
    def display_all_results(self):
        if self.all_results:
            self._safe_update_table(self.all_results)

    def _safe_update_table(self, results):
        try:
            self.results_table.setUpdatesEnabled(False)
            self.results_table.setSortingEnabled(False)
            if not results:
                self.results_table.setRowCount(0)
                self.results_table.setColumnCount(0)
                return

            is_high = 'p1_mark' in results[0]
            if is_high:
                headers = [
                    'Класс', 'Ученик', 'Предмет',
                    'ср. П1', 'П1', 'ср. П2', 'П2',
                    'ГПА', 'ср. балл', 'ГОД должно', 'ГОД выст.', 'Статус',
                ]
                period_keys = ('p1', 'p2')
            else:
                headers = [
                    'Класс', 'Ученик', 'Предмет',
                    'ср. Т1', 'Т1', 'ср. Т2', 'Т2', 'ср. Т3', 'Т3',
                    'ГПА', 'ср. балл', 'ГОД должно', 'ГОД выст.', 'Статус',
                ]
                period_keys = ('t1', 't2', 't3')

            self.results_table.setColumnCount(len(headers))
            self.results_table.setHorizontalHeaderLabels(headers)
            self.results_table.setRowCount(len(results))

            red_bg = QColor(255, 180, 180)
            for i, r in enumerate(results):
                col = 0
                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('class', '')))
                ); col += 1
                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('student', '')))
                ); col += 1
                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('subject', '')))
                ); col += 1

                for pk in period_keys:
                    avg_item = QTableWidgetItem(str(r.get(f'{pk}_avg', '-')))
                    self.results_table.setItem(i, col, avg_item); col += 1

                    mark_item = QTableWidgetItem(str(r.get(f'{pk}_mark', '-')))
                    if r.get(f'{pk}_mismatch'):
                        mark_item.setBackground(red_bg)
                    self.results_table.setItem(i, col, mark_item); col += 1

                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('gpa', '-')))
                ); col += 1
                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('calculated', '-')))
                ); col += 1
                self.results_table.setItem(
                    i, col, QTableWidgetItem(str(r.get('calculated_mark', '-')))
                ); col += 1

                actual_item = QTableWidgetItem(str(r.get('actual_mark', '-')))
                cm = str(r.get('calculated_mark', '-'))
                am = str(r.get('actual_mark', '-'))
                if cm != '-' and am != '-':
                    try:
                        if int(cm) != int(am):
                            actual_item.setBackground(red_bg)
                    except Exception:
                        pass
                self.results_table.setItem(i, col, actual_item); col += 1

                status_item = QTableWidgetItem(str(r.get('status', '')))
                st = r.get('status_type', '')
                if st == 'match':
                    status_item.setBackground(QColor(209, 250, 229))
                    status_item.setForeground(QColor(6, 95, 70))
                elif st == 'mismatch':
                    status_item.setBackground(QColor(254, 202, 202))
                    status_item.setForeground(QColor(153, 27, 27))
                elif st == 'missing_gpa':
                    status_item.setBackground(QColor(254, 240, 138))
                    status_item.setForeground(QColor(133, 77, 14))
                elif st in ('debt', 'npa'):
                    status_item.setBackground(QColor(254, 215, 170))
                    status_item.setForeground(QColor(124, 45, 18))
                elif st == 'missing_data':
                    status_item.setBackground(QColor(226, 232, 240))
                    status_item.setForeground(QColor(30, 41, 59))
                self.results_table.setItem(i, col, status_item)

            self.results_table.resizeColumnsToContents()
            if self.results_table.columnCount() > 2:
                self.results_table.setColumnWidth(0, 70)
                self.results_table.setColumnWidth(1, 220)
                self.results_table.setColumnWidth(2, 200)
        finally:
            self.results_table.setUpdatesEnabled(True)

    # ================================================================
    # Фильтры
    # ================================================================
    def apply_filter(self):
        if not self.all_results or self._is_loading:
            return
        student_f = self.filter_student.text().lower()
        status_f = self.filter_status.currentText()

        filtered = []
        for r in self.all_results:
            if student_f and student_f not in str(r.get('student', '')).lower():
                continue
            if status_f != "Все статусы" and r.get('status', '') != status_f:
                continue
            filtered.append(r)

        self._safe_update_table(filtered)
        self.append_to_status(
            f"Показано: {len(filtered)} из {len(self.all_results)} записей"
        )

    def clear_filters(self):
        self.filter_student.clear()
        self.filter_status.setCurrentIndex(0)
        self.display_all_results()

    # ================================================================
    # Экспорт
    # ================================================================
    def export_to_excel(self):
        if not self.all_results:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        is_high = 'p1_mark' in self.all_results[0]
        filename = (
            f"Проверка_итогов_online_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт", filename, "Excel files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Проверка итогов Online"

            if is_high:
                headers = [
                    'Класс', 'Ученик', 'Предмет',
                    'ср. П1', 'П1', 'ср. П2', 'П2',
                    'ГПА', 'ср. балл', 'ГОД должно', 'ГОД выст.', 'Статус',
                ]
                period_keys = ('p1', 'p2')
            else:
                headers = [
                    'Класс', 'Ученик', 'Предмет',
                    'ср. Т1', 'Т1', 'ср. Т2', 'Т2', 'ср. Т3', 'Т3',
                    'ГПА', 'ср. балл', 'ГОД должно', 'ГОД выст.', 'Статус',
                ]
                period_keys = ('t1', 't2', 't3')

            ws.append(headers)

            header_font = Font(bold=True, color="FFFFFF", size=10)
            header_fill = PatternFill(
                start_color="059669", end_color="059669", fill_type="solid"
            )
            header_align = Alignment(
                horizontal='center', vertical='center', wrap_text=True
            )
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col)
                c.font = header_font
                c.fill = header_fill
                c.alignment = header_align

            red_fill = PatternFill(
                start_color="FFB4B4", end_color="FFB4B4", fill_type="solid"
            )
            for r in self.all_results:
                row = [
                    r.get('class', ''), r.get('student', ''), r.get('subject', ''),
                ]
                for pk in period_keys:
                    row.append(r.get(f'{pk}_avg', '-'))
                    row.append(r.get(f'{pk}_mark', '-'))
                row.append(r.get('gpa', '-'))
                row.append(r.get('calculated', '-'))
                row.append(r.get('calculated_mark', '-'))
                row.append(r.get('actual_mark', '-'))
                row.append(r.get('status', ''))
                ws.append(row)

            actual_col_idx = len(headers) - 1
            calculated_col_idx = actual_col_idx - 1
            for row_idx in range(2, ws.max_row + 1):
                cm = ws.cell(row=row_idx, column=calculated_col_idx).value
                am = ws.cell(row=row_idx, column=actual_col_idx).value
                if cm and cm != '-' and am and am != '-':
                    try:
                        if int(str(cm)) != int(str(am)):
                            ws.cell(row=row_idx, column=actual_col_idx).fill = red_fill
                    except Exception:
                        pass

            last_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_letter}{ws.max_row}"
            ws.freeze_panes = 'A2'
            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 25

            wb.save(file_path)
            self.append_to_status(f"Отчёт сохранён: {file_path}")

            reply = QMessageBox.question(
                self, "Готово",
                f"Отчёт сохранён:\n{file_path}\n\nОткрыть файл?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить:\n{e}")