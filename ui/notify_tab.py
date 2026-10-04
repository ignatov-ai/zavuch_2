# -*- coding: utf-8 -*-
"""
Вкладка «📨 Уведомления родителям».
Генерация уведомлений в формате Word на основе данных из журналов.
"""
import os
import re
from collections import defaultdict
from datetime import datetime, timedelta
from PySide6.QtWidgets import *
from PySide6.QtCore import *
from PySide6.QtGui import *
from ui.console import ConsoleWidget
from academic_calendar import AcademicCalendar
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# Чекбокс ученика
# ============================================================
class StudentCheckBox(QCheckBox):
    """Чекбокс с данными ученика."""

    def __init__(self, student_name, student_data):
        super().__init__(student_name)
        self.student_data = student_data
        self.setStyleSheet("""
            QCheckBox {
                spacing: 7px;
                padding: 3px 4px;
                color: #334155;
            }
            QCheckBox::indicator {
                width: 17px;
                height: 17px;
            }
            QCheckBox:hover {
                color: #0d9488;
            }
        """)


# ============================================================
# Единый стиль (teal — уведомления)
# ============================================================
TAB_STYLE = """
    QWidget {
        font-size: 10pt;
        color: #1e293b;
    }
    QGroupBox {
        font-size: 11pt;
        font-weight: 700;
        color: #0d9488;
        border: 1px solid #5eead4;
        border-radius: 10px;
        margin-top: 12px;
        padding: 18px 14px 14px 14px;
        background: #ffffff;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 12px;
        padding: 0 6px;
    }
    QLineEdit {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QLineEdit:focus { border: 1px solid #0d9488; }
    QLineEdit:disabled {
        background: #f1f5f9;
        color: #94a3b8;
    }
    QPushButton {
        min-height: 30px;
        padding: 4px 12px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #f8fafc;
        color: #1e293b;
        font-weight: 500;
    }
    QPushButton:hover { background: #e2e8f0; }
    QPushButton:pressed { background: #cbd5e1; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
    QComboBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QSpinBox, QDoubleSpinBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QDateEdit {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QDateEdit:focus { border: 1px solid #0d9488; }
    QCheckBox {
        color: #334155;
        spacing: 6px;
    }
    QCheckBox::indicator {
        width: 16px;
        height: 16px;
    }
    QScrollArea {
        border: 1px solid #e2e8f0;
        background: #fafafa;
        border-radius: 6px;
    }
    QLabel { color: #334155; }
"""

PRIMARY_BTN = """
    QPushButton {
        background-color: #0d9488;
        color: #ffffff;
        font-weight: 700;
        border: none;
        border-radius: 7px;
    }
    QPushButton:hover { background-color: #0f766e; }
    QPushButton:pressed { background-color: #115e59; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""


# ============================================================
# Вкладка
# ============================================================
class NotifyTab(QWidget):
    """Компактная вкладка создания уведомлений для родителей."""
    log_signal = Signal(str)

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.journals_by_level = defaultdict(list)
        self.selected_folder = ""
        self.students_by_class = {}
        self.current_class = None
        self.notification_counter = 1
        self.curriculum_data = {}
        self.initUI()
        self.log_signal.connect(self.append_to_console)

    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(16, 14, 16, 14)

        # === Заголовок ===
        title = QLabel("📨  Уведомления родителям")
        title.setStyleSheet(
            "font-size: 16pt; font-weight: 700; color: #0f172a; padding: 2px 0 6px 0;"
        )
        main_layout.addWidget(title)

        subtitle = QLabel(
            "Генерация уведомлений в формате Word на основе данных из журналов."
        )
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet("color: #64748b; margin-bottom: 4px;")
        main_layout.addWidget(subtitle)

        # === Группа «Источники данных» ===
        sources_group = QGroupBox("📁  Источники данных")
        sources_layout = QGridLayout(sources_group)
        sources_layout.setVerticalSpacing(10)
        sources_layout.setHorizontalSpacing(12)

        sources_layout.addWidget(QLabel("📁 Журналы:"), 0, 0)
        folder_widget = QWidget()
        folder_layout = QHBoxLayout(folder_widget)
        folder_layout.setContentsMargins(0, 0, 0, 0)
        folder_layout.setSpacing(6)
        self.folder_edit = QLineEdit()
        self.folder_edit.setPlaceholderText("Путь к папке с журналами")
        self.folder_edit.textChanged.connect(self.on_folder_changed)
        folder_layout.addWidget(self.folder_edit, 1)
        self.folder_browse_btn = QPushButton("Обзор…")
        self.folder_browse_btn.setFixedWidth(90)
        self.folder_browse_btn.clicked.connect(self.browse_folder)
        folder_layout.addWidget(self.folder_browse_btn)
        sources_layout.addWidget(folder_widget, 0, 1)

        sources_layout.addWidget(QLabel("📚 Учебный план:"), 1, 0)
        curriculum_widget = QWidget()
        curriculum_layout = QHBoxLayout(curriculum_widget)
        curriculum_layout.setContentsMargins(0, 0, 0, 0)
        curriculum_layout.setSpacing(6)
        self.curriculum_edit = QLineEdit()
        self.curriculum_edit.setPlaceholderText("Сводная_таблица_УП.xlsx")
        self.curriculum_edit.setText("Сводная_таблица_УП_5-11_классы.xlsx")
        curriculum_layout.addWidget(self.curriculum_edit, 1)
        self.curriculum_browse_btn = QPushButton("Обзор…")
        self.curriculum_browse_btn.setFixedWidth(90)
        self.curriculum_browse_btn.clicked.connect(self.browse_curriculum)
        curriculum_layout.addWidget(self.curriculum_browse_btn)
        sources_layout.addWidget(curriculum_widget, 1, 1)

        main_layout.addWidget(sources_group)

        # === Группа «Период» ===
        period_group = QGroupBox("📅  Период уведомления")
        period_group.setStyleSheet("""
            QGroupBox {
                font-size: 11pt; font-weight: 700;
                color: #0369a1; border: 1px solid #7dd3fc;
                border-radius: 10px; margin-top: 12px;
                padding: 18px 14px 14px 14px; background: #ffffff;
            }
            QGroupBox::title {
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
            }
        """)
        period_layout = QHBoxLayout(period_group)
        period_layout.setSpacing(10)

        period_layout.addWidget(QLabel("Период:"))
        self.period_combo = QComboBox()
        self.period_combo.addItems([
            "Триместр 1 (сен-ноя)",
            "Триместр 2 (дек-фев)",
            "Триместр 3 (мар-май)",
            "Полугодие 1 (сен-дек)",
            "Полугодие 2 (янв-май)",
            "Произвольный",
        ])
        self.period_combo.setCurrentIndex(1)
        self.period_combo.setMinimumWidth(220)
        self.period_combo.currentIndexChanged.connect(self.on_period_changed)
        period_layout.addWidget(self.period_combo)

        period_layout.addWidget(QLabel("с:"))
        self.start_date_edit = QDateEdit()
        self.start_date_edit.setDate(QDate(2026, 1, 1))
        self.start_date_edit.setCalendarPopup(True)
        self.start_date_edit.setDisplayFormat("dd.MM.yyyy")
        self.start_date_edit.setEnabled(False)
        self.start_date_edit.setFixedWidth(110)
        period_layout.addWidget(self.start_date_edit)

        period_layout.addWidget(QLabel("по:"))
        self.end_date_edit = QDateEdit()
        self.end_date_edit.setDate(QDate(2026, 5, 31))
        self.end_date_edit.setCalendarPopup(True)
        self.end_date_edit.setDisplayFormat("dd.MM.yyyy")
        self.end_date_edit.setEnabled(False)
        self.end_date_edit.setFixedWidth(110)
        period_layout.addWidget(self.end_date_edit)

        period_layout.addStretch()
        main_layout.addWidget(period_group)

        # === ГОРИЗОНТАЛЬНЫЙ СПЛИТТЕР ===
        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setChildrenCollapsible(False)
        splitter.setStyleSheet("""
            QSplitter::handle {
                background-color: #e2e8f0;
                width: 2px;
            }
        """)

        # === ЛЕВАЯ ПАНЕЛЬ: ученики ===
        left_widget = QWidget()
        left_layout = QVBoxLayout(left_widget)
        left_layout.setContentsMargins(0, 0, 8, 0)
        left_layout.setSpacing(8)

        left_title = QLabel("👥  Ученики")
        left_title.setStyleSheet(
            "font-size: 11pt; font-weight: 700; color: #0d9488; padding: 2px 0;"
        )
        left_layout.addWidget(left_title)

        class_layout = QHBoxLayout()
        class_layout.setSpacing(6)
        class_layout.addWidget(QLabel("Класс:"))
        self.class_combo = QComboBox()
        self.class_combo.setMinimumWidth(150)
        self.class_combo.currentIndexChanged.connect(self.on_class_changed)
        class_layout.addWidget(self.class_combo)

        self.select_all_btn = QPushButton("✓ Все")
        self.select_all_btn.setFixedWidth(60)
        self.select_all_btn.setStyleSheet("""
            QPushButton {
                background: #f0fdfa; color: #0d9488;
                border: 1px solid #99f6e4; font-weight: 600;
            }
            QPushButton:hover { background: #ccfbf1; }
        """)
        self.select_all_btn.clicked.connect(self.select_all_students)
        class_layout.addWidget(self.select_all_btn)

        self.deselect_all_btn = QPushButton("✗ Снять")
        self.deselect_all_btn.setFixedWidth(70)
        self.deselect_all_btn.setStyleSheet("""
            QPushButton {
                background: #fff1f2; color: #be123c;
                border: 1px solid #fecdd3; font-weight: 600;
            }
            QPushButton:hover { background: #ffe4e6; }
        """)
        self.deselect_all_btn.clicked.connect(self.deselect_all_students)
        class_layout.addWidget(self.deselect_all_btn)
        class_layout.addStretch()
        left_layout.addLayout(class_layout)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setMinimumHeight(300)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.students_widget = QWidget()
        self.students_layout = QVBoxLayout(self.students_widget)
        self.students_layout.setSpacing(2)
        self.students_layout.setContentsMargins(4, 4, 4, 4)
        self.students_layout.addStretch()
        scroll.setWidget(self.students_widget)
        left_layout.addWidget(scroll, 1)

        # === ПРАВАЯ ПАНЕЛЬ: параметры ===
        right_widget = QWidget()
        right_layout = QVBoxLayout(right_widget)
        right_layout.setContentsMargins(8, 0, 0, 0)
        right_layout.setSpacing(10)

        right_title = QLabel("⚙️  Параметры уведомления")
        right_title.setStyleSheet(
            "font-size: 11pt; font-weight: 700; color: #0d9488; padding: 2px 0;"
        )
        right_layout.addWidget(right_title)

        # --- Строка 1: Дата и номер ---
        row1 = QGroupBox()
        row1.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row1_layout = QHBoxLayout(row1)
        row1_layout.setContentsMargins(10, 4, 10, 4)
        row1_layout.setSpacing(8)
        row1_layout.addWidget(QLabel("Дата:"))
        self.date_edit = QDateEdit()
        self.date_edit.setDate(QDate.currentDate())
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat("dd.MM.yyyy")
        self.date_edit.setFixedWidth(110)
        row1_layout.addWidget(self.date_edit)
        row1_layout.addWidget(QLabel("№:"))
        self.number_edit = QLineEdit()
        self.number_edit.setPlaceholderText("номер")
        self.number_edit.setFixedWidth(100)
        row1_layout.addWidget(self.number_edit)
        self.number_auto_btn = QPushButton("Авто")
        self.number_auto_btn.setFixedWidth(60)
        self.number_auto_btn.setStyleSheet("""
            QPushButton {
                background: #f0fdfa; color: #0d9488;
                border: 1px solid #99f6e4; font-weight: 600;
            }
            QPushButton:hover { background: #ccfbf1; }
        """)
        self.number_auto_btn.clicked.connect(self.generate_notification_number)
        row1_layout.addWidget(self.number_auto_btn)
        row1_layout.addStretch()
        right_layout.addWidget(row1)

        # --- Строка 2: Порог пропусков ---
        row2 = QGroupBox()
        row2.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row2_layout = QHBoxLayout(row2)
        row2_layout.setContentsMargins(10, 4, 10, 4)
        row2_layout.setSpacing(8)
        row2_layout.addWidget(QLabel("Пропуски >:"))
        self.absence_spin = QSpinBox()
        self.absence_spin.setRange(0, 100)
        self.absence_spin.setValue(20)
        self.absence_spin.setSuffix("%")
        self.absence_spin.setFixedWidth(80)
        row2_layout.addWidget(self.absence_spin)
        row2_layout.addWidget(QLabel("Комментарий:"))
        self.absence_comment = QLineEdit()
        self.absence_comment.setPlaceholderText("причина пропусков")
        self.absence_comment.setMinimumWidth(200)
        row2_layout.addWidget(self.absence_comment, 1)
        right_layout.addWidget(row2)

        # --- Строка 3: Порог среднего балла ---
        row3 = QGroupBox()
        row3.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row3_layout = QHBoxLayout(row3)
        row3_layout.setContentsMargins(10, 4, 10, 4)
        row3_layout.setSpacing(8)
        row3_layout.addWidget(QLabel("Ср.балл <:"))
        self.mark_threshold_spin = QDoubleSpinBox()
        self.mark_threshold_spin.setRange(2.0, 5.0)
        self.mark_threshold_spin.setValue(3.5)
        self.mark_threshold_spin.setSingleStep(0.1)
        self.mark_threshold_spin.setFixedWidth(80)
        row3_layout.addWidget(self.mark_threshold_spin)
        row3_layout.addStretch()
        right_layout.addWidget(row3)

        # --- Строка 4: Порог накопляемости ---
        row4 = QGroupBox()
        row4.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row4_layout = QHBoxLayout(row4)
        row4_layout.setContentsMargins(10, 4, 10, 4)
        row4_layout.setSpacing(8)
        row4_layout.addWidget(QLabel("Накопляемость <:"))
        self.accumulation_spin = QSpinBox()
        self.accumulation_spin.setRange(0, 100)
        self.accumulation_spin.setValue(70)
        self.accumulation_spin.setSuffix("%")
        self.accumulation_spin.setFixedWidth(80)
        self.accumulation_spin.setToolTip(
            "Включить предметы, где количество отметок меньше "
            "указанного процента от нормы"
        )
        row4_layout.addWidget(self.accumulation_spin)
        row4_layout.addStretch()
        right_layout.addWidget(row4)

        # --- Строка 5: Чекбоксы элементов ---
        row5 = QGroupBox()
        row5.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row5_layout = QHBoxLayout(row5)
        row5_layout.setContentsMargins(10, 4, 10, 4)
        row5_layout.setSpacing(12)
        self.include_absences = QCheckBox("Пропуски")
        self.include_absences.setChecked(True)
        row5_layout.addWidget(self.include_absences)
        self.include_az = QCheckBox("А/З")
        self.include_az.setChecked(True)
        row5_layout.addWidget(self.include_az)
        self.include_npa = QCheckBox("НПА")
        self.include_npa.setChecked(True)
        row5_layout.addWidget(self.include_npa)
        self.include_marks = QCheckBox("Средний балл")
        self.include_marks.setChecked(True)
        row5_layout.addWidget(self.include_marks)
        self.include_accumulation = QCheckBox("Накопляемость")
        self.include_accumulation.setChecked(True)
        row5_layout.addWidget(self.include_accumulation)
        self.include_low_marks = QCheckBox("Только < порога")
        self.include_low_marks.setChecked(True)
        row5_layout.addWidget(self.include_low_marks)
        row5_layout.addStretch()
        right_layout.addWidget(row5)

        # --- Строка 6: Срок ликвидации ---
        row6 = QGroupBox()
        row6.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row6_layout = QHBoxLayout(row6)
        row6_layout.setContentsMargins(10, 4, 10, 4)
        row6_layout.setSpacing(8)
        row6_layout.addWidget(QLabel("Срок ликвидации:"))
        self.deadline_date_edit = QDateEdit()
        self.deadline_date_edit.setDate(QDate.currentDate().addDays(14))
        self.deadline_date_edit.setCalendarPopup(True)
        self.deadline_date_edit.setDisplayFormat("dd.MM.yyyy")
        self.deadline_date_edit.setFixedWidth(110)
        row6_layout.addWidget(self.deadline_date_edit)
        row6_layout.addStretch()
        right_layout.addWidget(row6)

        # --- Строка 7: Руководитель ---
        row7 = QGroupBox()
        row7.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row7_layout = QHBoxLayout(row7)
        row7_layout.setContentsMargins(10, 4, 10, 4)
        row7_layout.setSpacing(8)
        row7_layout.addWidget(QLabel("Директор/Завуч:"))
        self.director_edit = QLineEdit()
        self.director_edit.setPlaceholderText("ФИО")
        self.director_edit.setMinimumWidth(200)
        row7_layout.addWidget(self.director_edit, 1)
        right_layout.addWidget(row7)

        # --- Строка 8: Папка сохранения ---
        row8 = QGroupBox()
        row8.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row8_layout = QHBoxLayout(row8)
        row8_layout.setContentsMargins(10, 4, 10, 4)
        row8_layout.setSpacing(8)
        row8_layout.addWidget(QLabel("Сохранить в:"))
        self.output_edit = QLineEdit()
        self.output_edit.setPlaceholderText("папка для уведомлений")
        row8_layout.addWidget(self.output_edit, 1)
        self.output_browse_btn = QPushButton("Обзор…")
        self.output_browse_btn.setFixedWidth(90)
        self.output_browse_btn.clicked.connect(self.browse_output_folder)
        row8_layout.addWidget(self.output_browse_btn)
        right_layout.addWidget(row8)

        # --- Строка 9: Шаблон ---
        row9 = QGroupBox()
        row9.setStyleSheet("""
            QGroupBox {
                border: 1px solid #e2e8f0;
                border-radius: 6px;
                background: #f8fafc;
                margin-top: 4px;
                padding: 8px;
            }
        """)
        row9_layout = QHBoxLayout(row9)
        row9_layout.setContentsMargins(10, 4, 10, 4)
        row9_layout.setSpacing(8)
        row9_layout.addWidget(QLabel("Шаблон:"))
        self.template_edit = QLineEdit()
        self.template_edit.setPlaceholderText("template.docx")
        row9_layout.addWidget(self.template_edit, 1)
        self.template_browse_btn = QPushButton("…")
        self.template_browse_btn.setFixedWidth(40)
        self.template_browse_btn.clicked.connect(self.browse_template)
        row9_layout.addWidget(self.template_browse_btn)
        right_layout.addWidget(row9)

        right_layout.addStretch()

        splitter.addWidget(left_widget)
        splitter.addWidget(right_widget)
        splitter.setSizes([320, 520])
        main_layout.addWidget(splitter, 1)

        # === КНОПКА СОЗДАНИЯ ===
        button_layout = QHBoxLayout()
        button_layout.setContentsMargins(0, 4, 0, 0)
        button_layout.addStretch()
        self.create_btn = QPushButton("📄  Создать уведомления")
        self.create_btn.setMinimumHeight(44)
        self.create_btn.setMaximumWidth(320)
        self.create_btn.setEnabled(False)
        self.create_btn.setStyleSheet(
            PRIMARY_BTN + "QPushButton { font-size: 12pt; }"
        )
        self.create_btn.clicked.connect(self.create_notifications)
        button_layout.addWidget(self.create_btn)
        button_layout.addStretch()
        main_layout.addLayout(button_layout)

        # === КОНСОЛЬ ===
        self.console = ConsoleWidget()
        main_layout.addWidget(self.console)

        self.setStyleSheet(TAB_STYLE)
        self.set_class_controls_enabled(False)

    def set_class_controls_enabled(self, enabled):
        self.class_combo.setEnabled(enabled)
        self.select_all_btn.setEnabled(enabled)
        self.deselect_all_btn.setEnabled(enabled)

    def on_period_changed(self, index):
        is_custom = (index == 5)
        self.start_date_edit.setEnabled(is_custom)
        self.end_date_edit.setEnabled(is_custom)

    def get_period_dates(self):
        period_index = self.period_combo.currentIndex()
        period_map = {
            0: "Т1", 1: "Т2", 2: "Т3",
            3: "П1", 4: "П2", 5: "custom",
        }
        period_type = period_map.get(period_index, "Т2")
        if period_type == "custom":
            start_qdate = self.start_date_edit.date()
            end_qdate = self.end_date_edit.date()
            start_str = start_qdate.toString("dd.MM.yyyy")
            end_str = end_qdate.toString("dd.MM.yyyy")
            start_date, end_date, period_name = AcademicCalendar.get_period_dates(
                "custom", start_str, end_str
            )
            start_date = datetime(
                start_qdate.year(), start_qdate.month(), start_qdate.day()
            )
            end_date = datetime(
                end_qdate.year(), end_qdate.month(), end_qdate.day()
            )
            period_name = "произвольный"
            return start_date, end_date, period_name
        start_date, end_date, period_name = AcademicCalendar.get_period_dates(
            period_type
        )
        self.log_signal.emit(
            f"  📅 Период: {start_date.strftime('%d.%m.%Y')} - "
            f"{end_date.strftime('%d.%m.%Y')}"
        )
        return start_date, end_date, period_name

    def generate_notification_number(self):
        current_date = QDate.currentDate()
        year = current_date.year()
        month = current_date.month()
        number = f"{year}-{month:02d}-{self.notification_counter:03d}"
        self.number_edit.setText(number)
        self.notification_counter += 1

    def browse_folder(self):
        folder = QFileDialog.getExistingDirectory(
            self, "Выберите папку с журналами"
        )
        if folder:
            self.folder_edit.setText(folder)
            self.scan_journals_folder(folder)

    def browse_curriculum(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл учебного плана", "", "Excel files (*.xlsx)"
        )
        if file_path:
            self.curriculum_edit.setText(file_path)
            self.load_curriculum_data(file_path)

    def load_curriculum_data(self, file_path):
        try:
            from openpyxl import load_workbook
            self.curriculum_data = {}
            if not os.path.exists(file_path):
                self.log_signal.emit(
                    f"⚠️ Файл учебного плана не найден: {file_path}"
                )
                return
            wb = load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                max_row = sheet.max_row
                max_col = sheet.max_column
                classes_in_sheet = []
                for col in range(2, max_col + 1):
                    cell_value = sheet.cell(row=1, column=col).value
                    if cell_value:
                        class_name = str(cell_value)
                        if len(class_name) > 1 and class_name[1].isalpha():
                            class_name = f'{class_name[0]}-{class_name[1:]}'
                        class_name = class_name.upper()
                        classes_in_sheet.append(class_name)
                for row in range(2, max_row + 1):
                    subject = sheet.cell(row=row, column=1).value
                    if not subject:
                        continue
                    subject = str(subject).strip()
                    for col_idx, class_name in enumerate(classes_in_sheet, start=2):
                        hours = sheet.cell(row=row, column=col_idx).value
                        if hours and isinstance(hours, (int, float)):
                            if class_name not in self.curriculum_data:
                                self.curriculum_data[class_name] = {}
                            self.curriculum_data[class_name][subject] = int(hours)
            wb.close()
            self.log_signal.emit(
                f"✅ Учебный план загружен: {len(self.curriculum_data)} классов"
            )
        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка загрузки учебного плана: {str(e)}")

    def browse_template(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл шаблона", "", "Word files (*.docx)"
        )
        if file_path:
            self.template_edit.setText(file_path)

    def browse_output_folder(self):
        folder = QFileDialog.getExistingDirectory(
            self, "Выберите папку для сохранения уведомлений"
        )
        if folder:
            self.output_edit.setText(folder)
            self.update_create_button()

    def on_folder_changed(self, text):
        if os.path.exists(text):
            self.scan_journals_folder(text)
            self.update_create_button()

    def scan_journals_folder(self, folder_path):
        self.journals_by_level.clear()
        self.students_by_class.clear()
        self.class_combo.clear()
        found_any = False
        for item in os.listdir(folder_path):
            item_path = os.path.join(folder_path, item)
            if os.path.isdir(item_path) and item.isdigit():
                level = int(item)
                if 1 <= level <= 11:
                    level_files = []
                    for filename in os.listdir(item_path):
                        if filename.endswith('_журнал.xlsx') and not filename.startswith('~$'):
                            class_name = filename.split('_журнал')[0].replace('_', '-')
                            level_files.append({
                                'path': os.path.join(item_path, filename),
                                'name': filename,
                                'level': level,
                                'class_name': class_name,
                            })
                    if level_files:
                        found_any = True
                        self.journals_by_level[level] = level_files
                        for journal in level_files:
                            display_name = f"{journal['class_name']} ({level})"
                            self.class_combo.addItem(display_name, journal)
        if found_any:
            self.log_signal.emit(
                f"📊 Найдено журналов: {sum(len(v) for v in self.journals_by_level.values())}"
            )
            self.set_class_controls_enabled(True)
            if self.curriculum_edit.text():
                self.load_curriculum_data(self.curriculum_edit.text())
        else:
            self.log_signal.emit("⚠️ Журналы не найдены")
            self.set_class_controls_enabled(False)
        self.update_create_button()

    def on_class_changed(self, index):
        if index < 0:
            return
        journal_data = self.class_combo.itemData(index)
        if not journal_data:
            return
        self.current_class = journal_data
        self.load_students_from_journal(journal_data['path'])

    def load_students_from_journal(self, journal_path):
        try:
            from openpyxl import load_workbook
            for i in reversed(range(self.students_layout.count())):
                item = self.students_layout.itemAt(i)
                if item and item.widget():
                    item.widget().deleteLater()
            book = load_workbook(journal_path, data_only=True)
            students_set = set()
            for sheet_name in book.sheetnames:
                sheet = book[sheet_name]
                start_row = 7
                max_row = min(sheet.max_row, 100)
                for row in range(start_row, max_row + 1):
                    student_cell = sheet.cell(row=row, column=2)
                    if student_cell and student_cell.value:
                        student_name = str(student_cell.value).strip()
                        if student_name:
                            students_set.add(student_name)
            book.close()
            sorted_students = sorted(students_set)
            for student_name in sorted_students:
                cb = StudentCheckBox(student_name, {})
                self.students_layout.addWidget(cb)
            self.students_layout.addStretch()
            self.log_signal.emit(f"👥 Загружено: {len(sorted_students)} уч.")
        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка: {str(e)}")

    def select_all_students(self):
        for i in range(self.students_layout.count()):
            widget = self.students_layout.itemAt(i).widget()
            if isinstance(widget, StudentCheckBox):
                widget.setChecked(True)
        self.update_create_button()

    def deselect_all_students(self):
        for i in range(self.students_layout.count()):
            widget = self.students_layout.itemAt(i).widget()
            if isinstance(widget, StudentCheckBox):
                widget.setChecked(False)
        self.update_create_button()

    def get_selected_students(self):
        selected = []
        for i in range(self.students_layout.count()):
            widget = self.students_layout.itemAt(i).widget()
            if isinstance(widget, StudentCheckBox) and widget.isChecked():
                selected.append(widget.text())
        return selected

    def update_create_button(self):
        has_output = bool(self.output_edit.text())
        has_students = bool(self.get_selected_students())
        has_journal = self.current_class is not None
        self.create_btn.setEnabled(has_output and has_students and has_journal)

    def append_to_console(self, text):
        self.console.append_text(text)
        QApplication.processEvents()

    def parse_date(self, day, month, year=None):
        return AcademicCalendar.parse_date(day, month, year)

    def parse_custom_date(self, date_str):
        return AcademicCalendar.parse_custom_date(date_str)

    def create_template_if_not_exists(self, template_path):
        if os.path.exists(template_path):
            return template_path
        try:
            doc = Document()
            style = doc.styles['Normal']
            style.font.name = 'Times New Roman'
            style.font.size = Pt(14)
            header = doc.add_paragraph()
            header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header_run = header.add_run(
                "ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ ОБЩЕОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ\n"
                "ГОРОДА МОСКВЫ \"ШКОЛА № 1234\""
            )
            header_run.font.size = Pt(14)
            header_run.font.bold = True
            doc.add_paragraph()
            title = doc.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_run = title.add_run("УВЕДОМЛЕНИЕ РОДИТЕЛЕЙ")
            title_run.font.size = Pt(16)
            title_run.font.bold = True
            title_run.underline = True
            doc.add_paragraph()
            doc.add_paragraph("от {DATE} г. № {NUMBER}")
            doc.add_paragraph()
            doc.add_paragraph("Кому: __________________________________")
            doc.add_paragraph(" (ФИО родителей/законных представителей)")
            doc.add_paragraph()
            doc.add_paragraph("ученика(цы) {CLASS} класса")
            doc.add_paragraph("{STUDENT_NAME}")
            doc.add_paragraph()
            doc.add_paragraph("Уважаемые родители!")
            doc.add_paragraph()
            doc.add_paragraph(
                "По состоянию на текущую дату у Вашего ребенка "
                "имеются следующие проблемы в обучении:"
            )
            doc.add_paragraph()
            doc.add_paragraph("{ABSENCES_INFO}")
            doc.add_paragraph("{ACCUMULATION_INFO}")
            doc.add_paragraph("{AZ_INFO}")
            doc.add_paragraph("{NPA_INFO}")
            doc.add_paragraph("{MARKS_INFO}")
            doc.add_paragraph()
            doc.add_paragraph("Просим принять необходимые меры.")
            doc.add_paragraph()
            doc.add_paragraph(
                "Директор школы ____________________ {DIRECTOR_NAME}"
            )
            doc.add_paragraph()
            doc.add_paragraph(
                "Классный руководитель ____________________"
            )
            doc.save(template_path)
            return template_path
        except Exception as e:
            return None

    def get_min_marks_required(self, lessons_per_week):
        if lessons_per_week == 0:
            return 0
        elif lessons_per_week == 1:
            return 3
        elif lessons_per_week == 2:
            return 5
        else:
            return 7

    def analyze_student(self, journal_path, student_name):
        from openpyxl import load_workbook
        results = {
            'absences': [],
            'az': [],
            'npa': [],
            'marks': [],
            'accumulation': [],
        }
        try:
            if not os.path.exists(journal_path):
                self.log_signal.emit("  ⚠️ Файл журнала не найден")
                return results
            book = load_workbook(journal_path, data_only=True)
            absence_threshold = self.absence_spin.value()
            mark_threshold = self.mark_threshold_spin.value()
            accumulation_threshold = self.accumulation_spin.value()
            period_start, period_end, period_name = self.get_period_dates()
            class_name = (
                self.current_class['class_name'].upper()
                if self.current_class else ""
            )
            for sheet_name in book.sheetnames:
                sheet = book[sheet_name]
                subject = "Неизвестно"
                if sheet.cell(row=1, column=1).value:
                    subject_cell = str(sheet.cell(row=1, column=1).value)
                    if "Предмет:" in subject_cell:
                        subject = subject_cell.replace("Предмет:", "").strip()
                student_row = None
                start_row = 7
                max_row = sheet.max_row
                for row in range(start_row, max_row + 1):
                    cell = sheet.cell(row=row, column=2)
                    if cell.value and str(cell.value).strip() == student_name:
                        student_row = row
                        break
                if not student_row:
                    continue
                marks_count = 0
                marks_sum = 0
                absences_count = 0
                sickness_count = 0
                lessons_in_period = 0
                lessons_per_week = 0
                if class_name in self.curriculum_data:
                    class_curriculum = self.curriculum_data[class_name]
                    subject_lower = subject.lower()
                    for curr_subject, hours in class_curriculum.items():
                        if (subject_lower == curr_subject.lower() or
                                subject_lower in curr_subject.lower() or
                                curr_subject.lower() in subject_lower):
                            lessons_per_week = hours
                            break
                for col in range(3, sheet.max_column + 1):
                    header_value = sheet.cell(row=6, column=col).value
                    if header_value == 'оц' or (
                        header_value and 'урок' in str(header_value).lower()
                    ):
                        month_val = sheet.cell(row=4, column=col).value
                        day_val = sheet.cell(row=5, column=col).value
                        if month_val and day_val:
                            try:
                                month = int(month_val)
                                day = int(day_val)
                                date_obj = self.parse_date(day, month)
                                if period_start <= date_obj <= period_end:
                                    lessons_in_period += 1
                                    cell_value = sheet.cell(
                                        row=student_row, column=col
                                    ).value
                                    if cell_value:
                                        cell_str = str(cell_value).lower()
                                        if 'н' in cell_str:
                                            absences_count += cell_str.count('н')
                                        if 'б' in cell_str:
                                            sickness_count += cell_str.count('б')
                                        parts = cell_str.split(',')
                                        for part in parts:
                                            part = part.strip()
                                            if '(' in part:
                                                try:
                                                    mark_part = part.split('(')[0].strip()
                                                    if mark_part in ['2', '3', '4', '5']:
                                                        marks_count += 1
                                                        marks_sum += int(mark_part)
                                                except Exception:
                                                    pass
                                            elif part in ['2', '3', '4', '5']:
                                                marks_count += 1
                                                marks_sum += int(part)
                            except Exception:
                                continue
                if lessons_in_period > 0:
                    absence_percent = (absences_count / lessons_in_period) * 100
                    if (self.include_absences.isChecked() and
                            absence_percent >= absence_threshold):
                        comment = self.absence_comment.text().strip()
                        if sickness_count > 0:
                            if comment:
                                results['absences'].append(
                                    f"{subject} - пропуски {absence_percent:.1f}% "
                                    f"(из них {sickness_count} по болезни) ({comment})"
                                )
                            else:
                                results['absences'].append(
                                    f"{subject} - пропуски {absence_percent:.1f}% "
                                    f"(из них {sickness_count} по болезни)"
                                )
                        else:
                            if comment:
                                results['absences'].append(
                                    f"{subject} - пропуски {absence_percent:.1f}% ({comment})"
                                )
                            else:
                                results['absences'].append(
                                    f"{subject} - пропуски {absence_percent:.1f}%"
                                )
                if (self.include_accumulation.isChecked() and
                        lessons_per_week > 0):
                    min_required = self.get_min_marks_required(lessons_per_week)
                    expected_marks = min_required
                    if expected_marks > 0:
                        percentage = (
                            (marks_count / expected_marks) * 100
                            if expected_marks > 0 else 0
                        )
                        if percentage < accumulation_threshold:
                            results['accumulation'].append(
                                f"{subject} - {marks_count} из {expected_marks} "
                                f"отметок ({percentage:.1f}%)"
                            )
                t_cols = {}
                for col in range(3, sheet.max_column + 1):
                    type_cell = sheet.cell(row=6, column=col)
                    period_cell = sheet.cell(row=4, column=col)
                    if (type_cell.value and str(type_cell.value) == 'итог' and
                            period_cell.value):
                        period = str(period_cell.value)
                        t_cols[period] = col
                for period, col in t_cols.items():
                    mark_cell = sheet.cell(row=student_row, column=col)
                    if mark_cell and mark_cell.value:
                        mark = str(mark_cell.value)
                        if mark == 'А/З' and self.include_az.isChecked():
                            results['az'].append(f"{subject} ({period})")
                        elif mark == 'НПА' and self.include_npa.isChecked():
                            results['npa'].append(f"{subject} ({period})")
                if self.include_marks.isChecked() and marks_count > 0:
                    avg_mark = marks_sum / marks_count
                    if self.include_low_marks.isChecked():
                        if avg_mark < mark_threshold:
                            results['marks'].append(
                                f"{subject} - {avg_mark:.2f}"
                            )
                    else:
                        results['marks'].append(f"{subject} - {avg_mark:.2f}")
            book.close()
        except Exception as e:
            self.log_signal.emit(f"⚠️ Ошибка анализа: {str(e)}")
            import traceback
            traceback.print_exc()
        return results

    def replace_placeholders_in_document(self, doc, replacements, sections_text):
        paragraphs_to_remove = []
        for paragraph in doc.paragraphs:
            full_text = ''.join(run.text for run in paragraph.runs)
            has_absence = '{ABSENCES_INFO}' in full_text
            has_az = '{AZ_INFO}' in full_text
            has_npa = '{NPA_INFO}' in full_text
            has_marks = '{MARKS_INFO}' in full_text
            has_accumulation = '{ACCUMULATION_INFO}' in full_text
            remove = False
            if has_absence and 'ABSENCES_INFO' not in sections_text:
                remove = True
            elif has_az and 'AZ_INFO' not in sections_text:
                remove = True
            elif has_npa and 'NPA_INFO' not in sections_text:
                remove = True
            elif has_marks and 'MARKS_INFO' not in sections_text:
                remove = True
            elif has_accumulation and 'ACCUMULATION_INFO' not in sections_text:
                remove = True
            if remove:
                paragraphs_to_remove.append(paragraph)
                continue
            for run in paragraph.runs:
                text = run.text
                new_text = text
                for placeholder, value in replacements.items():
                    if placeholder in new_text:
                        new_text = new_text.replace(placeholder, value)
                if new_text != text:
                    run.text = new_text
        for paragraph in paragraphs_to_remove:
            p = paragraph._element
            p.getparent().remove(p)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_paragraphs_to_remove = []
                    for paragraph in cell.paragraphs:
                        full_text = ''.join(run.text for run in paragraph.runs)
                        has_absence = '{ABSENCES_INFO}' in full_text
                        has_az = '{AZ_INFO}' in full_text
                        has_npa = '{NPA_INFO}' in full_text
                        has_marks = '{MARKS_INFO}' in full_text
                        has_accumulation = '{ACCUMULATION_INFO}' in full_text
                        remove = False
                        if has_absence and 'ABSENCES_INFO' not in sections_text:
                            remove = True
                        elif has_az and 'AZ_INFO' not in sections_text:
                            remove = True
                        elif has_npa and 'NPA_INFO' not in sections_text:
                            remove = True
                        elif has_marks and 'MARKS_INFO' not in sections_text:
                            remove = True
                        elif (has_accumulation and
                              'ACCUMULATION_INFO' not in sections_text):
                            remove = True
                        if remove:
                            cell_paragraphs_to_remove.append(paragraph)
                        else:
                            for run in paragraph.runs:
                                text = run.text
                                new_text = text
                                for placeholder, value in replacements.items():
                                    if placeholder in new_text:
                                        new_text = new_text.replace(
                                            placeholder, value
                                        )
                                if new_text != text:
                                    run.text = new_text
                    for paragraph in cell_paragraphs_to_remove:
                        p = paragraph._element
                        p.getparent().remove(p)
        return doc

    def create_notification_docx(self, student_name, class_name, data, output_path):
        template_path = self.template_edit.text().strip()
        if not template_path:
            template_path = "/Уведомление ШАБЛОН ГИА-11.docx"
        if not os.path.exists(template_path):
            template_path = self.create_template_if_not_exists(template_path)
            if not template_path:
                self.log_signal.emit("❌ Не удалось создать шаблон")
                return None
        try:
            doc = Document(template_path)
            notification_number = self.number_edit.text().strip()
            deadline_date = self.deadline_date_edit.date().toString("dd.MM.yyyy")
            current_date = self.date_edit.date().toString("dd.MM.yyyy")
            director_name = (
                self.director_edit.text().strip()
                if self.director_edit.text().strip() else "____________________"
            )
            absences_text = ""
            accumulation_text = ""
            az_text = ""
            npa_text = ""
            marks_text = ""
            sections_text = {}
            if self.include_absences.isChecked() and data['absences']:
                absences_text = "Пропуски предметов:\n"
                for item in data['absences']:
                    absences_text += f"• {item}\n"
                sections_text['ABSENCES_INFO'] = True
            if self.include_accumulation.isChecked() and data['accumulation']:
                accumulation_text = "Накопляемость отметок:\n"
                for item in data['accumulation']:
                    accumulation_text += f"• {item}\n"
                sections_text['ACCUMULATION_INFO'] = True
            if self.include_az.isChecked() and data['az']:
                az_text = "Академические задолженности:\n"
                for item in data['az']:
                    az_text += f"• {item}\n"
                az_text += f"\nСрок ликвидации: {deadline_date}\n"
                sections_text['AZ_INFO'] = True
            if self.include_npa.isChecked() and data['npa']:
                npa_text = "Непрохождение промежуточной аттестации:\n"
                for item in data['npa']:
                    npa_text += f"• {item}\n"
                npa_text += f"\nСрок ликвидации: {deadline_date}\n"
                sections_text['NPA_INFO'] = True
            if self.include_marks.isChecked() and data['marks']:
                marks_text = "Средние баллы по предметам:\n"
                for item in data['marks']:
                    marks_text += f"• {item}\n"
                sections_text['MARKS_INFO'] = True
            replacements = {
                '{DATE}': current_date,
                '{NUMBER}': notification_number if notification_number else "_____",
                '{CLASS}': class_name,
                '{STUDENT_NAME}': student_name,
                '{DIRECTOR_NAME}': director_name,
                '{DEADLINE_DATE}': deadline_date,
                '{ABSENCES_INFO}': absences_text,
                '{ACCUMULATION_INFO}': accumulation_text,
                '{AZ_INFO}': az_text,
                '{NPA_INFO}': npa_text,
                '{MARKS_INFO}': marks_text,
            }
            for paragraph in doc.paragraphs:
                original_text = ''.join(run.text for run in paragraph.runs)
                new_text = original_text
                for placeholder, value in replacements.items():
                    if placeholder in new_text:
                        new_text = new_text.replace(placeholder, value)
                if new_text != original_text:
                    for run in paragraph.runs:
                        run.text = ''
                    if paragraph.runs:
                        paragraph.runs[0].text = new_text
                    else:
                        paragraph.add_run(new_text)
            paragraphs_to_remove = []
            for paragraph in doc.paragraphs:
                text = ''.join(run.text for run in paragraph.runs).strip()
                if text == '>' or text == '':
                    paragraphs_to_remove.append(paragraph)
            for paragraph in paragraphs_to_remove:
                p = paragraph._element
                p.getparent().remove(p)
            os.makedirs(output_path, exist_ok=True)
            safe_name = re.sub(r'[\\/*?:"<>|]', '_', student_name)
            filename = f"Уведомление_{class_name}_{safe_name}.docx"
            filepath = os.path.join(output_path, filename)
            doc.save(filepath)
            self.log_signal.emit(f"  ✅ Уведомление сохранено: {filename}")
            return filepath
        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка: {str(e)}")
            import traceback
            traceback.print_exc()
            return None

    def create_empty_notification(self, student_name, class_name, output_path):
        template_path = self.template_edit.text().strip()
        if not template_path:
            template_path = "template.docx"
        if not os.path.exists(template_path):
            template_path = self.create_template_if_not_exists(template_path)
            if not template_path:
                self.log_signal.emit("❌ Не удалось создать шаблон")
                return None
        try:
            doc = Document(template_path)
            notification_number = self.number_edit.text().strip()
            current_date = self.date_edit.date().toString("dd.MM.yyyy")
            director_name = (
                self.director_edit.text().strip()
                if self.director_edit.text().strip() else "____________________"
            )
            replacements = {
                '{DATE}': current_date,
                '{NUMBER}': notification_number if notification_number else "_____",
                '{CLASS}': class_name,
                '{STUDENT_NAME}': student_name,
                '{DIRECTOR_NAME}': director_name,
                '{DEADLINE_DATE}': self.deadline_date_edit.date().toString("dd.MM.yyyy"),
                '{ABSENCES_INFO}': "",
                '{ACCUMULATION_INFO}': "",
                '{AZ_INFO}': "",
                '{NPA_INFO}': "",
                '{MARKS_INFO}': "",
            }
            empty_sections = {}
            doc = self.replace_placeholders_in_document(
                doc, replacements, empty_sections
            )
            for paragraph in doc.paragraphs:
                if "Уважаемые родители!" in paragraph.text:
                    new_paragraph = doc.add_paragraph()
                    new_paragraph.style = paragraph.style
                    run = new_paragraph.add_run(
                        "На текущий момент замечаний к успеваемости нет."
                    )
                    run.italic = True
                    paragraph._element.addnext(new_paragraph._element)
                    break
            os.makedirs(output_path, exist_ok=True)
            safe_name = re.sub(r'[\\/*?:"<>|]', '_', student_name)
            filename = f"Уведомление_{class_name}_{safe_name}_ЗАМЕЧАНИЙ_НЕТ.docx"
            filepath = os.path.join(output_path, filename)
            doc.save(filepath)
            return filepath
        except Exception as e:
            self.log_signal.emit(
                f"❌ Ошибка создания пустого уведомления для {student_name}: {str(e)}"
            )
            return None

    def create_notifications(self):
        selected_students = self.get_selected_students()
        if not selected_students:
            QMessageBox.warning(self, "Ошибка", "Выберите учеников!")
            return
        if not self.output_edit.text():
            QMessageBox.warning(self, "Ошибка", "Выберите папку для сохранения!")
            return
        if not self.current_class:
            QMessageBox.warning(self, "Ошибка", "Выберите класс!")
            return
        if not self.curriculum_data and self.include_accumulation.isChecked():
            reply = QMessageBox.question(
                self, "Учебный план",
                "Не загружены данные учебного плана. "
                "Накопляемость может быть рассчитана неточно. Продолжить?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.No:
                return
        self.create_btn.setEnabled(False)
        self.console.clear()
        period_start, period_end, period_name = self.get_period_dates()
        self.log_signal.emit("🚀 Начинаем создание уведомлений...")
        self.log_signal.emit(f"📊 Выбрано: {len(selected_students)} уч.")
        self.log_signal.emit(f"📁 Класс: {self.current_class['class_name']}")
        self.log_signal.emit(f"📅 Период: {period_name}")
        notifications_folder = os.path.join(
            self.output_edit.text(), "Уведомления"
        )
        os.makedirs(notifications_folder, exist_ok=True)
        created_count = 0
        no_issues_count = 0
        for i, student_name in enumerate(selected_students):
            self.log_signal.emit(
                f"\n👤 {i + 1}/{len(selected_students)}: {student_name}"
            )
            data = self.analyze_student(
                self.current_class['path'], student_name
            )
            has_data = (
                (self.include_absences.isChecked() and data['absences']) or
                (self.include_az.isChecked() and data['az']) or
                (self.include_npa.isChecked() and data['npa']) or
                (self.include_marks.isChecked() and data['marks']) or
                (self.include_accumulation.isChecked() and data['accumulation'])
            )
            if not has_data:
                self.log_signal.emit("  ⚠️ Нет замечаний")
                empty_filepath = self.create_empty_notification(
                    student_name,
                    self.current_class['class_name'],
                    notifications_folder,
                )
                if empty_filepath:
                    no_issues_count += 1
                    self.log_signal.emit("  📄 Создан пустой файл")
                QApplication.processEvents()
                continue
            filepath = self.create_notification_docx(
                student_name,
                self.current_class['class_name'],
                data,
                notifications_folder,
            )
            if filepath:
                created_count += 1
                self.log_signal.emit("  ✅ Уведомление создано")
            QApplication.processEvents()
        self.log_signal.emit("\n✅ Завершено!")
        self.log_signal.emit(f"   📊 С замечаниями: {created_count}")
        self.log_signal.emit(f"   📄 Без замечаний: {no_issues_count}")
        self.log_signal.emit(f"   👥 Всего: {len(selected_students)}")
        self.create_btn.setEnabled(True)