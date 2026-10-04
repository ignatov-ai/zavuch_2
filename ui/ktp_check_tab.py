# -*- coding: utf-8 -*-
"""
Вкладка проверки КТП внеурочной деятельности (ВД).
Единый стиль с индиго-акцентом.
"""
import os
import sys
from datetime import datetime
from PySide6.QtWidgets import *
from PySide6.QtCore import QThread, Signal, Qt
from PySide6.QtGui import QColor
from collector_ktp import KTPCollector


# ============================================================
# Потоки
# ============================================================
class KTPLoadThread(QThread):
    """Поток для загрузки списка групп."""
    finished = Signal(list)
    error = Signal(str)
    log_message = Signal(str)

    def __init__(self, auth, academic_year_id=14):
        super().__init__()
        self.auth = auth
        self.academic_year_id = academic_year_id

    def run(self):
        try:
            collector = KTPCollector(self.auth)
            collector.log_callback = lambda text: self.log_message.emit(text)
            groups = collector.get_all_groups()
            self.finished.emit(groups)
        except Exception as e:
            import traceback
            error_text = f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}"
            self.error.emit(error_text)
            self.finished.emit([])


class KTPCheckThread(QThread):
    """Поток для проверки КТП выбранных групп."""
    progress_update = Signal(int, int, str)
    row_checked = Signal(dict)
    finished = Signal(int)
    error = Signal(str)
    log_message = Signal(str)

    def __init__(self, auth, groups_to_check, academic_year_id=14):
        super().__init__()
        self.auth = auth
        self.groups_to_check = groups_to_check
        self.academic_year_id = academic_year_id
        self._is_running = True

    def stop(self):
        self._is_running = False

    def run(self):
        try:
            collector = KTPCollector(self.auth)
            collector.log_callback = lambda text: self.log_message.emit(text)
            total = len(self.groups_to_check)
            checked = 0
            for i, group in enumerate(self.groups_to_check):
                if not self._is_running:
                    self.log_message.emit("[KTP] ⏹️ Остановка проверки...")
                    break
                group_id = group.get('id')
                group_name = group.get('name', '')
                self.progress_update.emit(i + 1, total, group_name)
                try:
                    has_ktp, lessons_with_names, total_lessons = collector.has_ktp(
                        group_id, self.academic_year_id
                    )
                    group['has_ktp'] = has_ktp
                    group['lessons_with_names'] = lessons_with_names
                    group['total_lessons'] = total_lessons
                    group['is_checked'] = True
                    self.row_checked.emit(group.copy())
                    checked += 1
                except Exception as e:
                    self.log_message.emit(
                        f"[KTP] ❌ Ошибка для {group_name}: {e}"
                    )
                    group['is_checked'] = True
                    group['has_ktp'] = False
                    group['lessons_with_names'] = 0
                    group['total_lessons'] = 0
                    self.row_checked.emit(group.copy())
            self.finished.emit(checked)
        except Exception as e:
            import traceback
            error_text = f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}"
            self.error.emit(error_text)
            self.finished.emit(0)


# ============================================================
# Единый стиль (индиго — внеурочная деятельность)
# ============================================================
TAB_STYLE = """
    QWidget {
        font-size: 10pt;
        color: #1e293b;
    }
    QGroupBox {
        font-size: 11pt;
        font-weight: 700;
        color: #4f46e5;
        border: 1px solid #a5b4fc;
        border-radius: 10px;
        margin-top: 12px;
        padding: 18px 14px 14px 14px;
        background: #ffffff;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 12px;
        padding: 0 6px;
    }
    QLineEdit {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QLineEdit:focus { border: 1px solid #6366f1; }
    QLineEdit:disabled {
        background: #f1f5f9;
        color: #94a3b8;
    }
    QPushButton {
        min-height: 30px;
        padding: 4px 12px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #f8fafc;
        color: #1e293b;
        font-weight: 500;
    }
    QPushButton:hover { background: #e2e8f0; }
    QPushButton:pressed { background: #cbd5e1; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
    QComboBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QSpinBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QTableWidget {
        border: 1px solid #e2e8f0;
        border-radius: 6px;
        gridline-color: #e2e8f0;
        alternate-background-color: #eef2ff;
    }
    QTableWidget::item { padding: 3px 6px; }
    QHeaderView::section {
        background: #f1f5f9;
        color: #1e293b;
        padding: 5px;
        border: none;
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
        font-weight: 600;
    }
    QProgressBar {
        min-height: 22px;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        background: #f8fafc;
        text-align: center;
        color: #1e293b;
        font-weight: 600;
    }
    QProgressBar::chunk {
        background-color: #4f46e5;
        border-radius: 5px;
    }
    QLabel { color: #334155; }
"""

PRIMARY_BTN = """
    QPushButton {
        background-color: #4f46e5;
        color: #ffffff;
        font-weight: 700;
        border: none;
        border-radius: 7px;
    }
    QPushButton:hover { background-color: #4338ca; }
    QPushButton:pressed { background-color: #3730a3; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

SUCCESS_BTN = """
    QPushButton {
        background-color: #059669;
        color: #ffffff;
        font-weight: 600;
        border: none;
        border-radius: 6px;
    }
    QPushButton:hover { background-color: #047857; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

WARNING_BTN = """
    QPushButton {
        background-color: #ea580c;
        color: #ffffff;
        font-weight: 600;
        border: none;
        border-radius: 6px;
    }
    QPushButton:hover { background-color: #c2410c; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

DANGER_BTN = """
    QPushButton {
        background-color: #fff1f2;
        color: #be123c;
        border: 1px solid #fecdd3;
        font-weight: 600;
    }
    QPushButton:hover { background-color: #ffe4e6; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
"""


class KTPCheckTab(QWidget):
    """Вкладка проверки КТП внеурочной деятельности (ВД)."""

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.auth = None
        self.academic_year_id = 14
        self.all_groups = []
        self.filtered_groups = []
        self.initUI()

    def on_auth_updated(self, auth):
        self.auth = auth
        if auth:
            self.load_btn.setEnabled(True)
        else:
            self.load_btn.setEnabled(False)

    def on_academic_year_updated(self, aid):
        self.academic_year_id = aid
        if hasattr(self, 'year_spin'):
            self.year_spin.setValue(aid)

    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(16, 14, 16, 14)

        # === Заголовок ===
        title = QLabel("🔍  Проверка КТП (ВД)")
        title.setStyleSheet(
            "font-size: 16pt; font-weight: 700; color: #0f172a; padding: 2px 0 6px 0;"
        )
        main_layout.addWidget(title)

        subtitle = QLabel(
            "Проверка наличия КТП для групп внеурочной деятельности."
        )
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet("color: #64748b; margin-bottom: 4px;")
        main_layout.addWidget(subtitle)

        # === ПАНЕЛЬ УПРАВЛЕНИЯ ===
        control_group = QGroupBox("⚙️  Управление")
        control_layout = QHBoxLayout(control_group)
        control_layout.setSpacing(10)

        control_layout.addWidget(QLabel("Учебный год:"))
        self.year_spin = QSpinBox()
        self.year_spin.setRange(1, 99)
        self.year_spin.setValue(self.academic_year_id)
        self.year_spin.setMaximumWidth(80)
        self.year_spin.setToolTip("13 = 2025-2026, 14 = 2026-2027")
        control_layout.addWidget(self.year_spin)

        control_layout.addSpacing(10)

        self.load_btn = QPushButton("📋  1. Загрузить группы")
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(36)
        self.load_btn.setStyleSheet("""
            QPushButton {
                background: #eef2ff; color: #4f46e5;
                border: 1px solid #c7d2fe; font-weight: 600;
            }
            QPushButton:hover { background: #e0e7ff; }
        """)
        self.load_btn.clicked.connect(self.load_groups)
        control_layout.addWidget(self.load_btn)

        self.check_btn = QPushButton("🔍  2. Проверить выбранные")
        self.check_btn.setEnabled(False)
        self.check_btn.setMinimumHeight(36)
        self.check_btn.setStyleSheet(PRIMARY_BTN)
        self.check_btn.clicked.connect(self.check_ktp_selected)
        control_layout.addWidget(self.check_btn)

        self.check_unchecked_btn = QPushButton("🔍  Проверить не проверенные")
        self.check_unchecked_btn.setEnabled(False)
        self.check_unchecked_btn.setMinimumHeight(36)
        self.check_unchecked_btn.setStyleSheet(WARNING_BTN)
        self.check_unchecked_btn.clicked.connect(self.check_ktp_unchecked)
        control_layout.addWidget(self.check_unchecked_btn)

        self.check_all_btn = QPushButton("🔍  Проверить все")
        self.check_all_btn.setEnabled(False)
        self.check_all_btn.setMinimumHeight(36)
        self.check_all_btn.setStyleSheet("""
            QPushButton {
                background-color: #0369a1;
                color: #ffffff;
                font-weight: 600;
                border: none;
                border-radius: 6px;
            }
            QPushButton:hover { background-color: #075985; }
            QPushButton:disabled {
                background-color: #cbd5e1;
                color: #64748b;
            }
        """)
        self.check_all_btn.clicked.connect(self.check_ktp_all)
        control_layout.addWidget(self.check_all_btn)

        self.stop_btn = QPushButton("⏹️  Остановить")
        self.stop_btn.setEnabled(False)
        self.stop_btn.setMaximumWidth(130)
        self.stop_btn.setStyleSheet(DANGER_BTN)
        self.stop_btn.clicked.connect(self.stop_check)
        control_layout.addWidget(self.stop_btn)

        control_layout.addStretch()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("""
            color: #4338ca;
            background: #eef2ff;
            border-radius: 6px;
            padding: 6px 10px;
            font-weight: 700;
            font-size: 10pt;
        """)
        control_layout.addWidget(self.stats_label)

        main_layout.addWidget(control_group)

        # === ПРОГРЕСС-БАР ===
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        main_layout.addWidget(self.progress_bar)

        # === ФИЛЬТРЫ ===
        filter_group = QGroupBox("🔎  Фильтры")
        filter_group.setStyleSheet("""
            QGroupBox {
                font-size: 11pt; font-weight: 700;
                color: #0369a1; border: 1px solid #7dd3fc;
                border-radius: 10px; margin-top: 12px;
                padding: 18px 14px 14px 14px; background: #ffffff;
            }
            QGroupBox::title {
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
            }
        """)
        filter_layout = QHBoxLayout(filter_group)
        filter_layout.setSpacing(10)
        filter_layout.setContentsMargins(10, 6, 10, 10)

        filter_layout.addWidget(QLabel("Параллель:"))
        self.level_filter = QComboBox()
        self.level_filter.addItem("Все параллели", None)
        for level in range(1, 12):
            self.level_filter.addItem(f"{level} класс", level)
        self.level_filter.setMinimumWidth(130)
        self.level_filter.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.level_filter)

        filter_layout.addWidget(QLabel("Поиск:"))
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Название группы или активности...")
        self.search_edit.setMinimumWidth(220)
        self.search_edit.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.search_edit)

        filter_layout.addWidget(QLabel("Статус КТП:"))
        self.status_filter = QComboBox()
        self.status_filter.addItem("Все", "all")
        self.status_filter.addItem("✅ Есть КТП", "has_ktp")
        self.status_filter.addItem("❌ Нет КТП", "no_ktp")
        self.status_filter.addItem("❓ Не проверено", "not_checked")
        self.status_filter.setMinimumWidth(150)
        self.status_filter.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.status_filter)

        self.clear_filters_btn = QPushButton("Очистить")
        self.clear_filters_btn.setStyleSheet(DANGER_BTN)
        self.clear_filters_btn.clicked.connect(self.clear_filters)
        filter_layout.addWidget(self.clear_filters_btn)
        filter_layout.addStretch()

        main_layout.addWidget(filter_group)

        # === ПАНЕЛЬ ВЫБОРА ===
        select_group = QGroupBox("✓  Выбор групп")
        select_group.setStyleSheet("""
            QGroupBox {
                font-size: 11pt; font-weight: 700;
                color: #475569; border: 1px solid #cbd5e1;
                border-radius: 10px; margin-top: 12px;
                padding: 18px 14px 14px 14px; background: #f8fafc;
            }
            QGroupBox::title {
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
            }
        """)
        select_layout = QHBoxLayout(select_group)
        select_layout.setSpacing(8)

        self.select_all_btn = QPushButton("✓  Выбрать всё")
        self.select_all_btn.setEnabled(False)
        self.select_all_btn.clicked.connect(self.select_all)
        select_layout.addWidget(self.select_all_btn)

        self.deselect_all_btn = QPushButton("✗  Снять всё")
        self.deselect_all_btn.setEnabled(False)
        self.deselect_all_btn.clicked.connect(self.deselect_all)
        select_layout.addWidget(self.deselect_all_btn)

        select_layout.addSpacing(10)

        self.select_visible_btn = QPushButton("✓  Выбрать видимые")
        self.select_visible_btn.setEnabled(False)
        self.select_visible_btn.clicked.connect(self.select_visible)
        select_layout.addWidget(self.select_visible_btn)

        self.select_unchecked_btn = QPushButton("✓  Выбрать не проверенные")
        self.select_unchecked_btn.setEnabled(False)
        self.select_unchecked_btn.clicked.connect(self.select_unchecked)
        select_layout.addWidget(self.select_unchecked_btn)

        select_layout.addStretch()

        self.selected_count_label = QLabel("Выбрано: 0")
        self.selected_count_label.setStyleSheet("""
            color: #4f46e5;
            background: #eef2ff;
            border: 1px solid #c7d2fe;
            border-radius: 6px;
            padding: 5px 10px;
            font-weight: 700;
        """)
        select_layout.addWidget(self.selected_count_label)

        main_layout.addWidget(select_group)

        # === ТАБЛИЦА ===
        self.groups_table = QTableWidget()
        self.groups_table.setAlternatingRowColors(True)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.groups_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.groups_table.horizontalHeader().setStretchLastSection(True)
        self.groups_table.setColumnCount(8)
        self.groups_table.setHorizontalHeaderLabels([
            '✓', 'Группа', 'Активность',
            'Параллель', 'Класс(ы)',
            'Учеников', 'КТП', 'Статус'
        ])
        self.groups_table.verticalHeader().setDefaultSectionSize(70)
        self.groups_table.doubleClicked.connect(self.on_row_double_clicked)
        main_layout.addWidget(self.groups_table, 1)

        # === ЭКСПОРТ ===
        export_layout = QHBoxLayout()
        export_layout.addStretch()
        self.export_btn = QPushButton("📊  Экспорт в Excel")
        self.export_btn.setEnabled(False)
        self.export_btn.setStyleSheet(SUCCESS_BTN)
        self.export_btn.clicked.connect(self.export_to_excel)
        export_layout.addWidget(self.export_btn)
        main_layout.addLayout(export_layout)

        self.setStyleSheet(TAB_STYLE)
        self.set_controls_enabled(False)

    def set_controls_enabled(self, enabled):
        self.level_filter.setEnabled(enabled)
        self.search_edit.setEnabled(enabled)
        self.status_filter.setEnabled(enabled)
        self.clear_filters_btn.setEnabled(enabled)
        self.select_all_btn.setEnabled(enabled)
        self.deselect_all_btn.setEnabled(enabled)
        self.select_visible_btn.setEnabled(enabled)
        self.select_unchecked_btn.setEnabled(enabled)
        self.export_btn.setEnabled(enabled and bool(self.all_groups))

    # ==================== ЭТАП 1: ЗАГРУЗКА ====================
    def load_groups(self):
        if not self.auth:
            QMessageBox.warning(
                self, "Ошибка",
                "Сначала авторизуйтесь на вкладке 'Настройки'"
            )
            return
        self.load_btn.setEnabled(False)
        self.check_btn.setEnabled(False)
        self.check_all_btn.setEnabled(False)
        self.check_unchecked_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFormat("Загрузка списка групп...")
        self.groups_table.setRowCount(0)
        self.all_groups = []
        academic_year_id = self.year_spin.value()
        self.academic_year_id = academic_year_id
        self.load_thread = KTPLoadThread(self.auth, academic_year_id)
        self.load_thread.finished.connect(self.on_load_finished)
        self.load_thread.error.connect(self.on_load_error)
        self.load_thread.log_message.connect(self.on_log_message)
        self.load_thread.start()

    def on_load_finished(self, groups):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)
        if not groups:
            QMessageBox.warning(
                self, "Внимание", "Не удалось получить список групп."
            )
            return
        for group in groups:
            group['is_checked'] = False
            group['has_ktp'] = None
            group['lessons_with_names'] = 0
            group['total_lessons'] = 0
            group['is_selected'] = False
        self.all_groups = groups
        self.apply_filter()
        self.update_stats()
        self.set_controls_enabled(True)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)
        QMessageBox.information(
            self, "Готово",
            f"✅ Загружено групп: {len(groups)}\n\n"
            f"Выберите группы и нажмите 'Проверить выбранные'"
        )

    def on_load_error(self, error):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)
        QMessageBox.critical(self, "Ошибка загрузки", f"Ошибка:\n{error}")

    # ==================== ЭТАП 2: ПРОВЕРКА ====================
    def check_ktp_selected(self):
        selected = [g for g in self.all_groups if g.get('is_selected')]
        if not selected:
            QMessageBox.warning(self, "Ошибка", "Не выбрано ни одной группы.")
            return
        self._start_check(selected, "выбранных")

    def check_ktp_unchecked(self):
        unchecked = [g for g in self.all_groups if not g.get('is_checked')]
        if not unchecked:
            QMessageBox.information(
                self, "Внимание", "Все группы уже проверены!"
            )
            return
        self._start_check(unchecked, "не проверенных")

    def check_ktp_all(self):
        if not self.all_groups:
            QMessageBox.warning(
                self, "Ошибка", "Сначала загрузите список групп"
            )
            return
        self._start_check(self.all_groups, "всех")

    def _start_check(self, groups, label):
        reply = QMessageBox.question(
            self, "Подтверждение",
            f"Проверить КТП для {len(groups)} {label} групп?\n\n"
            f"Это может занять несколько минут.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        self.check_btn.setEnabled(False)
        self.check_all_btn.setEnabled(False)
        self.check_unchecked_btn.setEnabled(False)
        self.load_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        academic_year_id = self.year_spin.value()
        self.check_thread = KTPCheckThread(
            self.auth, groups, academic_year_id
        )
        self.check_thread.progress_update.connect(self.on_check_progress)
        self.check_thread.row_checked.connect(self.on_row_checked)
        self.check_thread.finished.connect(self.on_check_finished)
        self.check_thread.error.connect(self.on_check_error)
        self.check_thread.log_message.connect(self.on_log_message)
        self.check_thread.start()

    def stop_check(self):
        if hasattr(self, 'check_thread') and self.check_thread:
            self.check_thread.stop()

    def on_check_progress(self, current, total, name):
        if total > 0:
            percent = int(current / total * 100)
            self.progress_bar.setValue(percent)
            self.progress_bar.setFormat(
                f"Проверка {current}/{total}: {name[:40]} ({percent}%)"
            )
            QApplication.processEvents()

    def on_row_checked(self, group):
        for g in self.all_groups:
            if g['id'] == group['id']:
                g.update({
                    'has_ktp': group['has_ktp'],
                    'lessons_with_names': group['lessons_with_names'],
                    'total_lessons': group['total_lessons'],
                    'is_checked': True,
                })
                break
        self._update_row_in_table(group)
        self.update_stats()

    def _update_row_in_table(self, group):
        for i, g in enumerate(self.filtered_groups):
            if g['id'] == group['id']:
                g.update({
                    'has_ktp': group['has_ktp'],
                    'lessons_with_names': group['lessons_with_names'],
                    'total_lessons': group['total_lessons'],
                    'is_checked': True,
                })
                green_bg = QColor(209, 250, 229)
                red_bg = QColor(254, 202, 202)
                ktp_item = self.groups_table.item(i, 6)
                if ktp_item:
                    ktp_item.setText(
                        f"{group['lessons_with_names']} / {group['total_lessons']}"
                    )
                status_item = self.groups_table.item(i, 7)
                if status_item:
                    if group['has_ktp']:
                        status_item.setText("✅ Есть КТП")
                        status_item.setBackground(green_bg)
                        status_item.setForeground(QColor(6, 95, 70))
                        row_color = green_bg
                    else:
                        status_item.setText("❌ Нет КТП")
                        status_item.setBackground(red_bg)
                        status_item.setForeground(QColor(153, 27, 27))
                        row_color = red_bg
                    for col in [1, 2, 3, 4, 5, 6]:
                        item = self.groups_table.item(i, col)
                        if item:
                            item.setBackground(row_color)
                break

    def on_check_finished(self, checked_count):
        self.progress_bar.setVisible(False)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        self.apply_filter()
        self.update_stats()
        total_unchecked = sum(
            1 for g in self.all_groups if not g.get('is_checked')
        )
        if total_unchecked > 0:
            QMessageBox.information(
                self, "Проверка остановлена",
                f"✅ Проверено групп: {checked_count}\n"
                f"❓ Осталось не проверено: {total_unchecked}"
            )
        else:
            QMessageBox.information(
                self, "Проверка завершена",
                f"✅ Проверено групп: {checked_count}"
            )

    def on_check_error(self, error):
        self.progress_bar.setVisible(False)
        self.check_btn.setEnabled(True)
        self.check_all_btn.setEnabled(True)
        self.check_unchecked_btn.setEnabled(True)
        self.load_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        QMessageBox.critical(self, "Ошибка проверки", f"Ошибка:\n{error}")

    def on_log_message(self, text):
        try:
            sys.__stdout__.write(text + '\n')
            sys.__stdout__.flush()
        except Exception:
            pass

    # ==================== СТАТИСТИКА И ФИЛЬТРЫ ====================
    def update_stats(self):
        total = len(self.all_groups)
        checked = sum(1 for g in self.all_groups if g.get('is_checked'))
        with_ktp = sum(1 for g in self.all_groups if g.get('has_ktp') is True)
        without_ktp = sum(1 for g in self.all_groups if g.get('has_ktp') is False)
        self.stats_label.setText(
            f"Всего: {total}  |  Проверено: {checked}  |  "
            f"✅ {with_ktp}  |  ❌ {without_ktp}"
        )
        self.update_selected_count()

    def update_selected_count(self):
        selected = sum(1 for g in self.all_groups if g.get('is_selected'))
        self.selected_count_label.setText(f"Выбрано: {selected}")

    def apply_filter(self):
        level_filter = self.level_filter.currentData()
        search_text = self.search_edit.text().lower().strip()
        status_filter = self.status_filter.currentData()
        filtered = []
        for group in self.all_groups:
            if level_filter is not None:
                if level_filter not in group.get('class_levels', []):
                    continue
            if search_text:
                searchable = (
                    group.get('name', '').lower() + ' '
                    + group.get('activity_name', '').lower() + ' '
                    + group.get('activity_short_name', '').lower()
                )
                if search_text not in searchable:
                    continue
            if status_filter == 'has_ktp':
                if group.get('has_ktp') is not True:
                    continue
            elif status_filter == 'no_ktp':
                if group.get('has_ktp') is not False:
                    continue
            elif status_filter == 'not_checked':
                if group.get('is_checked'):
                    continue
            filtered.append(group)
        self.filtered_groups = filtered
        self.populate_table(filtered)

    def clear_filters(self):
        self.level_filter.setCurrentIndex(0)
        self.search_edit.clear()
        self.status_filter.setCurrentIndex(0)
        self.apply_filter()

    def populate_table(self, groups):
        self.groups_table.setSortingEnabled(False)
        self.groups_table.setUpdatesEnabled(False)
        self.groups_table.setRowCount(len(groups))

        green_bg = QColor(209, 250, 229)
        red_bg = QColor(254, 202, 202)
        gray_bg = QColor(241, 245, 249)

        for i, group in enumerate(groups):
            # Колонка 0: чекбокс
            cb_widget = QWidget()
            cb_layout = QHBoxLayout(cb_widget)
            cb_layout.setContentsMargins(0, 0, 0, 0)
            cb_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
            checkbox = QCheckBox()
            checkbox.setChecked(group.get('is_selected', False))
            checkbox.stateChanged.connect(
                lambda state, g=group: self.on_checkbox_changed(state, g)
            )
            cb_layout.addWidget(checkbox)
            self.groups_table.setCellWidget(i, 0, cb_widget)

            # Колонка 1: Название группы
            name_item = QTableWidgetItem(group.get('name', ''))
            name_item.setToolTip(f"ID группы: {group['id']}")
            self.groups_table.setItem(i, 1, name_item)

            # Колонка 2: Активность
            self.groups_table.setItem(
                i, 2, QTableWidgetItem(group.get('activity_name', ''))
            )

            # Колонка 3: Параллель
            class_levels = group.get('class_levels', [])
            parallel_text = (
                '\n'.join(str(l) for l in sorted(class_levels))
                if class_levels else ''
            )
            parallel_item = QTableWidgetItem(parallel_text)
            parallel_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 3, parallel_item)

            # Колонка 4: Класс(ы)
            class_names = group.get('class_unit_names', [])
            if not class_names:
                class_names = [
                    f"ID: {uid}" for uid in group.get('class_unit_ids', [])
                ]
            class_text = '\n'.join(sorted(class_names)) if class_names else ''
            class_item = QTableWidgetItem(class_text)
            class_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 4, class_item)

            # Колонка 5: Учеников
            students_item = QTableWidgetItem(
                str(group.get('student_count', 0))
            )
            students_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 5, students_item)

            # Колонка 6: КТП
            if group.get('is_checked'):
                ktp_text = (
                    f"{group['lessons_with_names']} / {group['total_lessons']}"
                )
            else:
                ktp_text = "—"
            ktp_item = QTableWidgetItem(ktp_text)
            ktp_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.groups_table.setItem(i, 6, ktp_item)

            # Колонка 7: Статус
            has_ktp = group.get('has_ktp')
            is_checked = group.get('is_checked', False)
            if not is_checked:
                status_item = QTableWidgetItem("❓ Не проверено")
                status_item.setBackground(gray_bg)
                status_item.setForeground(QColor(71, 85, 105))
                row_color = gray_bg
            elif has_ktp:
                status_item = QTableWidgetItem("✅ Есть КТП")
                status_item.setBackground(green_bg)
                status_item.setForeground(QColor(6, 95, 70))
                row_color = green_bg
            else:
                status_item = QTableWidgetItem("❌ Нет КТП")
                status_item.setBackground(red_bg)
                status_item.setForeground(QColor(153, 27, 27))
                row_color = red_bg
            self.groups_table.setItem(i, 7, status_item)

            for col in [1, 2, 3, 4, 5, 6]:
                item = self.groups_table.item(i, col)
                if item:
                    item.setBackground(row_color)

        self.groups_table.setUpdatesEnabled(True)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setColumnWidth(0, 40)
        self.groups_table.setColumnWidth(1, 200)
        self.groups_table.setColumnWidth(2, 250)
        self.groups_table.setColumnWidth(3, 90)
        self.groups_table.setColumnWidth(4, 130)
        self.groups_table.setColumnWidth(5, 80)
        self.groups_table.setColumnWidth(6, 100)
        self.groups_table.setColumnWidth(7, 140)

    def on_checkbox_changed(self, state, group):
        group['is_selected'] = (state == Qt.CheckState.Checked.value)
        self.update_selected_count()

    # ==================== ВЫБОР ====================
    def select_all(self):
        for group in self.all_groups:
            group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def deselect_all(self):
        for group in self.all_groups:
            group['is_selected'] = False
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def select_visible(self):
        for group in self.filtered_groups:
            group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def select_unchecked(self):
        for group in self.all_groups:
            if not group.get('is_checked'):
                group['is_selected'] = True
        self.populate_table(self.filtered_groups)
        self.update_selected_count()

    def on_row_double_clicked(self, index):
        row = index.row()
        if row < 0 or row >= len(self.filtered_groups):
            return
        group = self.filtered_groups[row]
        url = group.get('journal_url', '')
        if url:
            try:
                import webbrowser
                webbrowser.open(url)
            except Exception as e:
                QMessageBox.warning(
                    self, "Ошибка", f"Не удалось открыть браузер:\n{e}"
                )

    # ==================== ЭКСПОРТ ====================
    def export_to_excel(self):
        if not self.filtered_groups:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка", "Модуль openpyxl не установлен"
            )
            return

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"Проверка_КТП_ВД_{timestamp}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт", filename, "Excel files (*.xlsx)"
        )
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Проверка КТП (ВД)"
            headers = [
                'ID группы', 'Группа', 'Активность', 'Параллель', 'Класс(ы)',
                'Учеников', 'Уроков с названиями', 'Всего уроков',
                'Статус КТП', 'URL журнала'
            ]
            ws.append(headers)

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(
                start_color="4F46E5", end_color="4F46E5", fill_type="solid"
            )
            header_alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True
            )
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            green_fill = PatternFill(
                start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"
            )
            red_fill = PatternFill(
                start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"
            )
            gray_fill = PatternFill(
                start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"
            )

            for group in self.filtered_groups:
                if not group.get('is_checked'):
                    status = '❓ Не проверено'
                    fill = gray_fill
                elif group.get('has_ktp'):
                    status = '✅ Есть КТП'
                    fill = green_fill
                else:
                    status = '❌ Нет КТП'
                    fill = red_fill

                class_levels = group.get('class_levels', [])
                parallel_text = ', '.join(str(l) for l in sorted(class_levels))
                class_names = group.get('class_unit_names', [])
                class_text = ', '.join(sorted(class_names))

                row_data = [
                    group['id'],
                    group.get('name', ''),
                    group.get('activity_name', ''),
                    parallel_text,
                    class_text,
                    group.get('student_count', 0),
                    group.get('lessons_with_names', 0),
                    group.get('total_lessons', 0),
                    status,
                    group.get('journal_url', ''),
                ]
                ws.append(row_data)
                row_idx = ws.max_row
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col)
                    cell.fill = fill

            col_widths = {
                'A': 12, 'B': 25, 'C': 40, 'D': 12, 'E': 20,
                'F': 10, 'G': 18, 'H': 12, 'I': 15, 'J': 60,
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = 'A2'

            wb.save(file_path)
            reply = QMessageBox.question(
                self, "Готово",
                f"Отчёт сохранён:\n{file_path}\n\nОткрыть файл?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)
        except Exception as e:
            QMessageBox.critical(
                self, "Ошибка", f"Не удалось сохранить отчёт:\n{e}"
            )