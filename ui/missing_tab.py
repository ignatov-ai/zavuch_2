# -*- coding: utf-8 -*-
"""
Вкладка «Пропуски занятий».
Авторизация приходит из MainWindow через сигнал auth_updated.
Период задаётся двумя полями QDateEdit прямо в панели.
"""
import os
import csv
from collections import defaultdict
from datetime import datetime, timedelta
from PySide6.QtWidgets import *
from PySide6.QtCore import *
from PySide6.QtGui import *

from ui.console import ConsoleWidget


class MissingTab(QWidget):
    """Вкладка для просмотра и выгрузки пропусков занятий"""

    log_signal = Signal(str)

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.auth = None
        self.academic_year_id = 14
        self.classes_list = []
        self.selected_student = None
        self.selected_student_name = ""
        self.selected_class = None
        self.missing_data = []
        self.initUI()

        self.log_signal.connect(self.append_to_console)

    # ------------------------------------------------------------------
    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(10)
        main_layout.setContentsMargins(15, 15, 15, 15)

        # === ВЕРХНЯЯ ПАНЕЛЬ: ВЫБОР КЛАССА И УЧЕНИКА ===
        selector_group = QGroupBox("Выбор ученика и периода")
        selector_layout = QGridLayout(selector_group)
        selector_layout.setVerticalSpacing(10)
        selector_layout.setHorizontalSpacing(15)

        # --- Строка 1: Параллель / Класс / Ученик ---
        selector_layout.addWidget(QLabel("Параллель:"), 0, 0)
        self.level_combo = QComboBox()
        self.level_combo.setMinimumWidth(120)
        self.level_combo.currentIndexChanged.connect(self.on_level_changed)
        selector_layout.addWidget(self.level_combo, 0, 1)

        selector_layout.addWidget(QLabel("Класс:"), 0, 2)
        self.class_combo = QComboBox()
        self.class_combo.setMinimumWidth(150)
        self.class_combo.currentIndexChanged.connect(self.on_class_changed)
        selector_layout.addWidget(self.class_combo, 0, 3)

        selector_layout.addWidget(QLabel("Ученик:"), 0, 4)
        self.student_combo = QComboBox()
        self.student_combo.setMinimumWidth(200)
        self.student_combo.setEditable(True)
        self.student_combo.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self.student_combo.currentTextChanged.connect(self.on_student_changed)
        selector_layout.addWidget(self.student_combo, 0, 5)

        selector_layout.setColumnStretch(5, 1)

        # --- Строка 2: Период (с / по) ---
        selector_layout.addWidget(QLabel("Период с:"), 1, 0)

        self.start_date_edit = QDateEdit()
        self.start_date_edit.setCalendarPopup(True)
        self.start_date_edit.setDisplayFormat("dd.MM.yyyy")
        self.start_date_edit.setMinimumHeight(32)
        self.start_date_edit.setMinimumWidth(130)
        # По умолчанию — 1 сентября текущего учебного года
        today = datetime.now()
        if today.month >= 9:
            default_start = QDate(today.year, 9, 1)
        else:
            default_start = QDate(today.year - 1, 9, 1)
        self.start_date_edit.setDate(default_start)
        selector_layout.addWidget(self.start_date_edit, 1, 1)

        selector_layout.addWidget(QLabel("по:"), 1, 2)

        self.end_date_edit = QDateEdit()
        self.end_date_edit.setCalendarPopup(True)
        self.end_date_edit.setDisplayFormat("dd.MM.yyyy")
        self.end_date_edit.setMinimumHeight(32)
        self.end_date_edit.setMinimumWidth(130)
        # По умолчанию — сегодня
        self.end_date_edit.setDate(QDate.currentDate())
        selector_layout.addWidget(self.end_date_edit, 1, 3)

        # --- Кнопки периода ---
        period_btn_layout = QHBoxLayout()
        period_btn_layout.setSpacing(6)

        self.period_t1_btn = QPushButton("Т1")
        self.period_t1_btn.setToolTip("Сентябрь – Ноябрь")
        self.period_t1_btn.setMaximumWidth(50)
        self.period_t1_btn.clicked.connect(lambda: self._set_period("Т1"))
        period_btn_layout.addWidget(self.period_t1_btn)

        self.period_t2_btn = QPushButton("Т2")
        self.period_t2_btn.setToolTip("Декабрь – Февраль")
        self.period_t2_btn.setMaximumWidth(50)
        self.period_t2_btn.clicked.connect(lambda: self._set_period("Т2"))
        period_btn_layout.addWidget(self.period_t2_btn)

        self.period_t3_btn = QPushButton("Т3")
        self.period_t3_btn.setToolTip("Март – Май")
        self.period_t3_btn.setMaximumWidth(50)
        self.period_t3_btn.clicked.connect(lambda: self._set_period("Т3"))
        period_btn_layout.addWidget(self.period_t3_btn)

        self.period_p1_btn = QPushButton("П1")
        self.period_p1_btn.setToolTip("Сентябрь – Декабрь")
        self.period_p1_btn.setMaximumWidth(50)
        self.period_p1_btn.clicked.connect(lambda: self._set_period("П1"))
        period_btn_layout.addWidget(self.period_p1_btn)

        self.period_p2_btn = QPushButton("П2")
        self.period_p2_btn.setToolTip("Январь – Май")
        self.period_p2_btn.setMaximumWidth(50)
        self.period_p2_btn.clicked.connect(lambda: self._set_period("П2"))
        period_btn_layout.addWidget(self.period_p2_btn)

        self.period_year_btn = QPushButton("Год")
        self.period_year_btn.setToolTip("Весь учебный год")
        self.period_year_btn.setMaximumWidth(60)
        self.period_year_btn.clicked.connect(lambda: self._set_period("Год"))
        period_btn_layout.addWidget(self.period_year_btn)

        period_btn_widget = QWidget()
        period_btn_widget.setLayout(period_btn_layout)
        selector_layout.addWidget(period_btn_widget, 1, 4, 1, 2)

        # --- Строка 3: Кнопка загрузки ---
        self.load_btn = QPushButton("📥 Загрузить пропуски")
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(34)
        self.load_btn.setStyleSheet("""
            QPushButton { background-color: #2196F3; color: white;
                font-weight: bold; border-radius: 5px; }
            QPushButton:hover { background-color: #1976D2; }
            QPushButton:disabled { background-color: #cccccc; color: #666666; }
        """)
        self.load_btn.clicked.connect(self.load_missing_data)
        selector_layout.addWidget(self.load_btn, 2, 0, 1, 2)

        period_hint = QLabel(
            "ℹ️ Выберите даты вручную или нажмите Т1/Т2/Т3/П1/П2/Год "
            "для быстрой установки периода."
        )
        period_hint.setStyleSheet("color: #666; font-size: 9pt;")
        period_hint.setWordWrap(True)
        selector_layout.addWidget(period_hint, 2, 2, 1, 4)

        main_layout.addWidget(selector_group)

        # === ОСНОВНАЯ ИНФОРМАЦИЯ ОБ УЧЕНИКЕ ===
        info_group = QGroupBox("Информация об ученике")
        info_layout = QHBoxLayout(info_group)

        self.student_info_label = QLabel("Ученик не выбран")
        self.student_info_label.setStyleSheet("font-size: 12pt; font-weight: bold; color: #2c3e50;")
        info_layout.addWidget(self.student_info_label)

        self.total_missing_label = QLabel("Всего пропусков: 0")
        self.total_missing_label.setStyleSheet("color: #e74c3c; font-weight: bold;")
        info_layout.addWidget(self.total_missing_label)

        info_layout.addStretch()
        main_layout.addWidget(info_group)

        # === ТАБЛИЦА ПРЕДМЕТОВ С КОЛИЧЕСТВОМ ПРОПУСКОВ ===
        self.subjects_table = QTableWidget()
        self.subjects_table.setAlternatingRowColors(True)
        self.subjects_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.subjects_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.subjects_table.horizontalHeader().setStretchLastSection(True)
        self.subjects_table.setMinimumHeight(200)

        subjects_layout = QVBoxLayout()
        subjects_layout.addWidget(QLabel("📚 Пропуски по предметам:"))
        subjects_layout.addWidget(self.subjects_table)
        main_layout.addLayout(subjects_layout)

        # === ТАБЛИЦА ПРОПУСКОВ С ДЕТАЛИЗАЦИЕЙ ===
        details_group = QGroupBox("📋 Детализация пропусков")
        details_layout = QVBoxLayout(details_group)

        button_bar = QHBoxLayout()
        self.filter_edit = QLineEdit()
        self.filter_edit.setPlaceholderText("🔍 Фильтр по предмету...")
        self.filter_edit.textChanged.connect(self.filter_details_table)
        button_bar.addWidget(self.filter_edit)

        button_bar.addStretch()

        self.export_csv_btn = QPushButton("📎 Экспорт в CSV")
        self.export_csv_btn.setEnabled(False)
        self.export_csv_btn.clicked.connect(self.export_to_csv)
        button_bar.addWidget(self.export_csv_btn)

        self.export_excel_btn = QPushButton("📊 Экспорт в Excel")
        self.export_excel_btn.setEnabled(False)
        self.export_excel_btn.clicked.connect(self.export_to_excel)
        button_bar.addWidget(self.export_excel_btn)

        details_layout.addLayout(button_bar)

        self.details_table = QTableWidget()
        self.details_table.setAlternatingRowColors(True)
        self.details_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.details_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.details_table.horizontalHeader().setStretchLastSection(True)
        details_layout.addWidget(self.details_table)

        main_layout.addWidget(details_group)

        # === КОНСОЛЬ ===
        self.console = ConsoleWidget()
        self.console.setMaximumHeight(120)
        main_layout.addWidget(self.console)

        self.set_controls_enabled(False)

    # ------------------------------------------------------------------
    def set_controls_enabled(self, enabled):
        self.level_combo.setEnabled(enabled)
        self.class_combo.setEnabled(enabled)
        self.student_combo.setEnabled(enabled)
        self.start_date_edit.setEnabled(enabled)
        self.end_date_edit.setEnabled(enabled)
        self.period_t1_btn.setEnabled(enabled)
        self.period_t2_btn.setEnabled(enabled)
        self.period_t3_btn.setEnabled(enabled)
        self.period_p1_btn.setEnabled(enabled)
        self.period_p2_btn.setEnabled(enabled)
        self.period_year_btn.setEnabled(enabled)
        self.load_btn.setEnabled(enabled and self.selected_student is not None)

    # ==================================================================
    #  АВТОРИЗАЦИЯ ИЗ MAINWINDOW
    # ==================================================================
    def on_auth_updated(self, auth):
        self.auth = auth
        if auth:
            self.log_signal.emit("✅ Авторизация получена. Загружаю классы...")
            self.load_classes()
        else:
            self.set_controls_enabled(False)

    def on_academic_year_updated(self, aid):
        self.academic_year_id = aid
        if self.auth:
            self.auth.aid = str(aid)
            self.auth.curr_aid = str(aid)
            self.load_classes()

    # ------------------------------------------------------------------
    def load_classes(self):
        self.log_signal.emit("📚 Загрузка списка классов...")

        from collector import MarksDataCollector
        collector = MarksDataCollector(self.auth)
        self.classes_list = collector.get_classes()

        if not self.classes_list:
            self.log_signal.emit("⚠️ Не удалось загрузить список классов")
            return

        levels = sorted(set(c["level"] for c in self.classes_list if 1 <= c["level"] <= 11))

        self.level_combo.clear()
        self.level_combo.addItem("Выберите параллель", None)
        for level in levels:
            self.level_combo.addItem(f"{level} класс", level)

        self.set_controls_enabled(True)
        self.log_signal.emit(f"✅ Загружено классов: {len(self.classes_list)}")

    # ------------------------------------------------------------------
    def on_level_changed(self, index):
        self.class_combo.clear()
        self.class_combo.addItem("Выберите класс", None)
        self.student_combo.clear()
        self.student_combo.addItem("Выберите ученика", None)
        self.selected_student = None
        self.selected_class = None
        self.load_btn.setEnabled(False)

        level_data = self.level_combo.currentData()
        if not level_data:
            return

        filtered_classes = [c for c in self.classes_list if c["level"] == level_data]
        for class_info in filtered_classes:
            self.class_combo.addItem(class_info["name"], class_info)

    def on_class_changed(self, index):
        self.student_combo.clear()
        self.student_combo.addItem("Загрузка учеников...", None)
        self.selected_student = None
        self.selected_class = None
        self.load_btn.setEnabled(False)

        class_info = self.class_combo.currentData()
        if not class_info:
            self.student_combo.clear()
            self.student_combo.addItem("Выберите ученика", None)
            return

        self.selected_class = class_info
        self.log_signal.emit(f"👥 Загрузка учеников класса {class_info['name']}...")

        from collector import MarksDataCollector
        collector = MarksDataCollector(self.auth)

        groups = collector.get_groups_for_class(class_info["id"])
        if groups:
            students = collector.get_students_for_group(groups[0]["id"], class_info["id"])
            if students:
                self.student_combo.clear()
                for student in students:
                    full_name = self._get_student_full_name(student)
                    self.student_combo.addItem(full_name, student)
                self.log_signal.emit(f"✅ Загружено учеников: {len(students)}")
                return

        self.student_combo.clear()
        self.student_combo.addItem("Не удалось загрузить учеников", None)
        self.log_signal.emit("⚠️ Не удалось загрузить список учеников")

    def _get_student_full_name(self, student):
        last_name = student.get("last_name", "")
        first_name = student.get("first_name", "")
        middle_name = student.get("middle_name", "")

        if not last_name and student.get("user_name"):
            parts = student.get("user_name", "").split()
            if len(parts) >= 1:
                last_name = parts[0]
            if len(parts) >= 2:
                first_name = parts[1]
            if len(parts) >= 3:
                middle_name = parts[2]

        name_parts = [p for p in [last_name, first_name, middle_name] if p]
        return " ".join(name_parts) if name_parts else student.get("short_name", "")

    def on_student_changed(self, text):
        index = self.student_combo.currentIndex()
        if index < 0:
            self.selected_student = None
            self.load_btn.setEnabled(False)
            return

        student_data = self.student_combo.currentData()
        if student_data:
            self.selected_student = student_data
            self.selected_student_name = self.student_combo.currentText()
            self.student_info_label.setText(f"👤 {self.selected_student_name}")
            self.load_btn.setEnabled(True)
        else:
            self.selected_student = None
            self.load_btn.setEnabled(False)

    # ==================================================================
    #  БЫСТРАЯ УСТАНОВКА ПЕРИОДА
    # ==================================================================
    def _set_period(self, period_type: str):
        """Устанавливает даты в полях в соответствии с выбранным периодом."""
        today = datetime.now()
        if today.month >= 9:
            academic_start = today.year
        else:
            academic_start = today.year - 1
        academic_end = academic_start + 1

        # Периоды в формате (день, месяц) — как в academic_calendar.py
        periods = {
            "Т1": ((1, 9, academic_start), (30, 11, academic_start)),
            "Т2": ((1, 12, academic_start), (28, 2, academic_end)),
            "Т3": ((1, 3, academic_end), (31, 5, academic_end)),
            "П1": ((1, 9, academic_start), (31, 12, academic_start)),
            "П2": ((1, 1, academic_end), (31, 5, academic_end)),
            "Год": ((1, 9, academic_start), (31, 5, academic_end)),
        }

        if period_type not in periods:
            return

        (sd, sm, sy), (ed, em, ey) = periods[period_type]

        try:
            start_qdate = QDate(sy, sm, sd)
            end_qdate = QDate(ey, em, ed)
            self.start_date_edit.setDate(start_qdate)
            self.end_date_edit.setDate(end_qdate)
            self.log_signal.emit(
                f"[i] Установлен период {period_type}: "
                f"{start_qdate.toString('dd.MM.yyyy')} — {end_qdate.toString('dd.MM.yyyy')}"
            )
        except Exception as e:
            self.log_signal.emit(f"⚠️ Не удалось установить период {period_type}: {e}")

    # ==================================================================
    #  ЗАГРУЗКА ПРОПУСКОВ
    # ==================================================================
    def load_missing_data(self):
        if not self.selected_student or not self.auth:
            QMessageBox.warning(self, "Ошибка", "Выберите ученика!")
            return

        start_date = self.start_date_edit.date().toString("dd.MM.yyyy")
        end_date = self.end_date_edit.date().toString("dd.MM.yyyy")

        # Проверка, что начало не позже конца
        try:
            d1 = datetime.strptime(start_date, "%d.%m.%Y")
            d2 = datetime.strptime(end_date, "%d.%m.%Y")
            if d1 > d2:
                QMessageBox.warning(
                    self, "Ошибка",
                    "Дата начала не может быть позже даты окончания."
                )
                return
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Неверный формат даты: {e}")
            return

        self.load_btn.setEnabled(False)
        self.export_csv_btn.setEnabled(False)
        self.export_excel_btn.setEnabled(False)
        self.log_signal.emit(f"📅 Загрузка пропусков за период: {start_date} - {end_date}")
        self.log_signal.emit(f"👤 Ученик: {self.selected_student_name}")

        self.missing_data = []
        current_date = datetime.strptime(start_date, "%d.%m.%Y")
        end_date_obj = datetime.strptime(end_date, "%d.%m.%Y")

        profile_id = self.auth.pid
        student_id = self.selected_student.get("id")

        total_days = (end_date_obj - current_date).days + 1
        processed = 0

        while current_date <= end_date_obj:
            date_str = current_date.strftime("%d.%m.%Y")
            processed += 1

            if processed % 10 == 0 or processed == total_days:
                self.log_signal.emit(f"⏳ Обработка: {processed}/{total_days} дней")

            data = self.fetch_missing_for_date(date_str, profile_id, student_id)
            if data:
                for item in data:
                    self.missing_data.append({
                        "date": date_str,
                        "subject": item.get("subject_name", ""),
                        "teacher": item.get("teacher_fio", ""),
                        "topic": item.get("lesson_topic", ""),
                        "reason": item.get("nonattendance_reason_name", "Не указано"),
                        "teacher_id": item.get("teacher_id")
                    })

            current_date += timedelta(days=1)
            QApplication.processEvents()

        self.log_signal.emit(f"✅ Загружено пропусков: {len(self.missing_data)}")
        self.display_results()
        self.load_btn.setEnabled(True)

    def fetch_missing_for_date(self, date_str, profile_id, student_id):
        if not self.auth:
            return []

        try:
            url = "https://dnevnik.mos.ru/reports/api/missing/by_student/json"
            params = {
                "date": date_str,
                "pid": profile_id,
                "student_id": student_id
            }

            response = self.auth.fetch(url, params)
            if response and isinstance(response, list):
                return response
            return []
        except Exception as e:
            self.log_signal.emit(f"⚠️ Ошибка при загрузке {date_str}: {str(e)}")
            return []

    # ------------------------------------------------------------------
    def display_results(self):
        if not self.missing_data:
            self.subjects_table.setRowCount(0)
            self.details_table.setRowCount(0)
            self.total_missing_label.setText("Всего пропусков: 0")
            QMessageBox.information(self, "Информация", "Пропусков за выбранный период не найдено")
            return

        total = len(self.missing_data)
        self.total_missing_label.setText(f"Всего пропусков: {total}")

        subjects_count = defaultdict(int)
        for item in self.missing_data:
            subjects_count[item["subject"]] += 1

        self.subjects_table.setColumnCount(2)
        self.subjects_table.setHorizontalHeaderLabels(["Предмет", "Количество пропусков"])
        self.subjects_table.setRowCount(len(subjects_count))

        for i, (subject, count) in enumerate(sorted(subjects_count.items(), key=lambda x: x[1], reverse=True)):
            self.subjects_table.setItem(i, 0, QTableWidgetItem(subject))
            self.subjects_table.setItem(i, 1, QTableWidgetItem(str(count)))

        self.subjects_table.resizeColumnsToContents()
        self.subjects_table.setColumnWidth(0, 250)

        self.display_details_table(self.missing_data)
        self.export_csv_btn.setEnabled(True)
        self.export_excel_btn.setEnabled(True)

    def display_details_table(self, data):
        headers = ["Дата", "Предмет", "Учитель", "Тема урока", "Причина пропуска"]

        self.details_table.setColumnCount(len(headers))
        self.details_table.setHorizontalHeaderLabels(headers)
        self.details_table.setRowCount(len(data))

        for i, item in enumerate(data):
            self.details_table.setItem(i, 0, QTableWidgetItem(item["date"]))
            self.details_table.setItem(i, 1, QTableWidgetItem(item["subject"]))
            self.details_table.setItem(i, 2, QTableWidgetItem(item["teacher"]))
            self.details_table.setItem(i, 3, QTableWidgetItem(item["topic"]))
            self.details_table.setItem(i, 4, QTableWidgetItem(item["reason"]))

        self.details_table.resizeColumnsToContents()
        self.details_table.setColumnWidth(0, 100)
        self.details_table.setColumnWidth(1, 200)
        self.details_table.setColumnWidth(2, 180)
        self.details_table.setColumnWidth(3, 300)

    def filter_details_table(self):
        filter_text = self.filter_edit.text().lower()
        if not filter_text:
            self.display_details_table(self.missing_data)
            return

        filtered = [item for item in self.missing_data if filter_text in item["subject"].lower()]
        self.display_details_table(filtered)
        self.total_missing_label.setText(f"Показано: {len(filtered)} из {len(self.missing_data)} пропусков")

    # ------------------------------------------------------------------
    def export_to_csv(self):
        if not self.missing_data:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        class_name = self.selected_class["name"] if self.selected_class else "unknown"
        student_name = self.selected_student_name.replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Пропуски_{class_name}_{student_name}_{timestamp}.csv"

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить CSV", filename, "CSV files (*.csv)"
        )

        if not file_path:
            return

        try:
            with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f, delimiter=';')
                writer.writerow(["Класс", "ФИО ученика", "Дата", "Предмет", "Учитель", "Тема урока", "Причина пропуска"])

                for item in self.missing_data:
                    writer.writerow([
                        self.selected_class["name"] if self.selected_class else "",
                        self.selected_student_name,
                        item["date"],
                        item["subject"],
                        item["teacher"],
                        item["topic"],
                        item["reason"]
                    ])

            self.log_signal.emit(f"✅ CSV экспортирован: {file_path}")
            QMessageBox.information(self, "Готово", f"Файл сохранён:\n{file_path}")

        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка экспорта: {str(e)}")
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{str(e)}")

    def export_to_excel(self):
        if not self.missing_data:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта!")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Ошибка", "Модуль openpyxl не установлен")
            return

        class_name = self.selected_class["name"] if self.selected_class else "unknown"
        student_name = self.selected_student_name.replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Пропуски_{class_name}_{student_name}_{timestamp}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить Excel", filename, "Excel files (*.xlsx)"
        )

        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Пропуски занятий"

            headers = ["Класс", "ФИО ученика", "Дата", "Предмет", "Учитель", "Тема урока", "Причина пропуска"]
            ws.append(headers)

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
            header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for item in self.missing_data:
                ws.append([
                    self.selected_class["name"] if self.selected_class else "",
                    self.selected_student_name,
                    item["date"],
                    item["subject"],
                    item["teacher"],
                    item["topic"],
                    item["reason"]
                ])

            col_widths = {'A': 12, 'B': 35, 'C': 12, 'D': 25, 'E': 25, 'F': 40, 'G': 20}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            ws.auto_filter.ref = ws.dimensions

            wb.save(file_path)
            self.log_signal.emit(f"✅ Excel экспортирован: {file_path}")

            reply = QMessageBox.question(
                self, "Готово",
                f"Файл сохранён:\n{file_path}\n\nОткрыть файл?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)

        except Exception as e:
            self.log_signal.emit(f"❌ Ошибка экспорта: {str(e)}")
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{str(e)}")

    def append_to_console(self, text):
        self.console.append_text(text)
        QApplication.processEvents()