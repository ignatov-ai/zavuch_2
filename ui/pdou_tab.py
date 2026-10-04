# -*- coding: utf-8 -*-
"""
Вкладка «🎨 Кружки ПДОУ».
Загрузка списка групп ПДОУ через esz.mos.ru.
Требует ДВА токена + cookies:
• aupd_token (JWT RS256 из cookie aupd_token)
• esztoken  (JWT HS256 из Local Storage → eszToken)
• cookies   (session-cookie, mos_id, obr_id, subsystem_id, ...)
"""
import os
from datetime import datetime
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QLineEdit,
    QPushButton, QPlainTextEdit, QMessageBox, QFrame, QGroupBox,
    QTableWidget, QTableWidgetItem, QProgressBar, QComboBox,
    QFileDialog, QToolButton, QSizePolicy, QApplication,
)
from PySide6.QtCore import QThread, Signal, Qt
from PySide6.QtGui import QColor
from collector_pdou import PDOUCollector, PDOUToken, PDOUCookies


# ============================================================
# Сворачиваемый блок (единый стиль — фиолетовый)
# ============================================================
class CollapsibleGroupBox(QWidget):
    """Сворачиваемый блок с заголовком-кнопкой (▶ / ▼)."""

    def __init__(self, title: str, parent=None, collapsed: bool = True):
        super().__init__(parent)
        self._title_text = title
        self._status_suffix = ""
        self._collapsed = collapsed

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self._header_btn = QToolButton()
        self._header_btn.setToolButtonStyle(
            Qt.ToolButtonStyle.ToolButtonTextBesideIcon
        )
        self._header_btn.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed
        )
        self._header_btn.setCheckable(True)
        self._header_btn.setChecked(not collapsed)
        self._header_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self._header_btn.clicked.connect(self._on_header_clicked)
        self._header_btn.setStyleSheet("""
            QToolButton {
                border: 1px solid #c4b5fd;
                border-radius: 8px;
                background-color: #f5f3ff;
                padding: 9px 14px;
                font-weight: 700;
                font-size: 11pt;
                color: #6d28d9;
                text-align: left;
            }
            QToolButton:hover {
                background-color: #ede9fe;
                color: #5b21b6;
            }
            QToolButton:checked {
                border-bottom-left-radius: 0px;
                border-bottom-right-radius: 0px;
                border-bottom: none;
                background-color: #ede9fe;
            }
        """)
        main_layout.addWidget(self._header_btn)

        self._content = QWidget()
        self._content.setObjectName("CollapsibleContent")
        self._content.setStyleSheet("""
            QWidget#CollapsibleContent {
                background-color: #ffffff;
                border: 1px solid #c4b5fd;
                border-top: none;
                border-bottom-left-radius: 8px;
                border-bottom-right-radius: 8px;
            }
        """)
        self._content_layout = QVBoxLayout(self._content)
        self._content_layout.setContentsMargins(14, 12, 14, 14)
        self._content_layout.setSpacing(8)
        main_layout.addWidget(self._content)

        self._apply_collapsed_state()
        self._update_header_text()

    def _on_header_clicked(self):
        self._collapsed = not self._header_btn.isChecked()
        self._apply_collapsed_state()
        self._update_header_text()

    def _apply_collapsed_state(self):
        if self._collapsed:
            self._content.setVisible(False)
            self._header_btn.setArrowType(Qt.ArrowType.RightArrow)
        else:
            self._content.setVisible(True)
            self._header_btn.setArrowType(Qt.ArrowType.DownArrow)

    def _update_header_text(self):
        arrow = "▶" if self._collapsed else "▼"
        suffix = f"  —  {self._status_suffix}" if self._status_suffix else ""
        self._header_btn.setText(f"{arrow}  {self._title_text}{suffix}")

    def set_title(self, title: str):
        self._title_text = title
        self._update_header_text()

    def set_collapsed(self, collapsed: bool):
        self._collapsed = collapsed
        self._header_btn.setChecked(not collapsed)
        self._apply_collapsed_state()
        self._update_header_text()

    def is_collapsed(self) -> bool:
        return self._collapsed

    def content_layout(self) -> QVBoxLayout:
        return self._content_layout

    def set_status_suffix(self, text: str):
        self._status_suffix = text or ""
        self._update_header_text()


# ============================================================
# Потоки
# ============================================================
class PDOUTokenCheckThread(QThread):
    finished = Signal(bool, str, list, str)
    log = Signal(str)

    def __init__(self, aupd_token, esztoken=""):
        super().__init__()
        self.aupd_token = aupd_token
        self.esztoken = esztoken

    def run(self):
        try:
            collector = PDOUCollector(self.aupd_token, self.esztoken)
            collector.log_callback = lambda t: self.log.emit(t)
            ok, user_name, roles, reason = collector.check_token()
            self.finished.emit(ok, user_name, roles, reason)
        except Exception as e:
            self.finished.emit(False, "", [], f"Исключение: {e}")


class PDOULoadThread(QThread):
    finished = Signal(list)
    error = Signal(str)
    log_message = Signal(str)

    def __init__(self, aupd_token, esztoken):
        super().__init__()
        self.aupd_token = aupd_token
        self.esztoken = esztoken

    def run(self):
        try:
            collector = PDOUCollector(self.aupd_token, self.esztoken)
            collector.log_callback = lambda t: self.log_message.emit(t)
            groups = collector.get_all_groups()
            self.finished.emit(groups)
        except Exception as e:
            import traceback
            err = f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}"
            self.error.emit(err)
            self.finished.emit([])


# ============================================================
# Единый стиль (фиолетовый — ПДОУ)
# ============================================================
TAB_STYLE = """
    QWidget {
        font-size: 10pt;
        color: #1e293b;
    }
    QGroupBox {
        font-size: 11pt;
        font-weight: 700;
        color: #7c3aed;
        border: 1px solid #c4b5fd;
        border-radius: 10px;
        margin-top: 12px;
        padding: 18px 14px 14px 14px;
        background: #ffffff;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 12px;
        padding: 0 6px;
    }
    QLineEdit {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QLineEdit:focus { border: 1px solid #7c3aed; }
    QLineEdit:disabled {
        background: #f1f5f9;
        color: #94a3b8;
    }
    QPushButton {
        min-height: 30px;
        padding: 4px 12px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #f8fafc;
        color: #1e293b;
        font-weight: 500;
    }
    QPushButton:hover { background: #e2e8f0; }
    QPushButton:pressed { background: #cbd5e1; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
    QComboBox {
        min-height: 30px;
        padding: 2px 8px;
        border: 1px solid #cbd5e1;
        border-radius: 5px;
        background: #ffffff;
    }
    QPlainTextEdit {
        font-family: Consolas, "Courier New", monospace;
        font-size: 9.5pt;
    }
    QTableWidget {
        border: 1px solid #e2e8f0;
        border-radius: 6px;
        gridline-color: #e2e8f0;
    }
    QTableWidget::item { padding: 3px 6px; }
    QHeaderView::section {
        background: #f1f5f9;
        color: #1e293b;
        padding: 5px;
        border: none;
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
        font-weight: 600;
    }
    QProgressBar {
        min-height: 22px;
        border: 1px solid #cbd5e1;
        border-radius: 6px;
        background: #f8fafc;
        text-align: center;
        color: #1e293b;
        font-weight: 600;
    }
    QProgressBar::chunk {
        background-color: #7c3aed;
        border-radius: 5px;
    }
    QLabel { color: #334155; }
"""

PRIMARY_BTN = """
    QPushButton {
        background-color: #7c3aed;
        color: #ffffff;
        font-weight: 700;
        border: none;
        border-radius: 7px;
    }
    QPushButton:hover { background-color: #6d28d9; }
    QPushButton:pressed { background-color: #5b21b6; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

SUCCESS_BTN = """
    QPushButton {
        background-color: #059669;
        color: #ffffff;
        font-weight: 600;
        border: none;
        border-radius: 6px;
    }
    QPushButton:hover { background-color: #047857; }
    QPushButton:disabled {
        background-color: #cbd5e1;
        color: #64748b;
    }
"""

DANGER_BTN = """
    QPushButton {
        background-color: #fff1f2;
        color: #be123c;
        border: 1px solid #fecdd3;
        font-weight: 600;
    }
    QPushButton:hover { background-color: #ffe4e6; }
    QPushButton:disabled {
        background: #f1f5f9;
        color: #94a3b8;
        border-color: #e2e8f0;
    }
"""


# ============================================================
# Вкладка
# ============================================================
class PDOUTab(QWidget):
    """Вкладка просмотра групп ПДОУ (кружки и секции)."""

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent
        self.current_token = ""
        self.current_esztoken = ""
        self.current_user_name = ""
        self.current_roles = []
        self.all_groups = []
        self.filtered_groups = []
        self.token_check_thread = None
        self.load_thread = None
        self._log_count = 0
        self.initUI()
        self._load_saved_token()
        self._load_saved_cookies()

    # ================================================================
    # UI
    # ================================================================
    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(12)
        main_layout.setContentsMargins(16, 14, 16, 14)

        # === Заголовок ===
        title = QLabel("🎨  Кружки ПДОУ")
        title.setStyleSheet(
            "font-size: 16pt; font-weight: 700; color: #0f172a; padding: 2px 0 6px 0;"
        )
        main_layout.addWidget(title)

        subtitle = QLabel(
            "Загрузка списка кружков и секций через esz.mos.ru. "
            "Требуются токены aupd_token и esztoken."
        )
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet("color: #64748b; margin-bottom: 4px;")
        main_layout.addWidget(subtitle)

        # === СВОРАЧИВАЕМЫЙ БЛОК ТОКЕНОВ ===
        self.token_group = CollapsibleGroupBox(
            "🔑 Токены ПДОУ (aupd_token + esztoken)", collapsed=True
        )
        token_widget = QWidget()
        token_grid = QGridLayout(token_widget)
        token_grid.setVerticalSpacing(10)
        token_grid.setHorizontalSpacing(10)
        token_grid.setContentsMargins(0, 0, 0, 0)

        token_grid.addWidget(QLabel("aupd_token:"), 0, 0)
        self.token_edit = QLineEdit()
        self.token_edit.setPlaceholderText("JWT RS256 (cookie aupd_token)")
        self.token_edit.setMinimumHeight(34)
        token_grid.addWidget(self.token_edit, 0, 1)
        self.paste_btn = QPushButton("📋 Из буфера")
        self.paste_btn.setMaximumWidth(140)
        self.paste_btn.setMinimumHeight(32)
        self.paste_btn.setStyleSheet("""
            QPushButton {
                background: #f5f3ff; color: #7c3aed;
                border: 1px solid #ddd6fe; font-weight: 600;
            }
            QPushButton:hover { background: #ede9fe; }
        """)
        self.paste_btn.clicked.connect(self.on_paste_token)
        token_grid.addWidget(self.paste_btn, 0, 2)

        token_grid.addWidget(QLabel("esztoken:"), 1, 0)
        self.esztoken_edit = QLineEdit()
        self.esztoken_edit.setPlaceholderText(
            "JWT HS256 (Local Storage → eszToken)"
        )
        self.esztoken_edit.setMinimumHeight(34)
        token_grid.addWidget(self.esztoken_edit, 1, 1)
        self.paste_esztoken_btn = QPushButton("📋 Из буфера")
        self.paste_esztoken_btn.setMaximumWidth(140)
        self.paste_esztoken_btn.setMinimumHeight(32)
        self.paste_esztoken_btn.setStyleSheet("""
            QPushButton {
                background: #f5f3ff; color: #7c3aed;
                border: 1px solid #ddd6fe; font-weight: 600;
            }
            QPushButton:hover { background: #ede9fe; }
        """)
        self.paste_esztoken_btn.clicked.connect(self.on_paste_esztoken)
        token_grid.addWidget(self.paste_esztoken_btn, 1, 2)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(6)
        self.check_token_btn = QPushButton("🔍 Проверить")
        self.check_token_btn.setMinimumHeight(32)
        self.check_token_btn.setStyleSheet("""
            QPushButton {
                background: #f5f3ff; color: #7c3aed;
                border: 1px solid #ddd6fe; font-weight: 600;
            }
            QPushButton:hover { background: #ede9fe; }
        """)
        self.check_token_btn.clicked.connect(self.on_check_token)
        btn_row.addWidget(self.check_token_btn)

        self.save_token_btn = QPushButton("💾 Сохранить")
        self.save_token_btn.setMinimumHeight(32)
        self.save_token_btn.setStyleSheet(SUCCESS_BTN)
        self.save_token_btn.clicked.connect(self.on_save_token)
        btn_row.addWidget(self.save_token_btn)

        self.clear_token_btn = QPushButton("🗑 Очистить")
        self.clear_token_btn.setMinimumHeight(32)
        self.clear_token_btn.setStyleSheet(DANGER_BTN)
        self.clear_token_btn.clicked.connect(self.on_clear_token)
        btn_row.addWidget(self.clear_token_btn)
        btn_row.addStretch()
        token_grid.addLayout(btn_row, 2, 0, 1, 3)

        self.token_status_label = QLabel("⏳ Токены не заданы")
        self.token_status_label.setStyleSheet("""
            color: #6d28d9; background: #f5f3ff;
            border-radius: 6px; padding: 7px 10px;
            font-weight: 600;
        """)
        self.token_status_label.setWordWrap(True)
        token_grid.addWidget(self.token_status_label, 3, 0, 1, 3)

        self.token_group.content_layout().addWidget(token_widget)
        main_layout.addWidget(self.token_group)

        # === СВОРАЧИВАЕМЫЙ БЛОК COOKIES ===
        self.cookies_group = CollapsibleGroupBox(
            "🍪 Cookies ПДОУ", collapsed=True
        )
        cookies_widget = QWidget()
        cookies_layout = QVBoxLayout(cookies_widget)
        cookies_layout.setContentsMargins(0, 0, 0, 0)
        cookies_layout.setSpacing(8)

        hint = QLabel(
            "ℹ️  Вставьте cookies (F12 → Network → любой запрос к esz.mos.ru → "
            "Request Headers → Cookie:)."
        )
        hint.setStyleSheet("color: #64748b; font-size: 9pt;")
        hint.setWordWrap(True)
        cookies_layout.addWidget(hint)

        self.cookies_edit = QPlainTextEdit()
        self.cookies_edit.setPlaceholderText(
            "session-cookie=...; Ltpatoken2=...; mos_id=...; ..."
        )
        self.cookies_edit.setMaximumHeight(80)
        self.cookies_edit.setStyleSheet("""
            QPlainTextEdit {
                background-color: #0f172a;
                color: #e2e8f0;
                border: 1px solid #334155;
                border-radius: 6px;
                padding: 6px;
                font-family: Consolas, "Courier New", monospace;
                font-size: 9pt;
            }
        """)
        cookies_layout.addWidget(self.cookies_edit)

        cookies_btn_row = QHBoxLayout()
        cookies_btn_row.setSpacing(6)
        self.cookies_paste_btn = QPushButton("📋 Из буфера")
        self.cookies_paste_btn.setMinimumHeight(30)
        self.cookies_paste_btn.setStyleSheet("""
            QPushButton {
                background: #f5f3ff; color: #7c3aed;
                border: 1px solid #ddd6fe; font-weight: 600;
            }
            QPushButton:hover { background: #ede9fe; }
        """)
        self.cookies_paste_btn.clicked.connect(self.on_cookies_paste)
        cookies_btn_row.addWidget(self.cookies_paste_btn)

        self.cookies_save_btn = QPushButton("💾 Сохранить cookies")
        self.cookies_save_btn.setMinimumHeight(30)
        self.cookies_save_btn.setStyleSheet(SUCCESS_BTN)
        self.cookies_save_btn.clicked.connect(self.on_cookies_save)
        cookies_btn_row.addWidget(self.cookies_save_btn)

        self.cookies_clear_btn = QPushButton("🗑 Очистить")
        self.cookies_clear_btn.setMinimumHeight(30)
        self.cookies_clear_btn.setStyleSheet(DANGER_BTN)
        self.cookies_clear_btn.clicked.connect(self.on_cookies_clear)
        cookies_btn_row.addWidget(self.cookies_clear_btn)
        cookies_btn_row.addStretch()
        cookies_layout.addLayout(cookies_btn_row)

        self.cookies_status_label = QLabel("⏳ Cookies не заданы")
        self.cookies_status_label.setStyleSheet("""
            color: #6d28d9; background: #f5f3ff;
            border-radius: 6px; padding: 7px 10px;
            font-weight: 600;
        """)
        self.cookies_status_label.setWordWrap(True)
        cookies_layout.addWidget(self.cookies_status_label)

        self.cookies_group.content_layout().addWidget(cookies_widget)
        main_layout.addWidget(self.cookies_group)

        # === ПАНЕЛЬ УПРАВЛЕНИЯ ===
        control_group = QGroupBox("⚙️  Управление")
        control_layout = QHBoxLayout(control_group)
        control_layout.setSpacing(10)

        self.load_btn = QPushButton("📋  Загрузить список кружков/ПДОУ")
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(38)
        self.load_btn.setStyleSheet(PRIMARY_BTN)
        self.load_btn.clicked.connect(self.load_groups)
        control_layout.addWidget(self.load_btn)
        control_layout.addStretch()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("""
            color: #6d28d9; background: #f5f3ff;
            border-radius: 6px; padding: 6px 10px;
            font-weight: 700; font-size: 10pt;
        """)
        control_layout.addWidget(self.stats_label)
        main_layout.addWidget(control_group)

        # === ПРОГРЕСС-БАР ===
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        main_layout.addWidget(self.progress_bar)

        # === ФИЛЬТРЫ ===
        filter_group = QGroupBox("🔎  Фильтры")
        filter_group.setStyleSheet("""
            QGroupBox {
                font-size: 11pt; font-weight: 700;
                color: #0369a1; border: 1px solid #7dd3fc;
                border-radius: 10px; margin-top: 12px;
                padding: 18px 14px 14px 14px; background: #ffffff;
            }
            QGroupBox::title {
                subcontrol-origin: margin; left: 12px; padding: 0 6px;
            }
        """)
        filter_layout = QHBoxLayout(filter_group)
        filter_layout.setSpacing(10)
        filter_layout.setContentsMargins(10, 6, 10, 10)

        filter_layout.addWidget(QLabel("Поиск:"))
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText(
            "Название группы / программа / педагог..."
        )
        self.search_edit.setMinimumWidth(280)
        self.search_edit.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.search_edit)

        filter_layout.addWidget(QLabel("Статус:"))
        self.status_filter = QComboBox()
        self.status_filter.addItem("Все", "all")
        self.status_filter.addItem("Активна", 1)
        self.status_filter.addItem("Идёт обучение", 2)
        self.status_filter.addItem("Завершена", 3)
        self.status_filter.setMinimumWidth(160)
        self.status_filter.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.status_filter)

        self.clear_filters_btn = QPushButton("Очистить")
        self.clear_filters_btn.setStyleSheet(DANGER_BTN)
        self.clear_filters_btn.clicked.connect(self.clear_filters)
        filter_layout.addWidget(self.clear_filters_btn)
        filter_layout.addStretch()
        main_layout.addWidget(filter_group)

        # === ТАБЛИЦА ===
        self.groups_table = QTableWidget()
        self.groups_table.setAlternatingRowColors(True)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.groups_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.groups_table.horizontalHeader().setStretchLastSection(True)
        self.groups_table.setColumnCount(8)
        self.groups_table.setHorizontalHeaderLabels([
            "Код", "Название группы", "Педагог", "Программа",
            "Даты обучения", "Ёмкость", "Записано", "Статус",
        ])
        self.groups_table.verticalHeader().setDefaultSectionSize(50)
        main_layout.addWidget(self.groups_table, 1)

        # === ЭКСПОРТ ===
        export_layout = QHBoxLayout()
        export_layout.addStretch()
        self.export_btn = QPushButton("📊  Экспорт в Excel")
        self.export_btn.setEnabled(False)
        self.export_btn.setMinimumHeight(34)
        self.export_btn.setStyleSheet(SUCCESS_BTN)
        self.export_btn.clicked.connect(self.export_to_excel)
        export_layout.addWidget(self.export_btn)
        main_layout.addLayout(export_layout)

        # === СВОРАЧИВАЕМЫЙ ЖУРНАЛ ===
        self.log_group = CollapsibleGroupBox("📋  Журнал", collapsed=True)
        self.console = QPlainTextEdit()
        self.console.setReadOnly(True)
        self.console.setMinimumHeight(140)
        self.console.setMaximumHeight(220)
        self.console.setStyleSheet("""
            QPlainTextEdit {
                background-color: #0f172a;
                color: #e2e8f0;
                border: 1px solid #334155;
                border-radius: 6px;
                padding: 6px;
                font-family: Consolas, "Courier New", monospace;
                font-size: 9.5pt;
            }
        """)
        self.log_group.content_layout().addWidget(self.console)
        main_layout.addWidget(self.log_group)

        self.setStyleSheet(TAB_STYLE)

    # ================================================================
    # ЛОГ
    # ================================================================
    def _log(self, text):
        text = str(text)
        self.console.appendPlainText(text)
        sb = self.console.verticalScrollBar()
        sb.setValue(sb.maximum())
        self._log_count += 1
        self.log_group.set_status_suffix(f"{self._log_count} стр.")

    def on_auth_updated(self, auth):
        pass

    # ================================================================
    # ТОКЕНЫ
    # ================================================================
    def _load_saved_token(self):
        data = PDOUToken.load()
        token = data.get("aupd_token", "")
        esztoken = data.get("esztoken", "")
        user_name = data.get("user_name", "")
        roles = data.get("user_roles", [])
        if token:
            self.token_edit.setText(token)
            self.current_token = token
        if esztoken:
            self.esztoken_edit.setText(esztoken)
            self.current_esztoken = esztoken
        if token and esztoken:
            self.current_user_name = user_name
            self.current_roles = roles
            self._update_token_status(True, user_name, roles)
            self._log(f"[i] Загружены сохранённые токены ({user_name or '?'})")
        elif token:
            self._update_token_status(False)
            self._log("[i] aupd_token загружен, но esztoken отсутствует")
        else:
            self._update_token_status(False)
        self._update_load_btn_state()

    def on_paste_token(self):
        text = QApplication.clipboard().text().strip()
        text = text.replace("…", "").replace("...", "").strip()
        if text.startswith("aupd_token="):
            text = text[len("aupd_token="):].strip()
        if text:
            self.token_edit.clear()
            self.token_edit.setText(text)
            self._log(f"[clipboard] aupd_token ({len(text)} символов)")

    def on_paste_esztoken(self):
        text = QApplication.clipboard().text().strip()
        text = text.replace("…", "").replace("...", "").strip()
        for prefix in ("eszToken=", "esztoken=", "esztoken: "):
            if text.startswith(prefix):
                text = text[len(prefix):].strip()
                break
        text = text.strip('"').strip("'")
        if text:
            self.esztoken_edit.clear()
            self.esztoken_edit.setText(text)
            self._log(f"[clipboard] esztoken ({len(text)} символов)")

    def on_check_token(self):
        token = self.token_edit.text().strip().replace("…", "").replace("...", "").strip()
        esztoken = self.esztoken_edit.text().strip().replace("…", "").replace("...", "").strip()
        if not token:
            QMessageBox.warning(self, "Ошибка", "Введите aupd_token.")
            return
        self.token_edit.setText(token)
        self.esztoken_edit.setText(esztoken)
        self.check_token_btn.setEnabled(False)
        self._log("[i] Проверяю токены...")
        self.token_check_thread = PDOUTokenCheckThread(token, esztoken)
        self.token_check_thread.log.connect(self._log)
        self.token_check_thread.finished.connect(self._on_token_checked)
        self.token_check_thread.start()

    def _on_token_checked(self, ok, user_name, roles, reason):
        self.check_token_btn.setEnabled(True)
        if ok:
            self.current_token = self.token_edit.text().strip()
            self.current_esztoken = self.esztoken_edit.text().strip()
            self.current_user_name = user_name
            self.current_roles = roles
            self._update_token_status(True, user_name, roles)
            self._update_load_btn_state()
            roles_text = "\n".join(f"• {r}" for r in roles) if roles else "(нет)"
            has_pdou = any(
                "оператор" in r.lower() or "пдоу" in r.lower() or "круж" in r.lower()
                for r in roles
            )
            if has_pdou:
                QMessageBox.information(
                    self, "Проверка успешна",
                    f"✅ Токены рабочие.\n\nПользователь: {user_name}\n\n"
                    f"Роли ЕСЗ:\n{roles_text}"
                )
            else:
                QMessageBox.warning(
                    self, "Нет прав ПДОУ",
                    f"⚠️ Токены рабочие, но прав на ПДОУ не видно.\n\n"
                    f"Пользователь: {user_name}\n\nРоли ЕСЗ:\n{roles_text}"
                )
        else:
            self._update_token_status(False)
            self._update_load_btn_state()
            self._log(f"[!] {reason}")
            rl = reason.lower()
            if "«…»" in reason or "обрезан" in rl:
                QMessageBox.warning(
                    self, "Токен обрезан",
                    f"{reason}\n\nСкопируйте полное значение."
                )
            elif "401" in reason:
                QMessageBox.warning(
                    self, "401 Unauthorized",
                    f"{reason}\n\n"
                    "Проверьте: cookies ПДОУ вставлены? Токены не истекли?"
                )
            else:
                QMessageBox.warning(self, "Проверка не удалась", reason)

    def on_save_token(self):
        token = self.token_edit.text().strip().replace("…", "").replace("...", "").strip()
        esztoken = self.esztoken_edit.text().strip().replace("…", "").replace("...", "").strip()
        if not token:
            QMessageBox.warning(self, "Ошибка", "Введите aupd_token.")
            return
        if not esztoken:
            QMessageBox.warning(
                self, "Ошибка",
                "Введите esztoken (Local Storage → eszToken)."
            )
            return
        user_name = self.current_user_name or ""
        roles = self.current_roles or []
        if PDOUToken.save(token, user_name, roles, esztoken):
            self.current_token = token
            self.current_esztoken = esztoken
            self._update_token_status(True, user_name, roles)
            self._update_load_btn_state()
            self._log(f"[+] Токены сохранены ({user_name or 'без имени'})")
            QMessageBox.information(
                self, "Готово",
                f"✅ Токены сохранены.\n\nПользователь: {user_name or 'неизвестен'}"
            )
        else:
            QMessageBox.critical(self, "Ошибка", "Не удалось сохранить.")

    def on_clear_token(self):
        reply = QMessageBox.question(
            self, "Очистить токены?",
            "Удалить сохранённые токены ПДОУ?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        PDOUToken.clear()
        self.token_edit.clear()
        self.esztoken_edit.clear()
        self.current_token = ""
        self.current_esztoken = ""
        self.current_user_name = ""
        self.current_roles = []
        self._update_token_status(False)
        self._update_load_btn_state()
        self._log("[i] Токены удалены")

    def _update_token_status(self, ok, user_name="", roles=None):
        if ok:
            roles = roles or []
            has_pdou = any(
                "оператор" in r.lower() or "пдоу" in r.lower() or "круж" in r.lower()
                for r in roles
            )
            marker = "✅" if has_pdou else "⚠️"
            roles_text = ", ".join(roles) if roles else "нет ролей"
            self.token_status_label.setText(
                f"{marker} Пользователь: {user_name or 'неизвестен'}\n"
                f"Роли ЕСЗ: {roles_text}"
            )
            if has_pdou:
                self.token_status_label.setStyleSheet("""
                    color: #065f46; background: #ecfdf5;
                    border: 1px solid #a7f3d0;
                    border-radius: 6px; padding: 7px 10px;
                    font-weight: 600;
                """)
            else:
                self.token_status_label.setStyleSheet("""
                    color: #92400e; background: #fffbeb;
                    border: 1px solid #fde68a;
                    border-radius: 6px; padding: 7px 10px;
                    font-weight: 600;
                """)
            self.token_group.set_status_suffix(f"{marker} {user_name or 'OK'}")
        else:
            self.token_status_label.setText("❌ Токены не заданы или не работают")
            self.token_status_label.setStyleSheet("""
                color: #991b1b; background: #fef2f2;
                border: 1px solid #fecaca;
                border-radius: 6px; padding: 7px 10px;
                font-weight: 600;
            """)
            self.token_group.set_status_suffix("❌ не заданы")

    # ================================================================
    # COOKIES
    # ================================================================
    def _load_saved_cookies(self):
        cookies = PDOUCookies.load()
        if cookies:
            self._log(f"[i] Загружено cookies: {len(cookies)}")
            self._update_cookies_status(True, list(cookies.keys()))
        else:
            self._update_cookies_status(False)

    def on_cookies_paste(self):
        text = QApplication.clipboard().text().strip()
        if text:
            self.cookies_edit.setPlainText(text)
            self._log(f"[clipboard] Cookies вставлены ({len(text)} символов)")

    def on_cookies_save(self):
        text = self.cookies_edit.toPlainText().strip()
        if not text:
            QMessageBox.warning(self, "Ошибка", "Вставьте cookies.")
            return
        cookies = PDOUCookies.parse_cookie_string(text)
        if not cookies:
            QMessageBox.warning(
                self, "Ошибка",
                "Не удалось разобрать cookies.\nФормат: name1=value1; name2=value2"
            )
            return
        if PDOUCookies.save(cookies):
            self._log(f"[+] Сохранено {len(cookies)} cookies")
            self._update_cookies_status(True, list(cookies.keys()))
            self._update_load_btn_state()
            has_ltpa = "Ltpatoken2" in cookies
            msg = f"✅ Сохранено cookies: {len(cookies)}"
            if not has_ltpa:
                msg += ("\n\n⚠️ Ltpatoken2 отсутствует. "
                        "Для чтения групп это не критично, но для других "
                        "операций (заявления) он нужен.")
            QMessageBox.information(self, "Готово", msg)
        else:
            QMessageBox.critical(self, "Ошибка", "Не удалось сохранить cookies.")

    def on_cookies_clear(self):
        reply = QMessageBox.question(
            self, "Очистить cookies?",
            "Удалить сохранённые cookies ПДОУ?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return
        PDOUCookies.clear()
        self.cookies_edit.clear()
        self._update_cookies_status(False)
        self._update_load_btn_state()
        self._log("[i] Cookies очищены")

    def _update_cookies_status(self, ok, keys=None):
        if ok and keys:
            has_ltpa = "Ltpatoken2" in keys
            marker = "✅" if has_ltpa else "⚠️"
            self.cookies_status_label.setText(
                f"{marker} Cookies: {len(keys)}\n"
                f"Ключи: {', '.join(keys[:6])}..."
            )
            if has_ltpa:
                self.cookies_status_label.setStyleSheet("""
                    color: #065f46; background: #ecfdf5;
                    border: 1px solid #a7f3d0;
                    border-radius: 6px; padding: 7px 10px;
                    font-weight: 600;
                """)
            else:
                self.cookies_status_label.setStyleSheet("""
                    color: #92400e; background: #fffbeb;
                    border: 1px solid #fde68a;
                    border-radius: 6px; padding: 7px 10px;
                    font-weight: 600;
                """)
            self.cookies_group.set_status_suffix(f"{marker} {len(keys)} шт.")
        else:
            self.cookies_status_label.setText("❌ Cookies не заданы")
            self.cookies_status_label.setStyleSheet("""
                color: #991b1b; background: #fef2f2;
                border: 1px solid #fecaca;
                border-radius: 6px; padding: 7px 10px;
                font-weight: 600;
            """)
            self.cookies_group.set_status_suffix("❌ не заданы")

    # ================================================================
    # ЗАГРУЗКА ГРУПП
    # ================================================================
    def _update_load_btn_state(self):
        has_token = bool(self.current_token)
        has_esztoken = bool(self.current_esztoken)
        self.load_btn.setEnabled(has_token and has_esztoken)

    def load_groups(self):
        if not self.current_token or not self.current_esztoken:
            self.token_group.set_collapsed(False)
            QMessageBox.warning(
                self, "Ошибка",
                "Введите и сохраните оба токена (aupd_token и esztoken)."
            )
            return
        self.load_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setFormat("Загрузка групп ПДОУ...")
        self.groups_table.setRowCount(0)
        self.all_groups = []
        self.console.clear()
        self._log_count = 0
        self.log_group.set_status_suffix("0 стр.")
        self.load_thread = PDOULoadThread(
            self.current_token, self.current_esztoken
        )
        self.load_thread.finished.connect(self.on_load_finished)
        self.load_thread.error.connect(self.on_load_error)
        self.load_thread.log_message.connect(self._log)
        self.load_thread.start()

    def on_load_finished(self, groups):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)
        if not groups:
            QMessageBox.warning(
                self, "Внимание",
                "Не удалось получить группы. Подробности в журнале."
            )
            return
        self.all_groups = groups
        self.apply_filter()
        self.update_stats()
        self.export_btn.setEnabled(True)
        QMessageBox.information(
            self, "Готово", f"✅ Загружено групп: {len(groups)}"
        )

    def on_load_error(self, error):
        self.progress_bar.setVisible(False)
        self.progress_bar.setRange(0, 100)
        self.load_btn.setEnabled(True)
        QMessageBox.critical(self, "Ошибка загрузки", f"Ошибка:\n{error}")

    # ================================================================
    # ФИЛЬТРЫ / СТАТИСТИКА
    # ================================================================
    def update_stats(self):
        total = len(self.all_groups)
        active = sum(1 for g in self.all_groups if g.get("serviceClassStatus") == 1)
        running = sum(1 for g in self.all_groups if g.get("serviceClassStatus") == 2)
        done = sum(1 for g in self.all_groups if g.get("serviceClassStatus") == 3)
        self.stats_label.setText(
            f"Всего: {total}  |  ✅ {active}  |  ▶ {running}  |  ✓ {done}"
        )

    def apply_filter(self):
        search_text = self.search_edit.text().lower().strip()
        status_filter = self.status_filter.currentData()
        filtered = []
        for g in self.all_groups:
            if search_text:
                haystack = " ".join([
                    str(g.get("name", "")),
                    str(g.get("serviceName", "")),
                    str(g.get("supervisorPerson", "")),
                    str(g.get("code", "")),
                ]).lower()
                if search_text not in haystack:
                    continue
            if status_filter != "all":
                if g.get("serviceClassStatus") != status_filter:
                    continue
            filtered.append(g)
        self.filtered_groups = filtered
        self.populate_table(filtered)

    def clear_filters(self):
        self.search_edit.clear()
        self.status_filter.setCurrentIndex(0)
        self.apply_filter()

    # ================================================================
    # ТАБЛИЦА
    # ================================================================
    def populate_table(self, groups):
        self.groups_table.setSortingEnabled(False)
        self.groups_table.setUpdatesEnabled(False)
        self.groups_table.setRowCount(len(groups))

        green_bg = QColor(209, 250, 229)
        blue_bg = QColor(219, 234, 254)
        gray_bg = QColor(241, 245, 249)

        for i, g in enumerate(groups):
            status = g.get("serviceClassStatus")
            if status == 1:
                row_color = green_bg
            elif status == 2:
                row_color = blue_bg
            elif status == 3:
                row_color = gray_bg
            else:
                row_color = None

            def make_item(text, align=Qt.AlignmentFlag.AlignLeft
                          | Qt.AlignmentFlag.AlignVCenter):
                it = QTableWidgetItem(str(text) if text is not None else "")
                it.setTextAlignment(align)
                if row_color:
                    it.setBackground(row_color)
                return it

            self.groups_table.setItem(
                i, 0, make_item(g.get("code", ""), Qt.AlignmentFlag.AlignCenter)
            )
            name_item = make_item((g.get("name", "") or "").strip())
            name_item.setToolTip(f"ID: {g.get('id')}")
            self.groups_table.setItem(i, 1, name_item)
            self.groups_table.setItem(i, 2, make_item(g.get("supervisorPerson", "")))
            prog_item = make_item((g.get("serviceName", "") or "").strip())
            prog_item.setToolTip(f"serviceId: {g.get('serviceId')}")
            self.groups_table.setItem(i, 3, prog_item)
            self.groups_table.setItem(
                i, 4, make_item(g.get("trainDates", ""), Qt.AlignmentFlag.AlignCenter)
            )
            self.groups_table.setItem(
                i, 5, make_item(g.get("capacity", 0), Qt.AlignmentFlag.AlignCenter)
            )
            self.groups_table.setItem(
                i, 6, make_item(g.get("included", 0), Qt.AlignmentFlag.AlignCenter)
            )
            status_text = PDOUCollector.get_status_text(status)
            status_item = make_item(status_text, Qt.AlignmentFlag.AlignCenter)
            if status == 1:
                status_item.setForeground(QColor(6, 95, 70))
            elif status == 2:
                status_item.setForeground(QColor(30, 64, 175))
            elif status == 3:
                status_item.setForeground(QColor(71, 85, 105))
            self.groups_table.setItem(i, 7, status_item)

        self.groups_table.setUpdatesEnabled(True)
        self.groups_table.setSortingEnabled(True)
        self.groups_table.setColumnWidth(0, 100)
        self.groups_table.setColumnWidth(1, 280)
        self.groups_table.setColumnWidth(2, 220)
        self.groups_table.setColumnWidth(3, 320)
        self.groups_table.setColumnWidth(4, 190)
        self.groups_table.setColumnWidth(5, 90)
        self.groups_table.setColumnWidth(6, 90)
        self.groups_table.setColumnWidth(7, 130)

    # ================================================================
    # ЭКСПОРТ
    # ================================================================
    def export_to_excel(self):
        if not self.filtered_groups:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Ошибка", "Модуль openpyxl не установлен")
            return
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Кружки_ПДОУ_{timestamp}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт", filename, "Excel files (*.xlsx)"
        )
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Кружки ПДОУ"
            headers = ["Код", "Название группы", "Педагог", "Программа",
                       "Даты обучения", "Ёмкость", "Записано", "Статус"]
            ws.append(headers)
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(
                start_color="7C3AED", end_color="7C3AED", fill_type="solid"
            )
            header_alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col)
                c.font = header_font
                c.fill = header_fill
                c.alignment = header_alignment
            for g in self.filtered_groups:
                status_text = PDOUCollector.get_status_text(
                    g.get("serviceClassStatus")
                )
                ws.append([
                    g.get("code", ""),
                    (g.get("name", "") or "").strip(),
                    g.get("supervisorPerson", ""),
                    (g.get("serviceName", "") or "").strip(),
                    g.get("trainDates", ""),
                    g.get("capacity", 0),
                    g.get("included", 0),
                    status_text,
                ])
            col_widths = {"A": 12, "B": 40, "C": 28, "D": 45,
                          "E": 24, "F": 10, "G": 10, "H": 16}
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
            wb.save(file_path)
            reply = QMessageBox.question(
                self, "Готово",
                f"Отчёт сохранён:\n{file_path}\n\nОткрыть?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить:\n{e}")