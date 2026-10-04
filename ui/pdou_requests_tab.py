# -*- coding: utf-8 -*-
"""
Вкладка «📋 Заявления ПДОУ».

Таблица заявлений с фильтрами и пагинацией.
Данные тянутся из /Request/Search + /Request?id + /Learner/Education/List.
"""
import os
from datetime import datetime

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QLineEdit,
    QPushButton, QPlainTextEdit, QMessageBox, QGroupBox,
    QTableWidget, QTableWidgetItem, QProgressBar, QComboBox,
    QFileDialog, QToolButton, QSizePolicy, QApplication, QSpinBox
)
from PySide6.QtCore import QThread, Signal, Qt
from PySide6.QtGui import QColor

from collector_pdou_requests import PDOURequestsCollector
from collector_pdou import PDOUToken, PDOUCookies


# ============================================================
#  Поток: загрузка одной страницы
# ============================================================
class RequestsPageLoadThread(QThread):
    """
    Загружает одну страницу заявлений и обогащает каждую запись.
    Прогресс: (done, total_on_page).
    """
    finished = Signal(list, int)     # (enriched, total)
    error = Signal(str)
    log = Signal(str)
    page_progress = Signal(int, int)  # (done, total_on_page)

    def __init__(self, aupd_token, esztoken, page, page_size,
                 search_text="", status_id=None):
        super().__init__()
        self.aupd_token = aupd_token
        self.esztoken = esztoken
        self.page = page
        self.page_size = page_size
        self.search_text = search_text
        self.status_id = status_id

    def run(self):
        try:
            collector = PDOURequestsCollector(self.aupd_token, self.esztoken)
            collector.log_callback = lambda t: self.log.emit(t)

            def on_progress(done, total):
                self.page_progress.emit(done, total)

            items, total = collector.load_page(
                page_number=self.page,
                page_size=self.page_size,
                search_text=self.search_text,
                status_id=self.status_id,
                progress_callback=on_progress,
            )
            self.finished.emit(items, total)
        except Exception as e:
            import traceback
            self.error.emit(f"{type(e).__name__}: {e}\n\n{traceback.format_exc()}")
            self.finished.emit([], 0)


# ============================================================
#  Сворачиваемый блок (как в pdou_tab.py)
# ============================================================
class CollapsibleGroupBox(QWidget):
    def __init__(self, title: str, parent=None, collapsed: bool = True):
        super().__init__(parent)
        self._title_text = title
        self._status_suffix = ""
        self._collapsed = collapsed

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self._header_btn = QToolButton()
        self._header_btn.setToolButtonStyle(
            Qt.ToolButtonStyle.ToolButtonTextBesideIcon
        )
        self._header_btn.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed
        )
        self._header_btn.setCheckable(True)
        self._header_btn.setChecked(not collapsed)
        self._header_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self._header_btn.clicked.connect(self._on_header_clicked)
        self._header_btn.setStyleSheet("""
            QToolButton {
                border: 2px solid #8b5cf6;
                border-radius: 10px;
                background-color: #f5f3ff;
                padding: 8px 12px;
                font-weight: bold;
                font-size: 12pt;
                color: #4c1d95;
                text-align: left;
            }
            QToolButton:hover { background-color: #ede9fe; }
            QToolButton:checked {
                border-bottom-left-radius: 0px;
                border-bottom-right-radius: 0px;
                border-bottom: none;
            }
        """)
        main_layout.addWidget(self._header_btn)

        self._content = QWidget()
        self._content.setObjectName("CollapsibleContent")
        self._content.setStyleSheet("""
            QWidget#CollapsibleContent {
                background-color: #ffffff;
                border: 2px solid #8b5cf6;
                border-top: none;
                border-bottom-left-radius: 10px;
                border-bottom-right-radius: 10px;
            }
        """)
        self._content_layout = QVBoxLayout(self._content)
        self._content_layout.setContentsMargins(12, 10, 12, 12)
        self._content_layout.setSpacing(8)

        main_layout.addWidget(self._content)

        self._apply_collapsed_state()
        self._update_header_text()

    def _on_header_clicked(self):
        self._collapsed = not self._header_btn.isChecked()
        self._apply_collapsed_state()
        self._update_header_text()

    def _apply_collapsed_state(self):
        if self._collapsed:
            self._content.setVisible(False)
            self._header_btn.setArrowType(Qt.ArrowType.RightArrow)
        else:
            self._content.setVisible(True)
            self._header_btn.setArrowType(Qt.ArrowType.DownArrow)

    def _update_header_text(self):
        arrow = "▶" if self._collapsed else "▼"
        suffix = f"  —  {self._status_suffix}" if self._status_suffix else ""
        self._header_btn.setText(f"{arrow}  {self._title_text}{suffix}")

    def set_collapsed(self, collapsed: bool):
        self._collapsed = collapsed
        self._header_btn.setChecked(not collapsed)
        self._apply_collapsed_state()
        self._update_header_text()

    def content_layout(self) -> QVBoxLayout:
        return self._content_layout

    def set_status_suffix(self, text: str):
        self._status_suffix = text or ""
        self._update_header_text()


# ============================================================
#  Вкладка
# ============================================================
class PDOURequestsTab(QWidget):
    """Вкладка «📋 Заявления ПДОУ»."""

    def __init__(self, parent):
        super().__init__()
        self.main_window = parent

        self.current_token = ""
        self.current_esztoken = ""
        self.current_user_name = ""
        self.current_roles = []

        self.current_page = 1
        self.page_size = 50
        self.total_records = 0
        self.total_pages = 1
        self.current_items = []
        self.filtered_items = []

        self.load_thread = None
        self._log_count = 0

        self.initUI()
        self._load_saved_token()

    # ================================================================
    #  UI
    # ================================================================
    def initUI(self):
        main_layout = QVBoxLayout(self)
        main_layout.setSpacing(10)
        main_layout.setContentsMargins(15, 15, 15, 15)

        # === Токены ===
        self.token_group = CollapsibleGroupBox(
            "🔑 Токены ПДОУ (aupd_token + esztoken)", collapsed=True
        )
        token_widget = QWidget()
        token_grid = QGridLayout(token_widget)
        token_grid.setVerticalSpacing(8)
        token_grid.setHorizontalSpacing(10)
        token_grid.setContentsMargins(0, 0, 0, 0)

        token_grid.addWidget(QLabel("aupd_token:"), 0, 0)
        self.token_edit = QLineEdit()
        self.token_edit.setPlaceholderText("JWT RS256 (cookie aupd_token)")
        self.token_edit.setMinimumHeight(34)
        token_grid.addWidget(self.token_edit, 0, 1)

        self.paste_btn = QPushButton("📋 Из буфера")
        self.paste_btn.setMaximumWidth(140)
        self.paste_btn.clicked.connect(self.on_paste_token)
        token_grid.addWidget(self.paste_btn, 0, 2)

        token_grid.addWidget(QLabel("esztoken:"), 1, 0)
        self.esztoken_edit = QLineEdit()
        self.esztoken_edit.setPlaceholderText(
            "JWT HS256 (Local Storage → eszToken)"
        )
        self.esztoken_edit.setMinimumHeight(34)
        token_grid.addWidget(self.esztoken_edit, 1, 1)

        self.paste_esztoken_btn = QPushButton("📋 Из буфера")
        self.paste_esztoken_btn.setMaximumWidth(140)
        self.paste_esztoken_btn.clicked.connect(self.on_paste_esztoken)
        token_grid.addWidget(self.paste_esztoken_btn, 1, 2)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(6)

        self.save_token_btn = QPushButton("💾 Сохранить")
        self.save_token_btn.setMinimumHeight(32)
        self.save_token_btn.setStyleSheet("""
            QPushButton { background-color: #059669; color: white;
                font-weight: bold; border-radius: 6px; }
            QPushButton:hover { background-color: #047857; }
        """)
        self.save_token_btn.clicked.connect(self.on_save_token)
        btn_row.addWidget(self.save_token_btn)
        btn_row.addStretch()
        token_grid.addLayout(btn_row, 2, 0, 1, 3)

        self.token_status_label = QLabel("⏳ Токены не заданы")
        self.token_status_label.setStyleSheet(
            "color: #666; font-weight: bold; padding: 4px;"
        )
        self.token_status_label.setWordWrap(True)
        token_grid.addWidget(self.token_status_label, 3, 0, 1, 3)

        self.token_group.content_layout().addWidget(token_widget)
        main_layout.addWidget(self.token_group)

        # === Управление ===
        control_group = QGroupBox("Управление")
        control_layout = QHBoxLayout(control_group)

        self.load_btn = QPushButton("📋 Загрузить заявления")
        self.load_btn.setEnabled(False)
        self.load_btn.setMinimumHeight(34)
        self.load_btn.setStyleSheet("""
            QPushButton { background-color: #2196F3; color: white;
                font-weight: bold; border-radius: 6px; }
            QPushButton:hover { background-color: #1976D2; }
            QPushButton:disabled { background-color: #cccccc; color: #666666; }
        """)
        self.load_btn.clicked.connect(self.load_page)
        control_layout.addWidget(self.load_btn)

        control_layout.addStretch()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("font-weight: bold;")
        control_layout.addWidget(self.stats_label)

        main_layout.addWidget(control_group)

        # === Прогресс-бар ===
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.progress_bar.setStyleSheet("""
            QProgressBar { border: 1px solid #cccccc; border-radius: 4px;
                text-align: center; height: 24px; }
            QProgressBar::chunk { background-color: #4CAF50; border-radius: 4px; }
        """)
        main_layout.addWidget(self.progress_bar)

        # === Фильтры ===
        filter_group = QGroupBox("Фильтры (применяются к текущей странице)")
        filter_layout = QHBoxLayout(filter_group)

        filter_layout.addWidget(QLabel("Поиск:"))
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("ФИО / группа / программа")
        self.search_edit.setMinimumWidth(220)
        self.search_edit.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.search_edit)

        filter_layout.addWidget(QLabel("ФИО:"))
        self.filter_fio = QLineEdit()
        self.filter_fio.setPlaceholderText("фамилия")
        self.filter_fio.setMaximumWidth(140)
        self.filter_fio.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_fio)

        filter_layout.addWidget(QLabel("Группа:"))
        self.filter_group = QLineEdit()
        self.filter_group.setPlaceholderText("имя группы")
        self.filter_group.setMaximumWidth(140)
        self.filter_group.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_group)

        filter_layout.addWidget(QLabel("Класс:"))
        self.filter_class = QLineEdit()
        self.filter_class.setPlaceholderText("10а")
        self.filter_class.setMaximumWidth(80)
        self.filter_class.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_class)

        filter_layout.addWidget(QLabel("Корпус:"))
        self.filter_building = QComboBox()
        self.filter_building.addItem("Все", None)
        self.filter_building.setMinimumWidth(160)
        self.filter_building.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_building)

        filter_layout.addWidget(QLabel("Статус:"))
        self.filter_status = QComboBox()
        self.filter_status.addItem("Все", None)
        self.filter_status.setMinimumWidth(180)
        self.filter_status.currentIndexChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.filter_status)

        self.clear_filters_btn = QPushButton("Очистить")
        self.clear_filters_btn.clicked.connect(self.clear_filters)
        filter_layout.addWidget(self.clear_filters_btn)

        filter_layout.addStretch()
        main_layout.addWidget(filter_group)

        # === Пагинация ===
        page_group = QGroupBox("Пагинация")
        page_layout = QHBoxLayout(page_group)

        self.prev_btn = QPushButton("◀ Предыдущая")
        self.prev_btn.setMinimumHeight(32)
        self.prev_btn.clicked.connect(self.prev_page)
        page_layout.addWidget(self.prev_btn)

        page_layout.addWidget(QLabel("Страница:"))
        self.page_spin = QSpinBox()
        self.page_spin.setMinimum(1)
        self.page_spin.setMaximum(1)
        self.page_spin.setValue(1)
        self.page_spin.setMaximumWidth(90)
        self.page_spin.valueChanged.connect(self.on_page_spin_changed)
        page_layout.addWidget(self.page_spin)

        self.total_pages_label = QLabel("/ 1")
        self.total_pages_label.setStyleSheet("font-weight: bold;")
        page_layout.addWidget(self.total_pages_label)

        page_layout.addWidget(QLabel("  Записей на странице:"))
        self.page_size_combo = QComboBox()
        self.page_size_combo.addItems(["25", "50", "100", "200"])
        self.page_size_combo.setCurrentText("50")
        self.page_size_combo.setMaximumWidth(90)
        self.page_size_combo.currentTextChanged.connect(self.on_page_size_changed)
        page_layout.addWidget(self.page_size_combo)

        self.next_btn = QPushButton("Следующая ▶")
        self.next_btn.setMinimumHeight(32)
        self.next_btn.clicked.connect(self.next_page)
        page_layout.addWidget(self.next_btn)

        page_layout.addStretch()

        self.total_records_label = QLabel("Всего: 0")
        self.total_records_label.setStyleSheet("font-weight: bold;")
        page_layout.addWidget(self.total_records_label)

        main_layout.addWidget(page_group)

        # === Таблица ===
        self.table = QTableWidget()
        self.table.setAlternatingRowColors(True)
        self.table.setSortingEnabled(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels([
            "Класс обучения", "Класс группы", "ФИО",
            "Кружок", "Группа", "Статус",
        ])
        self.table.verticalHeader().setDefaultSectionSize(44)
        main_layout.addWidget(self.table)

        # === Экспорт ===
        export_layout = QHBoxLayout()
        export_layout.addStretch()

        self.export_page_btn = QPushButton("📊 Экспорт текущей страницы")
        self.export_page_btn.setEnabled(False)
        self.export_page_btn.setMinimumHeight(32)
        self.export_page_btn.clicked.connect(self.export_current_page)
        export_layout.addWidget(self.export_page_btn)

        self.export_all_btn = QPushButton("📊 Экспорт всех (медленно!)")
        self.export_all_btn.setEnabled(False)
        self.export_all_btn.setMinimumHeight(32)
        self.export_all_btn.setToolTip(
            "Загрузит ВСЕ заявления через пагинацию. Это ~8000 записей\n"
            "и несколько минут работы."
        )
        self.export_all_btn.clicked.connect(self.export_all)
        export_layout.addWidget(self.export_all_btn)

        main_layout.addLayout(export_layout)

        # === Журнал ===
        self.log_group = CollapsibleGroupBox("📋 Журнал", collapsed=True)
        self.console = QPlainTextEdit()
        self.console.setReadOnly(True)
        self.console.setMinimumHeight(120)
        self.console.setMaximumHeight(200)
        self.console.setStyleSheet(
            "font-family: Consolas, monospace; font-size: 10px; "
            "background-color: #1e1e1e; color: #d4d4d4; "
            "border-radius: 6px; padding: 6px;"
        )
        self.log_group.content_layout().addWidget(self.console)
        main_layout.addWidget(self.log_group)

    # ================================================================
    #  Лог
    # ================================================================
    def _log(self, text):
        text = str(text)
        self.console.appendPlainText(text)
        sb = self.console.verticalScrollBar()
        sb.setValue(sb.maximum())
        self._log_count += 1
        self.log_group.set_status_suffix(f"{self._log_count} стр.")

    def on_auth_updated(self, auth):
        # ПДОУ работает на своём токене, ЭЖД не нужен
        pass

    # ================================================================
    #  Токены
    # ================================================================
    def _load_saved_token(self):
        data = PDOUToken.load()
        token = data.get("aupd_token", "")
        esztoken = data.get("esztoken", "")
        user_name = data.get("user_name", "")

        if token:
            self.token_edit.setText(token)
            self.current_token = token
        if esztoken:
            self.esztoken_edit.setText(esztoken)
            self.current_esztoken = esztoken

        if token and esztoken:
            self.current_user_name = user_name
            self._update_token_status(True, user_name)
            self._log(f"[i] Загружены сохранённые токены ({user_name or '?'})")
        else:
            self._update_token_status(False)

        self._update_load_btn_state()

    def on_paste_token(self):
        text = QApplication.clipboard().text().strip()
        text = text.replace("…", "").replace("...", "").strip()
        if text.startswith("aupd_token="):
            text = text[len("aupd_token="):].strip()
        if text:
            self.token_edit.setText(text)
            self._log(f"[clipboard] aupd_token ({len(text)} символов)")

    def on_paste_esztoken(self):
        text = QApplication.clipboard().text().strip()
        text = text.replace("…", "").replace("...", "").strip()
        for prefix in ("eszToken=", "esztoken=", "esztoken: "):
            if text.startswith(prefix):
                text = text[len(prefix):].strip()
                break
        text = text.strip('"').strip("'")
        if text:
            self.esztoken_edit.setText(text)
            self._log(f"[clipboard] esztoken ({len(text)} символов)")

    def on_save_token(self):
        token = self.token_edit.text().strip()
        esztoken = self.esztoken_edit.text().strip()
        if not token or not esztoken:
            QMessageBox.warning(self, "Ошибка", "Введите оба токена.")
            return
        PDOUToken.save(token, self.current_user_name, [], esztoken)
        self.current_token = token
        self.current_esztoken = esztoken
        self._update_token_status(True, self.current_user_name)
        self._update_load_btn_state()
        QMessageBox.information(self, "Готово", "✅ Токены сохранены.")

    def _update_token_status(self, ok, user_name=""):
        if ok:
            self.token_status_label.setText(
                f"✅ Пользователь: {user_name or 'неизвестен'}"
            )
            self.token_status_label.setStyleSheet(
                "color: #059669; font-weight: bold; padding: 4px;"
            )
            self.token_group.set_status_suffix(f"✅ {user_name or 'OK'}")
        else:
            self.token_status_label.setText("❌ Токены не заданы")
            self.token_status_label.setStyleSheet(
                "color: #dc2626; font-weight: bold; padding: 4px;"
            )
            self.token_group.set_status_suffix("❌ не заданы")

    def _update_load_btn_state(self):
        has_token = bool(self.current_token)
        has_esztoken = bool(self.current_esztoken)
        self.load_btn.setEnabled(has_token and has_esztoken)

    # ================================================================
    #  Загрузка страницы
    # ================================================================
    def load_page(self):
        if not self.current_token or not self.current_esztoken:
            self.token_group.set_collapsed(False)
            QMessageBox.warning(self, "Ошибка",
                                "Введите и сохраните оба токена.")
            return

        self.load_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat(
            f"Загрузка страницы {self.current_page}..."
        )
        self.table.setRowCount(0)

        self.load_thread = RequestsPageLoadThread(
            self.current_token, self.current_esztoken,
            self.current_page, self.page_size,
            search_text="",
            status_id=None,
        )
        self.load_thread.log.connect(self._log)
        self.load_thread.page_progress.connect(self._on_page_progress)
        self.load_thread.finished.connect(self.on_page_loaded)
        self.load_thread.error.connect(self.on_load_error)
        self.load_thread.start()

    def _on_page_progress(self, done, total):
        if total <= 0:
            return
        percent = int(done / total * 100)
        self.progress_bar.setValue(percent)
        self.progress_bar.setFormat(
            f"Загрузка заявлений: {done}/{total} ({percent}%)"
        )
        QApplication.processEvents()

    def on_page_loaded(self, items, total):
        self.progress_bar.setVisible(False)
        self.load_btn.setEnabled(True)

        self.current_items = items
        self.total_records = total

        if total > 0:
            self.total_pages = max(
                1, (total + self.page_size - 1) // self.page_size
            )
        else:
            self.total_pages = 1

        self.page_spin.blockSignals(True)
        self.page_spin.setMaximum(self.total_pages)
        self.page_spin.setValue(self.current_page)
        self.page_spin.blockSignals(False)
        self.total_pages_label.setText(f"/ {self.total_pages}")
        self.total_records_label.setText(f"Всего: {total}")

        self.prev_btn.setEnabled(self.current_page > 1)
        self.next_btn.setEnabled(self.current_page < self.total_pages)

        if items:
            self.export_page_btn.setEnabled(True)
            self.export_all_btn.setEnabled(True)

        self._rebuild_filter_options()
        self.apply_filter()

        self._log(
            f"[+] Страница {self.current_page}/{self.total_pages}: "
            f"{len(items)} записей (всего {total})"
        )

    def on_load_error(self, error):
        self.progress_bar.setVisible(False)
        self.load_btn.setEnabled(True)
        QMessageBox.critical(self, "Ошибка загрузки", f"Ошибка:\n{error}")

    # ================================================================
    #  Пагинация
    # ================================================================
    def prev_page(self):
        if self.current_page > 1:
            self.current_page -= 1
            self.load_page()

    def next_page(self):
        if self.current_page < self.total_pages:
            self.current_page += 1
            self.load_page()

    def on_page_spin_changed(self, value):
        if value != self.current_page:
            self.current_page = value
            self.load_page()

    def on_page_size_changed(self, text):
        try:
            new_size = int(text)
        except ValueError:
            return
        if new_size != self.page_size:
            self.page_size = new_size
            self.current_page = 1
            if self.current_items:
                self.load_page()

    # ================================================================
    #  Фильтры (клиентские)
    # ================================================================
    def _rebuild_filter_options(self):
        # Корпуса
        current_b = self.filter_building.currentData()
        self.filter_building.blockSignals(True)
        self.filter_building.clear()
        self.filter_building.addItem("Все", None)
        buildings = set()
        for it in self.current_items:
            b = it.get("building", "")
            if b:
                buildings.add(b)
        for b in sorted(buildings):
            self.filter_building.addItem(b, b)
        idx = self.filter_building.findData(current_b)
        if idx >= 0:
            self.filter_building.setCurrentIndex(idx)
        self.filter_building.blockSignals(False)

        # Статусы
        current_s = self.filter_status.currentData()
        self.filter_status.blockSignals(True)
        self.filter_status.clear()
        self.filter_status.addItem("Все", None)
        statuses = set()
        for it in self.current_items:
            s = it.get("status_name", "")
            if s:
                statuses.add(s)
        for s in sorted(statuses):
            self.filter_status.addItem(s, s)
        idx = self.filter_status.findData(current_s)
        if idx >= 0:
            self.filter_status.setCurrentIndex(idx)
        self.filter_status.blockSignals(False)

    def apply_filter(self):
        search = self.search_edit.text().lower().strip()
        fio = self.filter_fio.text().lower().strip()
        group = self.filter_group.text().lower().strip()
        klass = self.filter_class.text().lower().strip()
        building = self.filter_building.currentData()
        status = self.filter_status.currentData()

        filtered = []
        for it in self.current_items:
            if search:
                haystack = " ".join([
                    it.get("child_name", ""),
                    it.get("applicant_name", ""),
                    it.get("service_name", ""),
                    it.get("group_name", ""),
                    it.get("service_class_code", ""),
                ]).lower()
                if search not in haystack:
                    continue

            if fio and fio not in it.get("child_name", "").lower():
                continue
            if group and group not in it.get("group_name", "").lower():
                continue

            if klass:
                gc = it.get("group_class", "").lower()
                lc = it.get("learner_class", "").lower()
                if klass not in gc and klass not in lc:
                    continue

            if building is not None:
                if it.get("building", "") != building:
                    continue

            if status is not None:
                if it.get("status_name", "") != status:
                    continue

            filtered.append(it)

        self.filtered_items = filtered
        self._populate_table(filtered)

    def clear_filters(self):
        self.search_edit.clear()
        self.filter_fio.clear()
        self.filter_group.clear()
        self.filter_class.clear()
        self.filter_building.setCurrentIndex(0)
        self.filter_status.setCurrentIndex(0)
        self.apply_filter()

    # ================================================================
    #  Таблица
    # ================================================================
    @staticmethod
    def _building_color(building: str) -> QColor:
        if not building:
            return QColor(255, 235, 235)
        palette = {
            "Маршала Захарова": QColor(219, 234, 254),
            "Домодедовская": QColor(220, 252, 231),
            "Совхоз им. Ленина": QColor(254, 249, 195),
            "ЗИЛ / Лихачёва": QColor(237, 233, 254),
            "Елецкая": QColor(255, 237, 213),
            "Шипиловская": QColor(226, 232, 240),
        }
        return palette.get(building, QColor(241, 245, 249))

    def _populate_table(self, items):
        self.table.setSortingEnabled(False)
        self.table.setUpdatesEnabled(False)
        self.table.setRowCount(len(items))

        for i, it in enumerate(items):
            def make(text, align=Qt.AlignmentFlag.AlignLeft
                     | Qt.AlignmentFlag.AlignVCenter):
                item = QTableWidgetItem(str(text) if text is not None else "")
                item.setTextAlignment(align)
                return item

            # 1. Класс обучения
            self.table.setItem(i, 0, make(
                it.get("learner_class", "") or "—",
                Qt.AlignmentFlag.AlignCenter,
            ))

            # 2. Класс группы
            self.table.setItem(i, 1, make(
                it.get("group_class", "") or "—",
                Qt.AlignmentFlag.AlignCenter,
            ))

            # 3. ФИО
            fio_item = make(it.get("child_name", ""))
            fio_item.setToolTip(
                f"Заявитель: {it.get('applicant_name', '')}\n"
                f"Телефон: {it.get('applicant_phone', '')}\n"
                f"Заявка: {it.get('request_number', '')}"
            )
            self.table.setItem(i, 2, fio_item)

            # 4. Кружок (serviceName)
            service = it.get("service_name", "")
            s_item = make(service)
            s_item.setToolTip(service)
            b = it.get("building", "")
            if b:
                s_item.setBackground(self._building_color(b))
            self.table.setItem(i, 3, s_item)

            # 5. Группа (serviceClassName из деталей)
            group_name = it.get("group_name", "") or "—"
            g_item = make(group_name)
            g_item.setToolTip(
                f"Код группы: {it.get('service_class_code', '')}\n"
                f"Группа: {group_name}"
            )
            self.table.setItem(i, 4, g_item)

            # 6. Статус
            status = it.get("status_name", "")
            st_item = make(status, Qt.AlignmentFlag.AlignCenter)
            st_id = it.get("status_id")
            if st_id == 16:
                st_item.setBackground(QColor(255, 245, 200))
            elif st_id == 23:
                st_item.setBackground(QColor(255, 230, 200))
            elif st_id == 1:
                st_item.setBackground(QColor(220, 255, 220))
            self.table.setItem(i, 5, st_item)

        self.table.setUpdatesEnabled(True)
        self.table.setSortingEnabled(True)
        self.table.setColumnWidth(0, 100)
        self.table.setColumnWidth(1, 100)
        self.table.setColumnWidth(2, 220)
        self.table.setColumnWidth(3, 320)
        self.table.setColumnWidth(4, 260)
        self.table.setColumnWidth(5, 220)

        self.stats_label.setText(
            f"Показано: {len(items)} | Всего на странице: {len(self.current_items)} "
            f"| Всего записей: {self.total_records}"
        )

    # ================================================================
    #  Экспорт
    # ================================================================
    def export_current_page(self):
        if not self.filtered_items:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта.")
            return
        self._export_items(self.filtered_items, "текущей_страницы")

    def export_all(self):
        if not self.current_token or not self.current_esztoken:
            QMessageBox.warning(self, "Ошибка", "Нужны токены.")
            return

        reply = QMessageBox.question(
            self, "Экспорт всех заявлений",
            f"Будет загружено ~{self.total_records} записей.\n"
            f"Это может занять несколько минут.\n\nПродолжить?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat("Загрузка всех заявлений...")

        try:
            collector = PDOURequestsCollector(
                self.current_token, self.current_esztoken
            )
            collector.log_callback = self._log

            def on_progress(done, total):
                if total > 0:
                    percent = int(done / total * 100)
                    self.progress_bar.setValue(percent)
                    self.progress_bar.setFormat(
                        f"Загрузка: {done}/{total} ({percent}%)"
                    )
                    QApplication.processEvents()

            items, _ = collector.load_all(
                page_size=self.page_size,
                progress_callback=on_progress,
            )
        finally:
            self.progress_bar.setVisible(False)

        if not items:
            QMessageBox.warning(self, "Ошибка", "Ничего не загружено.")
            return

        self._export_items(items, "все")

    def _export_items(self, items, suffix):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Ошибка", "openpyxl не установлен.")
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Заявления_ПДОУ_{suffix}_{timestamp}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить отчёт", filename, "Excel files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Заявления ПДОУ"

            headers = [
                "№ заявки", "Класс обучения", "Класс группы", "ФИО",
                "Заявитель", "Телефон", "Корпус", "Кружок", "Группа",
                "Код группы", "Статус", "Дата заявки", "Номер договора",
            ]
            ws.append(headers)

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(
                start_color="8b5cf6", end_color="8b5cf6", fill_type="solid"
            )
            header_alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col)
                c.font = header_font
                c.fill = header_fill
                c.alignment = header_alignment

            for it in items:
                ws.append([
                    it.get("request_number", ""),
                    it.get("learner_class", ""),
                    it.get("group_class", ""),
                    it.get("child_name", ""),
                    it.get("applicant_name", ""),
                    it.get("applicant_phone", ""),
                    it.get("building", ""),
                    it.get("service_name", ""),
                    it.get("group_name", ""),
                    it.get("service_class_code", ""),
                    it.get("status_name", ""),
                    it.get("request_date", ""),
                    it.get("contract_number", ""),
                ])

            col_widths = {
                "A": 14, "B": 14, "C": 12, "D": 30, "E": 30,
                "F": 16, "G": 20, "H": 45, "I": 35, "J": 14,
                "K": 30, "L": 20, "M": 16,
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
            wb.save(file_path)

            reply = QMessageBox.question(
                self, "Готово",
                f"Отчёт сохранён:\n{file_path}\n\nОткрыть?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)

        except Exception as e:
            QMessageBox.critical(self, "Ошибка",
                                 f"Не удалось сохранить:\n{e}")