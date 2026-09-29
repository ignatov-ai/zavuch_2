# -*- coding: utf-8 -*-
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

from academic_calendar import AcademicCalendar  # Добавляем импорт


class JournalChecker:
    """Класс для проверки журналов"""

    def __init__(self):
        self.data_str = ['1', '2', '3', '4', '5']
        self.curriculum_data = {}
        self.teacher_cache = {}  # Кэш для имен учителей
        self.debug_mode = False  # Режим отладки (по умолчанию выключен)

        # Даты триместров (для 5-9 классов) - используем для совместимости со старым кодом
        self.TRIMESTER_DATES = {
            'Т1': {'start': (1, 9), 'end': (30, 11), 'name': 'Триместр 1'},
            'Т2': {'start': (1, 12), 'end': (28, 2), 'name': 'Триместр 2'},
            'Т3': {'start': (1, 3), 'end': (31, 5), 'name': 'Триместр 3'}
        }

        # Даты полугодий (для 10-11 классов)
        self.SEMESTER_DATES = {
            'П1': {'start': (1, 9), 'end': (31, 12), 'name': 'Полугодие 1'},
            'П2': {'start': (1, 1), 'end': (31, 5), 'name': 'Полугодие 2'}
        }

        # Объединяем все периоды
        self.PERIOD_DATES = {}
        self.PERIOD_DATES.update(self.TRIMESTER_DATES)
        self.PERIOD_DATES.update(self.SEMESTER_DATES)

    def get_teacher_from_sheet(self, sheet):
        """Извлечение имени учителя из листа"""
        try:
            if sheet.cell(row=3, column=1).value:
                teacher_cell = str(sheet.cell(row=3, column=1).value)
                # Проверяем разные форматы
                if "Учитель:" in teacher_cell:
                    teacher = teacher_cell.replace("Учитель:", "").strip()
                    if teacher and teacher != "Не указан":
                        return teacher
                elif teacher_cell and teacher_cell != "Не указан" and teacher_cell.strip():
                    # Если это не "Не указан" и не пустая строка, возвращаем как есть
                    return teacher_cell.strip()
        except:
            pass
        return "Не указан"

    def load_curriculum_data(self, curriculum_file):
        """Загружает данные учебного плана из файла"""
        curriculum_dict = {}

        try:
            if not os.path.exists(curriculum_file):
                print(f"Файл учебного плана не найден: {curriculum_file}")
                return {}

            wb = load_workbook(curriculum_file, data_only=True)
            for sheet in wb.worksheets:
                max_row = sheet.max_row
                max_col = sheet.max_column

                classes_in_sheet = []
                for col in range(2, max_col + 1):
                    cell_value = sheet.cell(row=1, column=col).value
                    if cell_value:
                        class_name = str(cell_value)
                        # Обработка формата класса (5А, 5-А, 5а и т.д.)
                        if len(class_name) > 1 and class_name[1].isalpha():
                            class_name = f'{class_name[0]}-{class_name[1:]}'
                        if '-з' in class_name.lower():
                            class_name = class_name.upper().replace('-З', '-З')
                        else:
                            class_name = class_name.upper()
                        classes_in_sheet.append(class_name)

                for row in range(2, max_row + 1):
                    subject = sheet.cell(row=row, column=1).value
                    if not subject:
                        continue

                    subject = str(subject).strip()

                    for col_idx, class_name in enumerate(classes_in_sheet, start=2):
                        hours = sheet.cell(row=row, column=col_idx).value
                        if hours and isinstance(hours, (int, float)):
                            if class_name not in curriculum_dict:
                                curriculum_dict[class_name] = {}
                            curriculum_dict[class_name][subject] = int(hours)

            wb.close()
            return curriculum_dict

        except Exception as e:
            print(f"Ошибка загрузки учебного плана: {e}")
            return {}

    def parse_date(self, day, month, year=None):
        """Преобразует день и месяц в объект datetime"""
        return AcademicCalendar.parse_date(day, month, year)

    def parse_custom_date(self, date_str):
        """Парсит дату из строки в формате ДД.ММ.ГГГГ"""
        return AcademicCalendar.parse_custom_date(date_str)

    def find_all_mark_columns(self, sheet, log_callback=None):
        """Находит все столбцы с отметками на листе"""
        mark_columns = []
        max_col = min(sheet.max_column, 1000)

        for col in range(3, max_col + 1):
            if sheet.cell(row=6, column=col).value == 'оц':
                month_val = sheet.cell(row=4, column=col).value
                day_val = sheet.cell(row=5, column=col).value

                if month_val is not None and day_val is not None:
                    try:
                        month = int(month_val)
                        day = int(day_val)
                        date_obj = self.parse_date(day, month)

                        if date_obj:
                            mark_columns.append({
                                'col': col,
                                'month': month,
                                'day': day,
                                'date_obj': date_obj,
                                'date_key': f"{day:02d}.{month:02d}"
                            })
                    except (ValueError, TypeError):
                        continue

        return mark_columns

    def get_min_marks_required(self, lessons_per_week):
        """Рассчитывает минимальное количество отметок"""
        if lessons_per_week == 0:
            return 0
        elif lessons_per_week == 1:
            return 3
        elif lessons_per_week == 2:
            return 5
        else:
            return 7

    def calculate_average(self, marks_sum, marks_count):
        """Вычисляет средний балл"""
        if marks_count == 0:
            return 0
        return round(marks_sum / marks_count, 2)

    def get_grade_from_average(self, avg_mark):
        """Определяет итоговую отметку по среднему баллу"""
        if avg_mark == 0:
            return 'НПА'
        elif avg_mark < 2.6:
            return 'А/З'
        elif avg_mark < 3.6:
            return 3
        elif avg_mark < 4.6:
            return 4
        else:
            return 5

    def parse_mark_cell(self, cell_value):
        """
        Парсит ячейку с отметками.
        """
        if not cell_value:
            return [], [], [], False

        cell_str = str(cell_value).strip()

        marks = []
        absences = []
        sickness = []
        has_absence = False

        # Разделяем по запятой
        parts = cell_str.split(',')

        for part in parts:
            part = part.strip()

            # Обработка пропуска
            if part == 'н':
                absences.append('н')
                has_absence = True
                continue

            # Обработка болезни
            if part == 'б':
                sickness.append('б')
                has_absence = True
                continue

            # Обработка отметки с весом
            if '(' in part and ')' in part:
                try:
                    mark_part = part.split('(')[0].strip()
                    coeff_part = part.split('(')[1].split(')')[0].strip()

                    coeff_clean = ''
                    for char in coeff_part:
                        if char.isdigit():
                            coeff_clean += char

                    is_control = 'К' in coeff_part

                    if mark_part in ['2', '3', '4', '5']:
                        mark_int = int(mark_part)
                        coeff = int(coeff_clean) if coeff_clean else 1
                        marks.append((mark_int, coeff, is_control))
                except:
                    pass

            # Обработка простой отметки
            elif part in ['2', '3', '4', '5']:
                marks.append((int(part), 1, False))

        return marks, absences, sickness, has_absence

    def get_period_dates(self, check_mode, start_date, end_date, log_callback=None):
        """
        Определяет границы периода проверки
        """
        if check_mode == 'custom' and start_date and end_date:
            if isinstance(start_date, str):
                period_start = self.parse_custom_date(start_date)
            else:
                period_start = start_date

            if isinstance(end_date, str):
                period_end = self.parse_custom_date(end_date)
            else:
                period_end = end_date

            if period_start and period_end:
                period_name = f"{period_start.strftime('%d.%m.%Y')} - {period_end.strftime('%d.%m.%Y')}"
            else:
                period_name = "пользовательский период"
        else:
            # Используем AcademicCalendar для предопределенных периодов
            period_start, period_end, period_name = AcademicCalendar.get_period_dates(check_mode)

        return period_start, period_end, period_name

    def check_journal_file(self, file_path, curriculum_data, check_mode, start_date=None, end_date=None,
                           log_callback=None):
        """
        Проверка одного файла журнала с агрегацией по ученикам
        """
        try:
            file_name = os.path.basename(file_path).replace('.xlsx', '')

            if '_журнал' in file_name:
                class_part = file_name.split('_журнал')[0]
                class_name = class_part.replace('_', '-')
            else:
                class_name = file_name.replace('_', '-')

            if log_callback:
                log_callback(f"\n📄 Проверка: {class_name}")
                if class_name.upper() in curriculum_data:
                    subjects = list(curriculum_data[class_name.upper()].keys())
                    log_callback(f"    ✅ Класс найден в УП, предметов: {len(subjects)}")
                else:
                    log_callback(f"    ⚠️ Класс НЕ найден в учебном плане!")

            book = load_workbook(file_path, data_only=True)

            # Словарь для агрегации данных по ученикам и предметам (только для замечаний)
            student_data = defaultdict(lambda: {
                'marks_sum': 0,
                'marks_count': 0,
                'marks_count_clear': 0,
                'propuski_count': 0,
                'sickness_count': 0,
                'lessons_count': 0,
                'trimestr_mark': None,
                'lessons_per_week': 0,
                'min_marks_required': 0,
                'teacher_name': 'Не указан',
                'seen_groups': set()
            })

            # Словарь для полной статистики по всем ученикам (для второго листа)
            full_stats = defaultdict(lambda: {
                'marks_sum': 0,
                'marks_count': 0,
                'marks_count_clear': 0,
                'propuski_count': 0,
                'sickness_count': 0,
                'lessons_count': 0,
                'lessons_per_week': 0,
                'min_marks_required': 0,
                'avg_mark': 0,
                'predicted_mark': '',
                'teacher_name': 'Не указан',
                'subject_name': '',
                'seen_groups': set()
            })

            # Определяем границы периода проверки
            period_start, period_end, period_name = self.get_period_dates(
                check_mode, start_date, end_date, log_callback
            )

            if not period_start or not period_end:
                return [], [], "Не удалось определить период проверки"

            if log_callback:
                log_callback(
                    f"    📅 Период проверки: {period_start.strftime('%d.%m.%Y')} - {period_end.strftime('%d.%m.%Y')}")

            # Обрабатываем каждый лист (предмет)
            for sheet_name in book.sheetnames:
                sheet = book[sheet_name]

                subject_name = "Неизвестно"
                if sheet.cell(row=1, column=1).value:
                    subject_cell = str(sheet.cell(row=1, column=1).value)
                    if "Предмет:" in subject_cell:
                        subject_name = subject_cell.replace("Предмет:", "").strip()

                group_name = ""
                if sheet.cell(row=2, column=1).value:
                    group_cell = str(sheet.cell(row=2, column=1).value)
                    if "Группа:" in group_cell:
                        group_name = group_cell.replace("Группа:", "").strip()

                teacher_name = self.get_teacher_from_sheet(sheet)

                # Поиск столбцов с итоговыми отметками
                t1_col = t2_col = t3_col = p1_col = p2_col = None
                for col in range(3, 100):
                    type_val = sheet.cell(row=6, column=col).value
                    if type_val == 'итог':
                        period_val = sheet.cell(row=4, column=col).value
                        if period_val == 'Т1':
                            t1_col = col
                        elif period_val == 'Т2':
                            t2_col = col
                        elif period_val == 'Т3':
                            t3_col = col
                        elif period_val == 'П1':
                            p1_col = col
                        elif period_val == 'П2':
                            p2_col = col

                trimestr_column = None
                if check_mode == 'Т1':
                    trimestr_column = t1_col
                elif check_mode == 'Т2':
                    trimestr_column = t2_col
                elif check_mode == 'Т3':
                    trimestr_column = t3_col
                elif check_mode == 'П1':
                    trimestr_column = p1_col
                elif check_mode == 'П2':
                    trimestr_column = p2_col

                # Находим все столбцы с отметками
                all_mark_columns = self.find_all_mark_columns(sheet, log_callback)

                # Фильтруем столбцы по периоду
                period_columns = []
                lessons_in_period = set()

                for col_info in all_mark_columns:
                    date_obj = col_info['date_obj']
                    if date_obj and period_start <= date_obj <= period_end:
                        period_columns.append(col_info['col'])
                        lessons_in_period.add(col_info['date_key'])

                lessons_count = len(lessons_in_period)

                if lessons_count == 0:
                    continue

                # Получаем количество уроков в неделю из учебного плана
                lessons_per_week = 0
                class_upper = class_name.upper()
                if class_upper in curriculum_data:
                    class_curriculum = curriculum_data[class_upper]
                    lesson_lower = subject_name.lower()

                    for subject, hours in class_curriculum.items():
                        if lesson_lower in subject.lower() or subject.lower() in lesson_lower:
                            lessons_per_week = hours
                            break

                min_marks_required = self.get_min_marks_required(lessons_per_week)

                # Определяем количество учеников
                start_row = 7
                students_count = 0

                while sheet.cell(row=start_row + students_count, column=2).value:
                    students_count += 1
                    if students_count > 50:
                        break

                # Обрабатываем каждого ученика
                for row in range(start_row, start_row + students_count):
                    student_cell = sheet.cell(row=row, column=2)
                    if not student_cell.value:
                        continue

                    student_name = str(student_cell.value).strip()

                    # Ключ для агрегации (ученик + предмет)
                    agg_key = (student_name, subject_name)
                    agg = student_data[agg_key]

                    # Ключ для полной статистики (класс + ученик + предмет)
                    full_key = (class_name, student_name, subject_name)
                    full = full_stats[full_key]

                    # Проверяем, не обрабатывали ли уже эту группу
                    group_key = f"{sheet_name}_{group_name}"
                    if group_key in agg['seen_groups']:
                        continue

                    agg['seen_groups'].add(group_key)
                    full['seen_groups'].add(group_key)

                    # Заполняем данные
                    agg['lessons_per_week'] = lessons_per_week
                    agg['min_marks_required'] = min_marks_required
                    agg['lessons_count'] = max(agg['lessons_count'], lessons_count)
                    agg['teacher_name'] = teacher_name

                    full['lessons_per_week'] = lessons_per_week
                    full['min_marks_required'] = min_marks_required
                    full['lessons_count'] = max(full['lessons_count'], lessons_count)
                    full['teacher_name'] = teacher_name
                    full['subject_name'] = subject_name

                    # Получаем итоговую отметку за период, если есть
                    if trimestr_column and agg['trimestr_mark'] is None:
                        trimestr_cell = sheet.cell(row=row, column=trimestr_column)
                        if trimestr_cell.value:
                            agg['trimestr_mark'] = str(trimestr_cell.value)

                    # Обрабатываем отметки за каждый урок в периоде
                    for col in period_columns:
                        cell_value = sheet.cell(row=row, column=col).value

                        # Парсим ячейку - теперь возвращает 4 значения
                        marks_list, absences, sickness, has_absence = self.parse_mark_cell(cell_value)

                        if has_absence:
                            agg['propuski_count'] += len(absences) + len(sickness)
                            full['propuski_count'] += len(absences) + len(sickness)

                            agg['sickness_count'] += len(sickness)
                            full['sickness_count'] += len(sickness)

                        # Обрабатываем отметки - каждый элемент: (mark, weight, is_control)
                        for mark_int, coeff, is_control in marks_list:
                            agg['marks_sum'] += mark_int * coeff
                            agg['marks_count'] += coeff
                            agg['marks_count_clear'] += 1

                            full['marks_sum'] += mark_int * coeff
                            full['marks_count'] += coeff
                            full['marks_count_clear'] += 1

            book.close()

            # Формируем результаты (только с замечаниями)
            all_results = []
            for (student_name, subject_name), agg in student_data.items():
                if agg['marks_count'] == 0 and agg['propuski_count'] == 0:
                    continue

                avg_mark = self.calculate_average(agg['marks_sum'], agg['marks_count'])
                correct_mark = self.get_grade_from_average(avg_mark)
                propuski_percent = round(agg['propuski_count'] / agg['lessons_count'] * 100, 0) if agg[
                                                                                                       'lessons_count'] > 0 else 0

                comments = []
                if agg['marks_count'] == 0:
                    comments.append('Нет отметок')
                elif agg['marks_count_clear'] < agg['min_marks_required'] and agg['min_marks_required'] > 0:
                    comments.append(f'Не хватает отметок: {agg["min_marks_required"] - agg["marks_count_clear"]}')
                elif avg_mark < 2.6 and agg['marks_count'] > 0:
                    comments.append('Выходит А/З')

                if comments:
                    result = {
                        'class': class_name,
                        'student': student_name,
                        'subject': subject_name,
                        'teacher': agg['teacher_name'],
                        'marks_count': agg['marks_count_clear'],
                        'avg_mark': avg_mark,
                        'trimestr_mark': agg['trimestr_mark'] if agg['trimestr_mark'] else '',
                        'correct_mark': correct_mark,
                        'lessons_count': agg['lessons_count'],
                        'lessons_per_week': agg['lessons_per_week'],
                        'min_marks_required': agg['min_marks_required'],
                        'propuski_count': agg['propuski_count'],
                        'sickness_count': agg['sickness_count'],
                        'propuski_percent': propuski_percent,
                        'comments': '\n'.join(comments)
                    }
                    all_results.append(result)

            # Формируем полную статистику
            all_stats = []
            for (class_name_key, student_name, subject_name), full in full_stats.items():
                if full['marks_count'] == 0 and full['propuski_count'] == 0:
                    continue

                avg_mark = self.calculate_average(full['marks_sum'], full['marks_count'])
                predicted_mark = self.get_grade_from_average(avg_mark)
                propuski_percent = round(full['propuski_count'] / full['lessons_count'] * 100, 0) if full[
                                                                                                         'lessons_count'] > 0 else 0

                stat = {
                    'class': class_name,
                    'student': student_name,
                    'subject': subject_name,
                    'teacher': full['teacher_name'],
                    'marks_count': full['marks_count_clear'],
                    'avg_mark': avg_mark,
                    'predicted_mark': predicted_mark,
                    'lessons_count': full['lessons_count'],
                    'lessons_per_week': full['lessons_per_week'],
                    'min_marks_required': full['min_marks_required'],
                    'propuski_count': full['propuski_count'],
                    'sickness_count': full['sickness_count'],
                    'propuski_percent': propuski_percent
                }
                all_stats.append(stat)

            return all_results, all_stats, None

        except Exception as e:
            import traceback
            traceback.print_exc()
            return [], [], f"Ошибка при обработке {file_path}: {str(e)}"

    def save_results_to_excel(self, results, stats, output_folder, grade, period_name):
        """Сохранение результатов проверки в Excel с двумя листами"""
        os.makedirs(output_folder, exist_ok=True)

        wb = Workbook()

        # === ЛИСТ 1: Замечания ===
        ws_errors = wb.active
        ws_errors.title = "Замечания"

        headers_errors = [
            'Класс', 'Ученик', 'Предмет', 'Учитель', 'Кол-во отметок', 'Средний балл',
            'Выставленная отметка', 'Правильная отметка', 'Кол-во уроков',
            'Уроков в неделю по плану', 'Мин. отметок требуется',
            'Кол-во пропусков', 'Из них по болезни', 'Процент пропусков', 'Комментарии'
        ]

        ws_errors.append(headers_errors)

        col_widths_errors = [8, 40, 40, 35, 10, 10, 12, 12, 10, 12, 12, 10, 10, 10, 50]
        for i, width in enumerate(col_widths_errors, 1):
            ws_errors.column_dimensions[get_column_letter(i)].width = width

        header_font = Font(bold=True, size=10)
        header_alignment = Alignment(
            horizontal='center',
            vertical='center',
            text_rotation=90,
            wrap_text=True
        )

        for col in range(1, len(headers_errors) + 1):
            cell = ws_errors.cell(row=1, column=col)
            cell.font = header_font
            cell.alignment = header_alignment
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

        ws_errors.auto_filter.ref = ws_errors.dimensions

        for result in results:
            ws_errors.append([
                result['class'],
                result['student'],
                result['subject'],
                result['teacher'],
                result['marks_count'],
                result['avg_mark'],
                result['trimestr_mark'],
                result['correct_mark'],
                result['lessons_count'],
                result['lessons_per_week'],
                result['min_marks_required'],
                result['propuski_count'],
                result['sickness_count'],
                result['propuski_percent'],
                result['comments']
            ])

        # === ЛИСТ 2: Полная статистика ===
        ws_stats = wb.create_sheet("Полная статистика")

        headers_stats = [
            'Класс', 'Ученик', 'Предмет', 'Учитель', 'Кол-во отметок', 'Средний балл',
            'Прогнозируемая отметка', 'Кол-во уроков', 'Уроков в неделю по плану',
            'Мин. отметок требуется', 'Кол-во пропусков', 'Из них по болезни', 'Процент пропусков'
        ]

        ws_stats.append(headers_stats)

        col_widths_stats = [8, 40, 40, 35, 10, 10, 12, 10, 12, 12, 10, 10, 10]
        for i, width in enumerate(col_widths_stats, 1):
            ws_stats.column_dimensions[get_column_letter(i)].width = width

        for col in range(1, len(headers_stats) + 1):
            cell = ws_stats.cell(row=1, column=col)
            cell.font = header_font
            cell.alignment = header_alignment
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

        ws_stats.auto_filter.ref = ws_stats.dimensions

        sorted_stats = sorted(stats, key=lambda x: (x['class'], x['student'], x['subject']))

        for stat in sorted_stats:
            ws_stats.append([
                stat['class'],
                stat['student'],
                stat['subject'],
                stat['teacher'],
                stat['marks_count'],
                stat['avg_mark'],
                stat['predicted_mark'],
                stat['lessons_count'],
                stat['lessons_per_week'],
                stat['min_marks_required'],
                stat['propuski_count'],
                stat['sickness_count'],
                stat['propuski_percent']
            ])

        if not stats:
            ws_stats.append(['', 'Нет данных за указанный период', '', '', '', '', '', '', '', '', '', '', ''])

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        data_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for row in range(2, len(results) + 2):
            for col in range(1, len(headers_errors) + 1):
                cell = ws_errors.cell(row=row, column=col)
                cell.alignment = data_alignment

                if col == 15:  # Комментарии
                    cell.alignment = left_alignment

                if row % 2 == 0:
                    cell.fill = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')

                cell.border = thin_border

        for row in range(2, len(sorted_stats) + 2):
            for col in range(1, len(headers_stats) + 1):
                cell = ws_stats.cell(row=row, column=col)
                cell.alignment = data_alignment

                if row % 2 == 0:
                    cell.fill = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')

                cell.border = thin_border

        safe_period_name = period_name.replace('/', '_').replace('\\', '_').replace(':', '_').replace('*', '_').replace(
            '?', '_').replace('"', '_').replace('<', '_').replace('>', '_').replace('|', '_')
        filename = f'РЕЗУЛЬТАТЫ ПРОВЕРКИ {safe_period_name} {grade}.xlsx'
        filepath = os.path.join(output_folder, filename)
        wb.save(filepath)

        return filepath

    def run_check(self, selected_journals, output_folder, curriculum_file, check_mode,
                  start_date=None, end_date=None, log_callback=None):
        """
        Запуск полной проверки для выбранных журналов
        """
        if log_callback:
            log_callback("\n📚 Загрузка учебного плана...")

        self.curriculum_data = self.load_curriculum_data(curriculum_file)

        if not self.curriculum_data:
            if log_callback:
                log_callback("❌ Ошибка загрузки учебного плана! Файл не найден или пуст.")
                log_callback(f"   Путь к файлу: {curriculum_file}")
            return 0, 0

        if log_callback:
            log_callback(f"✅ Учебный план загружен: {len(self.curriculum_data)} классов")
            # Покажем первые несколько классов для проверки
            sample_classes = list(self.curriculum_data.keys())[:10]
            log_callback(f"   Примеры классов в УП: {', '.join(sample_classes)}")

        start_date_obj = None
        end_date_obj = None
        period_name = ""

        if check_mode == 'custom':
            if not start_date or not end_date:
                if log_callback:
                    log_callback("❌ Для пользовательского режима необходимо указать даты начала и окончания!")
                return 0, 0

            start_date_obj = self.parse_custom_date(start_date)
            end_date_obj = self.parse_custom_date(end_date)

            if not start_date_obj or not end_date_obj:
                if log_callback:
                    log_callback("❌ Неверный формат даты! Используйте ДД.ММ.ГГГГ")
                return 0, 0

            if start_date_obj > end_date_obj:
                if log_callback:
                    log_callback("❌ Дата начала не может быть позже даты окончания!")
                return 0, 0

            period_name = f"custom_{start_date}_{end_date}"
        else:
            # Используем AcademicCalendar для получения имени периода
            _, _, period_name = AcademicCalendar.get_period_dates(check_mode)

        journals_by_level = {}
        for journal in selected_journals:
            level = journal['level']
            if level not in journals_by_level:
                journals_by_level[level] = []
            journals_by_level[level].append(journal)

        all_results = []
        all_stats = []
        total_files = len(selected_journals)
        error_files = 0

        for level, journals in journals_by_level.items():
            if log_callback:
                log_callback(f"\n📁 Проверка параллели {level}...")

            level_results = []
            level_stats = []

            for journal in journals:
                file_path = journal['path']

                if log_callback:
                    log_callback(f"\n  📄 {journal['class_name']}")

                results, stats, error = self.check_journal_file(
                    file_path,
                    self.curriculum_data,
                    check_mode,
                    start_date_obj,
                    end_date_obj,
                    log_callback
                )

                if error:
                    error_files += 1
                    if log_callback:
                        log_callback(f"    ⚠️ {error}")
                else:
                    level_results.extend(results)
                    level_stats.extend(stats)
                    if results and log_callback:
                        log_callback(f"    ✅ Найдено замечаний: {len(results)}")
                    if log_callback:
                        log_callback(f"    📊 Записей в статистике: {len(stats)}")

            if level_results or level_stats:
                all_results.extend(level_results)
                all_stats.extend(level_stats)

                # Создаем имя периода для файла
                if check_mode == 'custom':
                    if start_date and end_date:
                        start_clean = start_date.replace('.', '')
                        end_clean = end_date.replace('.', '')
                        file_period_name = f"с_{start_clean}_по_{end_clean}"
                    else:
                        file_period_name = "произвольный_период"
                else:
                    period_names = {
                        "Т1": "Триместр_1",
                        "Т2": "Триместр_2",
                        "Т3": "Триместр_3",
                        "П1": "Полугодие_1",
                        "П2": "Полугодие_2"
                    }
                    file_period_name = period_names.get(check_mode, check_mode)

                import re
                file_period_name = re.sub(r'[\\/*?:"<>|]', '_', file_period_name)

                output_path = self.save_results_to_excel(
                    level_results, level_stats, output_folder, level, file_period_name
                )
                if log_callback:
                    log_callback(f"\n  💾 Результаты сохранены: {output_path}")
            else:
                if log_callback:
                    log_callback(f"\n  ⚠️ Нет данных для сохранения по параллели {level}")

        if log_callback:
            log_callback(f"\n📊 Итоги проверки:")
            log_callback(f"  Обработано файлов: {total_files}")
            log_callback(f"  Найдено замечаний: {len(all_results)}")
            log_callback(f"  Записей в полной статистике: {len(all_stats)}")
            log_callback(f"  Файлов с ошибками: {error_files}")

        return total_files, len(all_results)